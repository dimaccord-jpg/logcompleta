"""Ingestão determinística de XLSX/CSV para o AgenteBI.

Lê a estrutura física do arquivo e monta uma preparação reutilizável.
Não classifica objetivo de negócio, não chama IA e não envia conteúdo
a provedor externo.

Os tetos desta etapa (3 arquivos, 10 MB, 200 mil células, 40 MB
descompactados) não afrouxam o validador documental já existente.
No ambiente atual o Excel continua limitado a 5 MB e o CSV a 2 MB,
porque esses são os tetos de `CleitonDocConfig`. O menor valor vence.

A gravação passa só por `criar_dataset`. Dataset que excede um único JSON,
mas cabe na ingestão, é fatiado em blocos da mesma versão. Se a governança
não autoriza o conteúdo intacto, nada é gravado. A decisão de envio à IA
não é tratada como autorização de persistência.
"""
from __future__ import annotations

import csv
import hashlib
import io
import logging
import math
import os
import time
import warnings
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time as hora_do_dia

from flask import request

from app.cleiton_doc_contracts import (
    ERROR_CORRUPTED_FILE,
    ERROR_DISABLED_TYPE,
    ERROR_EMPTY_FILE,
    ERROR_FILE_TOO_LARGE,
    ERROR_INVALID_EXTENSION,
    ERROR_INVALID_MIME,
    ERROR_UNSAFE_FILENAME,
    ERROR_UNSUPPORTED_TYPE,
    ERROR_UPLOAD_DISABLED,
)
from app.cleiton_doc_security import CleitonDocSecurityError, validate_upload_security
from app.models import AgenteBiDataset
from app.services.agente_bi_persistencia_service import (
    MAX_CHAVES,
    MAX_JSON_BYTES,
    MAX_LISTA,
    MAX_PROFUNDIDADE,
    MAX_TEXTO,
    AgenteBiErro,
    EscopoAgenteBi,
    calcular_impressao_dataset,
    criar_dataset,
)
from app.services.cleiton_doc_config_service import CleitonDocConfig, get_cleiton_doc_config

logger = logging.getLogger(__name__)

VERSAO_PREPARACAO = "preparar.v1"
ORIGEM_PREPARADA = "preparada"
DEPENDENCIA_ARMAZENAMENTO = "SCRUM-252"

TETO_ARQUIVOS = 3
TETO_BYTES = 10 * 1024 * 1024
TETO_CELULAS = 200_000
TETO_SEGUNDOS = 30
TETO_DESCOMPACTADO = 40 * 1024 * 1024
PADRAO_SEGUNDOS = 20
MAX_ENTRADAS_ZIP = 128
MAX_RAZAO_ZIP = 50
MAX_ABAS = 50
MAX_TEXTO_CELULA = 100_000
MAX_DIGITOS_NUMERO = 1024
MARGEM_CABECALHO = 3
MINIMO_CABECALHO = 4
LIMITE_CANDIDATOS_PUBLICADOS = 16
_SOBRECARGA_MULTIPART = 65_536
_ERROS_EXCEL = frozenset(
    {"#VALUE!", "#REF!", "#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#GETTING_DATA"}
)
_MAPA_SEGURANCA = {
    ERROR_FILE_TOO_LARGE: "limite_bytes",
    ERROR_EMPTY_FILE: "arquivo_vazio",
    ERROR_CORRUPTED_FILE: "arquivo_corrompido",
    ERROR_INVALID_EXTENSION: "arquivo_indevido",
    ERROR_INVALID_MIME: "arquivo_indevido",
    ERROR_UNSUPPORTED_TYPE: "arquivo_indevido",
    ERROR_UNSAFE_FILENAME: "arquivo_indevido",
    ERROR_DISABLED_TYPE: "tipo_nao_permitido",
    ERROR_UPLOAD_DISABLED: "tipo_nao_permitido",
}


class IngestaoErro(Exception):
    def __init__(self, codigo: str, status: int = 422):
        self.codigo = codigo
        self.status = status
        super().__init__(codigo)


@dataclass
class LimitesIngestao:
    max_arquivos: int
    max_bytes_xlsx: int
    max_bytes_csv: int
    max_bytes_total: int
    max_celulas: int
    max_segundos: float
    max_descompactado: int


@dataclass
class ArquivoBruto:
    nome: str
    mime: str | None
    conteudo: bytes


@dataclass
class Celula:
    tipo: str
    valor: object = None
    formula: str | None = None
    cache_estado: str | None = None
    tipo_cache: str | None = None
    bruto: str | None = None
    ordem_data: str | None = None


@dataclass
class _Cota:
    max_celulas: int
    max_segundos: float
    relogio: callable
    total: int = 0
    inicio: float = field(init=False)

    def __post_init__(self) -> None:
        self.inicio = self.relogio()

    def somar(self, quantidade: int = 1) -> None:
        if quantidade < 0:
            raise IngestaoErro("limite_processamento")
        self.total += quantidade
        if self.total > self.max_celulas:
            raise IngestaoErro("limite_celulas")
        self.verificar()

    def verificar(self) -> None:
        if self.relogio() - self.inicio > self.max_segundos:
            raise IngestaoErro("limite_tempo")


def resolver_limites(cfg: CleitonDocConfig | None = None) -> LimitesIngestao:
    """O teto configurável nunca ultrapassa o menor limite já vigente."""
    cfg = cfg or get_cleiton_doc_config()
    arquivos = min(_teto("AGENTE_BI_INGESTAO_MAX_ARQUIVOS", TETO_ARQUIVOS, TETO_ARQUIVOS), TETO_ARQUIVOS)
    bytes_teto = _teto("AGENTE_BI_INGESTAO_MAX_BYTES", TETO_BYTES, TETO_BYTES)
    return LimitesIngestao(
        max_arquivos=arquivos,
        max_bytes_xlsx=min(bytes_teto, int(cfg.excel_max_bytes)),
        max_bytes_csv=min(bytes_teto, int(cfg.csv_max_bytes)),
        max_bytes_total=min(int(cfg.session_max_bytes), arquivos * TETO_BYTES),
        max_celulas=_teto("AGENTE_BI_INGESTAO_MAX_CELULAS", TETO_CELULAS, TETO_CELULAS),
        max_segundos=float(_teto("AGENTE_BI_INGESTAO_MAX_SEGUNDOS", PADRAO_SEGUNDOS, TETO_SEGUNDOS)),
        max_descompactado=min(
            int(cfg.excel_max_bytes) * 20,
            _teto("AGENTE_BI_INGESTAO_MAX_DESCOMPACTADO", TETO_DESCOMPACTADO, TETO_DESCOMPACTADO),
        ),
    )


def ingerir_requisicao(escopo: EscopoAgenteBi, *, relogio=time.monotonic) -> dict:
    arquivos, chave, limites = _ler_requisicao()
    return ingerir(escopo, arquivos, chave=chave, limites=limites, relogio=relogio)


def ingerir(
    escopo: EscopoAgenteBi,
    arquivos: list[ArquivoBruto],
    *,
    chave: str | None,
    limites: LimitesIngestao | None = None,
    relogio=time.monotonic,
    persistir: bool = True,
) -> dict:
    limites = limites or resolver_limites()
    if not arquivos:
        raise IngestaoErro("payload_invalido")
    if len(arquivos) > limites.max_arquivos:
        raise IngestaoErro("limite_arquivos")
    if chave is not None and not _chave_valida(chave):
        raise IngestaoErro("payload_invalido")

    preparados = []
    dados_tipados: list[dict] = []
    ambiguidades: list[dict] = []
    alertas: list[str] = []
    cota = _Cota(limites.max_celulas, limites.max_segundos, relogio)
    cota.verificar()
    total_bytes = 0
    for indice, arquivo in enumerate(arquivos):
        cota.verificar()
        total_bytes += len(arquivo.conteudo)
        if total_bytes > limites.max_bytes_total:
            raise IngestaoErro("limite_bytes")
        preparado = _preparar_arquivo(indice, arquivo, limites, cota, dados_tipados, ambiguidades)
        preparados.append(preparado)
        alertas.extend(preparado["alertas"])

    cota.verificar()
    estrutura = _publicar(preparados, alertas, ambiguidades)
    perfil = {
        "arquivos": len(preparados),
        "abas": sum(len(item["abas"]) for item in preparados),
        "regioes": sum(len(aba["regioes"]) for item in preparados for aba in item["abas"]),
        "celulas": cota.total,
        "linhas_elegiveis": sum(
            regiao["linhas_elegiveis"]
            for item in preparados
            for aba in item["abas"]
            for regiao in aba["regioes"]
        ),
        "formulas": sum(
            regiao["formulas"]
            for item in preparados
            for aba in item["abas"]
            for regiao in aba["regioes"]
        ),
        "ambigua": bool(ambiguidades),
    }
    metadados = _metadados(limites)
    cota.verificar()
    chave_final = chave or _chave_de(arquivos)
    if not persistir:
        logger.info(
            "evento=agente_bi_ingestao usuario_id=%s arquivos=%s celulas=%s persistido=0 motivo=preparacao_em_memoria",
            escopo.usuario_id,
            len(preparados),
            cota.total,
        )
        return {
            "ok": True,
            "persistido": False,
            "reutilizado": False,
            "motivo": "preparacao_em_memoria",
            "dependencia": None,
            "preparacao": {
                "estrutura": estrutura,
                "perfil": perfil,
                "dados_tipados": dados_tipados,
                "metadados_preparacao": metadados,
            },
            "dataset": None,
        }
    persistido, reutilizado, motivo, dependencia, dataset = _persistir_se_permitido(
        escopo,
        chave_final,
        estrutura,
        perfil,
        dados_tipados,
        metadados,
        cota,
    )
    logger.info(
        "evento=agente_bi_ingestao usuario_id=%s arquivos=%s celulas=%s persistido=%s motivo=%s",
        escopo.usuario_id,
        len(preparados),
        cota.total,
        int(persistido),
        motivo or "-",
    )
    return {
        "ok": True,
        "persistido": persistido,
        "reutilizado": reutilizado,
        "motivo": motivo,
        "dependencia": dependencia,
        "preparacao": {
            "estrutura": estrutura,
            "perfil": perfil,
            "dados_tipados": dados_tipados,
            "metadados_preparacao": metadados,
        },
        "dataset": dataset,
    }


def _persistir_se_permitido(escopo, chave, estrutura, perfil, dados, metadados, cota: _Cota):
    cota.verificar()
    origem_ref = "prep-" + _sha256_bytes(_material_bytes(estrutura, perfil, dados, metadados))[:32]
    existente = _versao_da_origem(escopo, chave, origem_ref)
    if existente is not None:
        return True, True, None, None, _identidade(existente)
    ultima = _ultima_versao(escopo, chave)
    versao = 1 if ultima is None else int(ultima.versao) + 1
    anterior_id = None if ultima is None else int(ultima.id)
    cota.verificar()
    try:
        impressao = calcular_impressao_dataset(
            chave=chave,
            origem=ORIGEM_PREPARADA,
            origem_ref=origem_ref,
            versao=versao,
            estrutura=estrutura,
            perfil=perfil,
            dados_tipados=dados,
            metadados_preparacao=metadados,
        )
    except (TypeError, ValueError):
        return False, False, "limite_contrato_persistencia", None, None

    cota.verificar()
    if (
        _fora_do_contrato(estrutura)
        or _fora_do_contrato(perfil)
        or _fora_do_contrato(metadados)
        or not _dados_cabem_na_segmentacao(dados)
    ):
        return False, False, "limite_contrato_persistencia", None, None

    payload = {
        "chave": chave,
        "origem": ORIGEM_PREPARADA,
        "origem_ref": origem_ref,
        "versao": versao,
        "impressao": impressao,
        "estrutura": estrutura,
        "perfil": perfil,
        "dados_tipados": dados,
        "metadados_preparacao": metadados,
    }
    if anterior_id is not None:
        payload["versao_anterior_id"] = anterior_id
    try:
        criado = criar_dataset(escopo, payload, ponto_de_controle=cota.verificar)
    except AgenteBiErro as exc:
        if exc.codigo == "conteudo_nao_permitido":
            return False, False, exc.codigo, DEPENDENCIA_ARMAZENAMENTO, None
        if exc.codigo == "payload_invalido":
            return False, False, "limite_contrato_persistencia", None, None
        raise
    return True, False, None, None, {
        "id": criado["id"],
        "chave": criado["chave"],
        "versao": criado["versao"],
        "impressao": criado["impressao"],
        "versao_anterior_id": criado["versao_anterior_id"],
        "origem": criado["origem"],
        "origem_ref": criado["origem_ref"],
    }


def persistir_preparacao(escopo: EscopoAgenteBi, chave: str, preparacao: dict, *, relogio=time.monotonic) -> dict:
    """Grava uma preparação já lida, na chave autorizada, pelo versionamento existente."""
    if not isinstance(preparacao, dict):
        raise IngestaoErro("payload_invalido")
    limites = resolver_limites()
    cota = _Cota(limites.max_celulas, limites.max_segundos, relogio)
    persistido, reutilizado, motivo, dependencia, dataset = _persistir_se_permitido(
        escopo,
        chave,
        preparacao.get("estrutura"),
        preparacao.get("perfil"),
        preparacao.get("dados_tipados"),
        preparacao.get("metadados_preparacao"),
        cota,
    )
    if not persistido or dataset is None:
        raise IngestaoErro(motivo or "limite_contrato_persistencia")
    return {
        "persistido": True,
        "reutilizado": bool(reutilizado),
        "motivo": motivo,
        "dependencia": dependencia,
        "dataset": dataset,
    }


def _consulta_chave(escopo: EscopoAgenteBi, chave: str):
    return AgenteBiDataset.query.filter_by(
        conta_id=escopo.conta_id,
        franquia_id=escopo.franquia_id,
        usuario_id=escopo.usuario_id,
        chave=chave,
    )


def _ultima_versao(escopo: EscopoAgenteBi, chave: str):
    return _consulta_chave(escopo, chave).order_by(AgenteBiDataset.versao.desc(), AgenteBiDataset.id.desc()).first()


def _versao_da_origem(escopo: EscopoAgenteBi, chave: str, origem_ref: str):
    return (
        _consulta_chave(escopo, chave)
        .filter_by(origem_ref=origem_ref)
        .order_by(AgenteBiDataset.versao.desc(), AgenteBiDataset.id.desc())
        .first()
    )


def _identidade(linha: AgenteBiDataset) -> dict:
    return {
        "id": linha.id,
        "chave": linha.chave,
        "versao": linha.versao,
        "impressao": linha.impressao,
        "versao_anterior_id": linha.versao_anterior_id,
        "origem": linha.origem,
        "origem_ref": linha.origem_ref,
    }


def _preparar_arquivo(indice, arquivo, limites, cota, dados_tipados, ambiguidades) -> dict:
    cota.verificar()
    seguranca = _validar_seguranca(arquivo, limites)
    cota.verificar()
    ext = seguranca.extension
    if len(arquivo.conteudo) > (limites.max_bytes_xlsx if ext == ".xlsx" else limites.max_bytes_csv):
        raise IngestaoErro("limite_bytes")
    if ext == ".xlsx":
        abas = _ler_xlsx(arquivo.conteudo, limites, cota)
    else:
        abas = [_ler_csv(arquivo.conteudo, cota)]
    _numerar_dados(indice, abas, dados_tipados, ambiguidades)
    cota.verificar()
    alertas = _unicos([codigo for aba in abas for codigo in aba["alertas"]])
    return {
        "indice": indice,
        "nome": seguranca.safe_name,
        "extensao": ext,
        "mime": seguranca.mime_type,
        "bytes": seguranca.size_bytes,
        "sha256": hashlib.sha256(arquivo.conteudo).hexdigest(),
        "contem_grafico": any(aba.get("contem_grafico", False) for aba in abas),
        "abas": abas,
        "alertas": alertas,
    }


def _validar_seguranca(arquivo: ArquivoBruto, limites: LimitesIngestao):
    ext = _extensao(arquivo.nome)
    teto = limites.max_bytes_xlsx if ext == ".xlsx" else limites.max_bytes_csv
    if ext not in {".csv", ".xlsx"}:
        raise IngestaoErro("arquivo_indevido")
    if len(arquivo.nome) > 200:
        raise IngestaoErro("arquivo_indevido")
    if len(arquivo.conteudo) > teto:
        raise IngestaoErro("limite_bytes")
    try:
        return validate_upload_security(
            display_name=arquivo.nome,
            file_bytes=arquivo.conteudo,
            mime_type=arquivo.mime,
            extension=ext,
        )
    except CleitonDocSecurityError as exc:
        raise IngestaoErro(_MAPA_SEGURANCA.get(exc.error_code, "arquivo_indevido")) from exc


def _ler_csv(conteudo: bytes, cota: _Cota) -> dict:
    if _parece_binario(conteudo):
        raise IngestaoErro("arquivo_indevido")
    texto, encoding = _decodificar(conteudo)
    separador, ambiguo = _detectar_separador(texto)
    alertas: list[str] = []
    ambiguidades: list[dict] = []
    if ambiguo:
        alertas.append("separador_ambiguo")
        ambiguidades.append({"codigo": "separador_ambiguo", "aba": 0, "regiao": None})
        return _aba_csv(encoding, None, [], alertas, ambiguidades, 0, 0)
    grade = _grade_csv(texto, separador, cota)
    regioes = _montar_regioes(grade, cota, reinterpretar_texto=True, separador=separador)
    if not regioes:
        alertas.append("sem_regiao_tabular")
    if len(regioes) > 1:
        alertas.append("regioes_multiplas")
    alertas.extend(_unicos(codigo for regiao in regioes for codigo in regiao["alertas"]))
    for regiao in regioes:
        ambiguidades.extend(regiao.pop("_ambiguidades"))
    cota.verificar()
    return _aba_csv(
        encoding,
        separador if separador != "\x1f" else None,
        regioes,
        _unicos(alertas),
        ambiguidades,
        grade["linhas"],
        grade["colunas"],
    )


def _aba_csv(encoding, separador, regioes, alertas, ambiguidades, linhas, colunas) -> dict:
    return {
        "indice": 0,
        "nome": "csv",
        "estado": "visivel",
        "oculta": False,
        "indicio_auxiliar": False,
        "indicio_apresentacao": False,
        "contem_desenho": False,
        "contem_grafico": False,
        "encoding": encoding,
        "separador": {",": "virgula", ";": "ponto_e_virgula", "\t": "tabulacao", "|": "pipe"}.get(separador),
        "celulas_mescladas": 0,
        "linhas": linhas,
        "colunas": colunas,
        "regioes": regioes,
        "alertas": alertas,
        "_ambiguidades": ambiguidades,
    }


def _ler_xlsx(conteudo: bytes, limites: LimitesIngestao, cota: _Cota) -> list[dict]:
    cota.verificar()
    inspecao = _inspecionar_zip(conteudo, limites)
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
        from openpyxl.worksheet._reader import WorkSheetParser
    except ImportError as exc:
        raise IngestaoErro("arquivo_corrompido") from exc

    formulas = _mapa_xlsx(
        conteudo,
        data_only=False,
        cota=cota,
        load_workbook=load_workbook,
        invalido=InvalidFileException,
        parser_cls=WorkSheetParser,
        desenhos=inspecao["desenhos"],
    )
    caches = _mapa_xlsx(
        conteudo,
        data_only=True,
        cota=cota,
        load_workbook=load_workbook,
        invalido=InvalidFileException,
        parser_cls=WorkSheetParser,
        desenhos=set(),
        cobrar=False,
    )
    _aplicar_caches(formulas, caches)
    if len(formulas) > MAX_ABAS:
        raise IngestaoErro("limite_processamento")
    abas = []
    for indice, aba in enumerate(formulas):
        grade = {
            "celulas": aba["celulas"],
            "linhas": aba["linhas"],
            "colunas": aba["colunas"],
        }
        regioes = _montar_regioes(grade, cota, reinterpretar_texto=False, separador=None)
        alertas = list(aba["alertas"])
        if not regioes:
            alertas.append("sem_regiao_tabular")
        if len(regioes) > 1:
            alertas.append("regioes_multiplas")
        if aba["oculta"]:
            alertas.append("aba_oculta")
        elegiveis = sum(regiao["linhas_elegiveis"] for regiao in regioes)
        apresentacao = bool(aba["contem_desenho"] and elegiveis < 15)
        if apresentacao:
            alertas.append("indicio_apresentacao")
        alertas.extend(codigo for regiao in regioes for codigo in regiao["alertas"])
        ambiguidades = []
        for regiao in regioes:
            ambiguidades.extend(regiao.pop("_ambiguidades"))
        abas.append(
            {
                "indice": indice,
                "nome": aba["nome"],
                "estado": aba["estado"],
                "oculta": aba["oculta"],
                "indicio_auxiliar": aba["oculta"],
                "indicio_apresentacao": apresentacao,
                "contem_desenho": aba["contem_desenho"],
                "contem_grafico": inspecao["contem_grafico"],
                "encoding": None,
                "separador": None,
                "celulas_mescladas": aba["mescladas"],
                "linhas": aba["linhas"],
                "colunas": aba["colunas"],
                "regioes": regioes,
                "alertas": _unicos(alertas),
                "_ambiguidades": ambiguidades,
            }
        )
    cota.verificar()
    return abas


def _mapa_xlsx(conteudo, *, data_only, cota, load_workbook, invalido, parser_cls, desenhos, cobrar=True):
    abas = []
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            workbook = load_workbook(
                io.BytesIO(conteudo),
                read_only=True,
                data_only=data_only,
                keep_vba=False,
                keep_links=False,
            )
    except (zipfile.BadZipFile, invalido, OSError, ValueError) as exc:
        logger.info(
            "evento=agente_bi_ingestao_recusada codigo=arquivo_corrompido tipo=%s",
            exc.__class__.__name__,
        )
        raise IngestaoErro("arquivo_corrompido") from exc
    try:
        if len(workbook.worksheets) > MAX_ABAS:
            raise IngestaoErro("limite_processamento")
        for ws in workbook.worksheets:
            cota.verificar()
            estado, oculta = _estado_aba(ws.sheet_state)
            celulas: dict[tuple[int, int], Celula] = {}
            mescladas = 0
            ultima = 0
            max_linha = 0
            max_coluna = 0
            with workbook._archive.open(ws._worksheet_path) as origem:
                parser = parser_cls(
                    origem,
                    ws._shared_strings,
                    data_only=data_only,
                    epoch=workbook.epoch,
                    date_formats=workbook._date_formats,
                    timedelta_formats=workbook._timedelta_formats,
                )
                for row_idx, cells in parser.parse():
                    if row_idx < 1 or row_idx > 1_048_576:
                        raise IngestaoErro("arquivo_corrompido")
                    if not cobrar:
                        cota.verificar()
                    if cobrar and ultima and row_idx - ultima > 1:
                        cota.somar(row_idx - ultima - 1)
                    ultima = row_idx
                    if not cells and cobrar:
                        cota.somar(1)
                    for cell in cells:
                        coluna = cell["column"]
                        if coluna < 1 or coluna > 16384:
                            raise IngestaoErro("arquivo_corrompido")
                        montada = _celula_xlsx(cell, data_only=data_only)
                        if cobrar and montada is not None:
                            cota.somar(1)
                        if montada is None:
                            continue
                        celulas[(row_idx, coluna)] = montada
                        max_linha = max(max_linha, row_idx)
                        max_coluna = max(max_coluna, coluna)
                if parser.merged_cells is not None:
                    mescladas = len(parser.merged_cells.mergeCell)
            if not data_only:
                abas.append(
                    {
                        "nome": str(ws.title)[:200],
                        "estado": estado,
                        "oculta": oculta,
                        "contem_desenho": _normalizar_zip(ws._worksheet_path) in desenhos,
                        "mescladas": mescladas,
                        "celulas": celulas,
                        "linhas": max_linha,
                        "colunas": max_coluna,
                        "alertas": ["celulas_mescladas"] if mescladas else [],
                    }
                )
            else:
                abas.append({"celulas": celulas})
    except IngestaoErro:
        raise
    except (zipfile.BadZipFile, invalido, OSError, ValueError, KeyError) as exc:
        logger.info(
            "evento=agente_bi_ingestao_recusada codigo=arquivo_corrompido tipo=%s",
            exc.__class__.__name__,
        )
        raise IngestaoErro("arquivo_corrompido") from exc
    finally:
        workbook.close()
    return abas


def _aplicar_caches(formulas, caches) -> None:
    if len(formulas) != len(caches):
        raise IngestaoErro("arquivo_corrompido")
    for aba, cache in zip(formulas, caches):
        for coordenada, celula in aba["celulas"].items():
            if celula.tipo != "formula":
                continue
            observado = cache["celulas"].get(coordenada)
            if observado is None or observado.tipo in {"ausente", "formula"}:
                celula.cache_estado = "ausente"
                continue
            celula.cache_estado = "presente"
            celula.valor = observado.valor
            celula.tipo_cache = observado.tipo


def _celula_xlsx(cell: dict, *, data_only: bool) -> Celula | None:
    valor = cell.get("value")
    tipo = cell.get("data_type")
    if not data_only and tipo == "f":
        formula = _texto_formula(valor)
        if formula is None:
            raise IngestaoErro("limite_processamento")
        return Celula(tipo="formula", formula=formula, cache_estado="ausente")
    if valor is None:
        return None
    return _celula_observada(tipo, valor)


def _celula_observada(tipo_dado, valor) -> Celula | None:
    if isinstance(valor, str) and valor in _ERROS_EXCEL or tipo_dado == "e":
        return Celula(tipo="erro", valor=valor if isinstance(valor, str) else None)
    if isinstance(valor, bool):
        return Celula(tipo="booleano", valor=valor)
    if isinstance(valor, datetime):
        if valor.time() == hora_do_dia(0, 0, 0):
            return Celula(tipo="data", valor=valor.date().isoformat())
        return Celula(tipo="data", valor=valor.isoformat(sep=" "))
    if isinstance(valor, date):
        return Celula(tipo="data", valor=valor.isoformat())
    if isinstance(valor, hora_do_dia):
        return Celula(tipo="hora", valor=valor.isoformat())
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        if isinstance(valor, float) and not math.isfinite(valor):
            return Celula(tipo="erro", valor=None)
        if isinstance(valor, float) and valor.is_integer() and abs(valor) < 10**15:
            valor = int(valor)
        return Celula(tipo="numero", valor=valor)
    if isinstance(valor, str):
        if len(valor) > MAX_TEXTO_CELULA:
            raise IngestaoErro("limite_processamento")
        if valor.strip() == "":
            return None
        return Celula(tipo="texto", valor=valor)
    texto = _texto_rico(valor)
    if texto is None:
        return Celula(tipo="nao_suportado")
    if texto.strip() == "":
        return None
    return Celula(tipo="texto", valor=texto)


def _montar_regioes(grade, cota: _Cota, *, reinterpretar_texto: bool, separador: str | None) -> list[dict]:
    celulas = grade["celulas"]
    if not celulas:
        return []
    linhas = sorted({linha for linha, _coluna in celulas})
    bandas: list[list[int]] = []
    atual = [linhas[0]]
    for linha in linhas[1:]:
        if linha - atual[-1] > 1:
            bandas.append(atual)
            atual = [linha]
        else:
            atual.append(linha)
    bandas.append(atual)
    if len(bandas) > 1000:
        raise IngestaoErro("limite_processamento")
    regioes = []
    for banda in bandas:
        regioes.append(_regiao_da_banda(banda, celulas, cota, reinterpretar_texto, separador, len(regioes)))
        cota.verificar()
    return regioes


def _regiao_da_banda(banda, celulas, cota, reinterpretar_texto, separador, indice) -> dict:
    conjunto = set(banda)
    colunas = sorted(coluna for linha, coluna in celulas if linha in conjunto)
    linha_inicio, linha_fim = banda[0], banda[-1]
    coluna_inicio, coluna_fim = colunas[0], colunas[-1]
    altura = linha_fim - linha_inicio + 1
    largura = coluna_fim - coluna_inicio + 1
    ocupadas = sum(1 for linha, _coluna in celulas if linha in conjunto)
    buracos = altura * largura - ocupadas
    if buracos < 0 or altura * largura > cota.max_celulas:
        raise IngestaoErro("limite_celulas")
    cota.somar(buracos)
    matriz = []
    for linha in range(linha_inicio, linha_fim + 1):
        faixa = []
        for coluna in range(coluna_inicio, coluna_fim + 1):
            faixa.append(celulas.get((linha, coluna), Celula(tipo="ausente")))
        matriz.append(faixa)
    grupos = _grupos_laterais(colunas)
    divisao = None
    alertas = []
    ambiguidades = []
    if len(grupos) > 1:
        divisao = {
            "escolha": None,
            "candidatos": [{"coluna_inicio": grupo[0], "coluna_fim": grupo[-1]} for grupo in grupos],
        }
        alertas.append("divisao_lateral_ambigua")
        ambiguidades.append({"codigo": "divisao_lateral_ambigua", "regiao": indice})
    cabecalho, preambulo, repetidos = _decidir_cabecalho(
        matriz, linha_inicio, reinterpretar_texto, separador
    )
    if cabecalho["deteccao"] == "ambigua":
        alertas.append("cabecalho_ambiguo")
        ambiguidades.append({"codigo": "cabecalho_ambiguo", "regiao": indice})
    elif cabecalho["deteccao"] == "ausente":
        alertas.append("cabecalho_ausente")
    if reinterpretar_texto:
        _reinterpretar_dados(
            matriz, linha_inicio, cabecalho, preambulo, repetidos, separador, alertas, ambiguidades, indice
        )
    papeis = _classificar_linhas(matriz, linha_inicio, cabecalho, preambulo, repetidos)
    if papeis["pre_preenchidas"]:
        alertas.append("linha_pre_preenchida")
    if papeis["invalidas"]:
        alertas.append("linha_invalida")
    if papeis["repetidos"]:
        alertas.append("cabecalho_repetido")
    tipos = _tipos_colunas(matriz, papeis["elegiveis"])
    if "misto" in tipos:
        alertas.append("tipo_misto")
    if _coluna_parcial(matriz, papeis["elegiveis"]):
        alertas.append("coluna_parcialmente_vazia")
    if _linha_parcial(matriz, papeis["elegiveis"]):
        alertas.append("linha_parcialmente_vazia")
    formulas = _resumo_formulas(matriz, linha_inicio, coluna_inicio, set(papeis["elegiveis"]))
    if formulas["sem_cache"]:
        alertas.append("formula_sem_cache")
    preservadas = _preservar_fora_da_analise(
        matriz, linha_inicio, coluna_inicio, cabecalho, preambulo, repetidos, papeis
    )
    return {
        "indice": indice,
        "linha_inicio": linha_inicio,
        "linha_fim": linha_fim,
        "coluna_inicio": coluna_inicio,
        "coluna_fim": coluna_fim,
        "cabecalho": cabecalho,
        "preambulo": preambulo,
        "divisao_lateral": divisao,
        "tipos_fisicos": tipos,
        "linhas_ocupadas": len(banda),
        "linhas_elegiveis": len(papeis["elegiveis"]),
        "linhas_vazias": papeis["vazias"],
        "linhas_invalidas": len(papeis["invalidas"]),
        "linhas_pre_preenchidas": len(papeis["pre_preenchidas"]),
        "linhas_cabecalho_repetido": len(papeis["repetidos"]),
        "formulas": formulas["total"],
        "formulas_com_cache": formulas["com_cache"],
        "formulas_sem_cache": formulas["sem_cache"],
        "formulas_observadas": formulas["observadas"],
        "celulas_preservadas": preservadas,
        "alertas": _unicos(alertas),
        "_elegiveis": [
            {"linha": linha_inicio + deslocamento, "valores": [_valor_json(celula) for celula in matriz[deslocamento]]}
            for deslocamento in papeis["elegiveis"]
        ],
        "_ambiguidades": ambiguidades,
    }


def _decidir_cabecalho(matriz, linha_inicio, reinterpretar, separador):
    pontuacoes = []
    for indice, linha in enumerate(matriz):
        proxima = matriz[indice + 1] if indice + 1 < len(matriz) else None
        pontuacoes.append(_pontuar(linha, len(linha), proxima, reinterpretar, separador))
    candidatos = []
    vistos = set()
    for indice, pontos in enumerate(pontuacoes):
        if pontos < MINIMO_CABECALHO:
            continue
        rotulos = _rotulos(matriz[indice], reinterpretar, separador)
        chave_rotulo = tuple(rotulos)
        if chave_rotulo in vistos:
            continue
        vistos.add(chave_rotulo)
        candidatos.append(
            {
                "linha": linha_inicio + indice,
                "indice": indice,
                "pontuacao": pontos,
                "rotulos": [item if isinstance(item, str) or item is None else None for item in rotulos],
            }
        )
    if not candidatos:
        return (
            _cabecalho_publico("ausente", "nenhuma", None, []),
            [],
            set(),
        )
    candidatos.sort(key=lambda item: (-item["pontuacao"], item["indice"]))
    melhor = candidatos[0]["pontuacao"]
    segundo = candidatos[1]["pontuacao"] if len(candidatos) > 1 else None
    if segundo is not None and melhor - segundo < MARGEM_CABECALHO:
        return (
            _cabecalho_publico("ambigua", "nenhuma", None, candidatos),
            [],
            set(),
        )
    escolhido = candidatos[0]
    confianca = "alta" if escolhido["pontuacao"] >= 6 else "media"
    repetidos = set()
    rotulo_escolhido = _rotulos(matriz[escolhido["indice"]], reinterpretar, separador)
    for indice in range(escolhido["indice"] + 1, len(matriz)):
        if _rotulos(matriz[indice], reinterpretar, separador) == rotulo_escolhido:
            repetidos.add(indice)
    preambulo = []
    for indice in range(escolhido["indice"]):
        if any(celula.tipo != "ausente" for celula in matriz[indice]):
            preambulo.append(
                {
                    "linha": linha_inicio + indice,
                    "rotulos": [
                        item if isinstance(item, str) else None
                        for item in _rotulos(matriz[indice], reinterpretar, separador)
                    ],
                }
            )
    return (
        _cabecalho_publico("inequivoca", confianca, escolhido["linha"], candidatos),
        preambulo,
        repetidos,
    )


def _indices_bloqueados(linha_inicio, cabecalho, preambulo, repetidos) -> set[int]:
    bloqueadas = set(repetidos)
    if cabecalho["linha"] is not None:
        bloqueadas.update(range((cabecalho["linha"] - linha_inicio) + 1))
    for item in preambulo:
        bloqueadas.add(item["linha"] - linha_inicio)
    return bloqueadas


def _reinterpretar_dados(matriz, linha_inicio, cabecalho, preambulo, repetidos, separador, alertas, ambiguidades, regiao):
    bloqueadas = _indices_bloqueados(linha_inicio, cabecalho, preambulo, repetidos)
    for indice, linha in enumerate(matriz):
        if indice in bloqueadas:
            continue
        for celula in linha:
            if celula.tipo != "texto":
                continue
            _classificar_token(celula, separador)
    _resolver_datas(matriz, bloqueadas, alertas, ambiguidades, regiao)
    _resolver_numeros_ambiguos(matriz, alertas, ambiguidades, regiao)


def _preservar_fora_da_analise(matriz, linha_inicio, coluna_inicio, cabecalho, preambulo, repetidos, papeis) -> list[dict]:
    """Célula observada que não é registro analítico permanece, com lugar e papel.

    Lacuna não vira valor. Cabeçalho, preâmbulo e repetição não entram em
    `dados_tipados`.
    """
    elegiveis = set(papeis["elegiveis"])
    invalidas = set(papeis["invalidas"])
    pre_preenchidas = set(papeis["pre_preenchidas"])
    pre_linhas = {item["linha"] - linha_inicio for item in preambulo}
    cab_idx = None if cabecalho["linha"] is None else cabecalho["linha"] - linha_inicio
    saida = []
    for indice, linha in enumerate(matriz):
        if indice in elegiveis:
            continue
        if indice == cab_idx:
            papel = "cabecalho"
        elif indice in repetidos:
            papel = "cabecalho_repetido"
        elif indice in pre_linhas:
            papel = "preambulo"
        elif indice in invalidas or indice in pre_preenchidas:
            papel = "nao_classificado"
        else:
            continue
        for coluna, celula in enumerate(linha):
            if celula.tipo == "ausente":
                continue
            saida.append(
                {
                    "linha": linha_inicio + indice,
                    "coluna": coluna_inicio + coluna,
                    "papel": papel,
                    "celula": _valor_json(celula),
                }
            )
    return saida


def _classificar_linhas(matriz, linha_inicio, cabecalho, preambulo, repetidos):
    bloqueadas = _indices_bloqueados(linha_inicio, cabecalho, preambulo, repetidos)
    elegiveis = []
    invalidas = []
    pre_preenchidas = []
    vazias = 0
    for indice, linha in enumerate(matriz):
        if indice in bloqueadas:
            continue
        papel = _papel_linha(linha)
        if papel == "elegivel":
            elegiveis.append(indice)
        elif papel == "invalida":
            invalidas.append(indice)
        elif papel == "pre_preenchida":
            pre_preenchidas.append(indice)
        else:
            vazias += 1
    return {
        "elegiveis": elegiveis,
        "invalidas": invalidas,
        "pre_preenchidas": pre_preenchidas,
        "repetidos": repetidos,
        "vazias": vazias,
        "preambulo": preambulo,
    }


def _papel_linha(linha) -> str:
    observacao = False
    formula = False
    so_formula_sem_valor = True
    erro = False
    for celula in linha:
        if celula.tipo == "ausente":
            continue
        if celula.tipo == "formula":
            formula = True
            if celula.cache_estado == "presente" and celula.valor not in (None, ""):
                observacao = True
                so_formula_sem_valor = False
            elif celula.cache_estado == "presente":
                so_formula_sem_valor = False
            continue
        so_formula_sem_valor = False
        if celula.tipo == "erro":
            erro = True
            continue
        observacao = True
    if observacao:
        return "elegivel"
    if formula and so_formula_sem_valor:
        return "pre_preenchida"
    if erro:
        return "invalida"
    return "vazia"


def _tipos_colunas(matriz, elegiveis) -> list[str]:
    if not matriz:
        return []
    tipos = []
    for coluna in range(len(matriz[0])):
        vistos = []
        for indice in elegiveis:
            celula = matriz[indice][coluna]
            tipo = _tipo_observado(celula)
            if tipo == "ausente" or tipo in vistos:
                continue
            vistos.append(tipo)
        if not vistos:
            tipos.append("vazio")
        elif len(vistos) == 1:
            tipos.append(vistos[0])
        else:
            tipos.append("misto")
    return tipos


def _tipo_observado(celula: Celula) -> str:
    if celula.tipo == "formula":
        if celula.cache_estado == "presente" and celula.tipo_cache:
            return celula.tipo_cache
        return "formula"
    return celula.tipo


def _coluna_parcial(matriz, elegiveis) -> bool:
    if not matriz or not elegiveis:
        return False
    for coluna in range(len(matriz[0])):
        valores = [_tipo_observado(matriz[indice][coluna]) != "ausente" for indice in elegiveis]
        if any(valores) and not all(valores):
            return True
    return False


def _linha_parcial(matriz, elegiveis) -> bool:
    for indice in elegiveis:
        presencas = [_tipo_observado(celula) != "ausente" for celula in matriz[indice]]
        if any(presencas) and not all(presencas):
            return True
    return False


def _resumo_formulas(matriz, linha_inicio, coluna_inicio, elegiveis: set[int]) -> dict:
    total = com_cache = sem_cache = 0
    observadas = []
    for indice, linha in enumerate(matriz):
        for coluna, celula in enumerate(linha):
            if celula.tipo != "formula":
                continue
            total += 1
            if celula.cache_estado == "presente":
                com_cache += 1
            else:
                sem_cache += 1
            item = {
                "linha": linha_inicio + indice,
                "coluna": coluna_inicio + coluna,
                "formula": celula.formula,
                "cache_estado": celula.cache_estado or "ausente",
                "registro": indice in elegiveis,
            }
            if celula.cache_estado == "presente":
                item["valor"] = celula.valor
                item["tipo_cache"] = celula.tipo_cache
            observadas.append(item)
    return {"total": total, "com_cache": com_cache, "sem_cache": sem_cache, "observadas": observadas}


def _publicar(preparados, alertas, ambiguidades) -> dict:
    """Catálogo raso. O contrato de persistência rejeita árvore mais funda que 8."""
    arquivos = []
    abas = []
    regioes = []
    for preparado in preparados:
        arquivos.append(
            {
                "indice": preparado["indice"],
                "nome": preparado["nome"],
                "extensao": preparado["extensao"],
                "mime": preparado["mime"],
                "bytes": preparado["bytes"],
                "sha256": preparado["sha256"],
                "contem_grafico": preparado["contem_grafico"],
                "alertas": preparado["alertas"],
            }
        )
        for aba in preparado["abas"]:
            abas.append(
                {
                    "arquivo": preparado["indice"],
                    "indice": aba["indice"],
                    "nome": aba["nome"],
                    "estado": aba["estado"],
                    "oculta": aba["oculta"],
                    "indicio_auxiliar": aba["indicio_auxiliar"],
                    "indicio_apresentacao": aba["indicio_apresentacao"],
                    "contem_desenho": aba["contem_desenho"],
                    "contem_grafico": aba["contem_grafico"],
                    "encoding": aba.get("encoding"),
                    "separador": aba.get("separador"),
                    "celulas_mescladas": aba["celulas_mescladas"],
                    "linhas": aba["linhas"],
                    "colunas": aba["colunas"],
                    "alertas": aba["alertas"],
                }
            )
            for regiao in aba["regioes"]:
                regioes.append(_publicar_regiao(preparado["indice"], aba["indice"], regiao))
    return {
        "contrato": VERSAO_PREPARACAO,
        "arquivos": arquivos,
        "abas": abas,
        "regioes": regioes,
        "alertas": _unicos(alertas),
        "ambiguidades": ambiguidades,
    }


def _publicar_regiao(arquivo: int, aba: int, regiao: dict) -> dict:
    cabecalho = regiao["cabecalho"]
    divisao = regiao["divisao_lateral"]
    return {
        "arquivo": arquivo,
        "aba": aba,
        "indice": regiao["indice"],
        "linha_inicio": regiao["linha_inicio"],
        "linha_fim": regiao["linha_fim"],
        "coluna_inicio": regiao["coluna_inicio"],
        "coluna_fim": regiao["coluna_fim"],
        "cabecalho_deteccao": cabecalho["deteccao"],
        "cabecalho_confianca": cabecalho["confianca"],
        "cabecalho_linha": cabecalho["linha"],
        "candidatos": cabecalho["candidatos"],
        "candidatos_total": cabecalho["candidatos_total"],
        "candidatos_amostra": cabecalho["candidatos_amostra"],
        "preambulo": regiao["preambulo"],
        "divisao_escolha": None if divisao is None else divisao["escolha"],
        "divisao_candidatos": [] if divisao is None else divisao["candidatos"],
        "tipos_fisicos": regiao["tipos_fisicos"],
        "linhas_ocupadas": regiao["linhas_ocupadas"],
        "linhas_elegiveis": regiao["linhas_elegiveis"],
        "linhas_vazias": regiao["linhas_vazias"],
        "linhas_invalidas": regiao["linhas_invalidas"],
        "linhas_pre_preenchidas": regiao["linhas_pre_preenchidas"],
        "linhas_cabecalho_repetido": regiao["linhas_cabecalho_repetido"],
        "formulas": regiao["formulas"],
        "formulas_com_cache": regiao["formulas_com_cache"],
        "formulas_sem_cache": regiao["formulas_sem_cache"],
        "formulas_observadas": regiao["formulas_observadas"],
        "celulas_preservadas": regiao["celulas_preservadas"],
        "alertas": regiao["alertas"],
    }


def _numerar_dados(indice_arquivo, abas, dados_tipados, ambiguidades) -> None:
    for aba in abas:
        for extra in aba.pop("_ambiguidades", []):
            item = {"arquivo": indice_arquivo, "aba": aba["indice"]}
            item.update(extra)
            ambiguidades.append(item)
        for regiao in aba["regioes"]:
            for registro in regiao.pop("_elegiveis"):
                dados_tipados.append(
                    {
                        "arquivo": indice_arquivo,
                        "aba": aba["indice"],
                        "regiao": regiao["indice"],
                        "linha": registro["linha"],
                        "valores": registro["valores"],
                    }
                )


def _grade_csv(texto: str, separador: str | None, cota: _Cota) -> dict:
    delimitador = separador or "\x1f"
    if separador is None and "\x1f" in texto:
        raise IngestaoErro("arquivo_indevido")
    celulas: dict[tuple[int, int], Celula] = {}
    max_linha = 0
    max_coluna = 0
    leitor = csv.reader(io.StringIO(texto), delimiter=delimitador)
    for numero, campos in enumerate(leitor, start=1):
        if not any(campo.strip() for campo in campos):
            cota.somar(1)
            continue
        max_linha = numero
        for coluna, campo in enumerate(campos, start=1):
            if campo.strip() == "":
                continue
            if len(campo) > MAX_TEXTO_CELULA:
                raise IngestaoErro("limite_processamento")
            cota.somar(1)
            celulas[(numero, coluna)] = Celula(tipo="texto", valor=campo, bruto=campo)
            max_coluna = max(max_coluna, coluna)
    return {"celulas": celulas, "linhas": max_linha, "colunas": max_coluna}


def _detectar_separador(texto: str) -> tuple[str | None, bool]:
    amostra = texto[:100_000]
    linhas = []
    for linha in amostra.splitlines():
        if linha.strip():
            linhas.append(linha)
        if len(linhas) >= 20:
            break
    if not linhas:
        return None, False
    placares = []
    for separador in (",", ";", "\t", "|"):
        contagens = []
        for linha in linhas:
            if separador not in linha:
                continue
            contagens.append(len(next(csv.reader([linha], delimiter=separador))))
        if not contagens:
            continue
        moda = max(set(contagens), key=contagens.count)
        if moda < 2:
            continue
        consistencia = contagens.count(moda) / len(contagens)
        if consistencia >= 0.8:
            placares.append((consistencia, moda, separador))
    if not placares:
        return None, False
    placares.sort(reverse=True)
    melhor = placares[0]
    empatados = [item for item in placares if item[0] == melhor[0] and item[1] == melhor[1]]
    if len(empatados) > 1:
        return None, True
    if len(placares) > 1 and placares[1][0] == melhor[0] and placares[1][1] != melhor[1]:
        return None, True
    return melhor[2], False


def _classificar_token(celula: Celula, separador: str | None) -> None:
    bruto = celula.valor if isinstance(celula.valor, str) else ""
    token = bruto.strip()
    iso = _data_iso(token)
    if iso is not None:
        celula.tipo = "data"
        celula.valor = iso
        return
    curta = _data_curta(token)
    if curta is not None:
        celula.tipo = "data_pendente"
        celula.bruto = bruto
        celula.ordem_data = curta[0]
        celula.valor = curta[1]
        return
    numero = _parse_numero(token, separador)
    if numero[0] == "numero":
        celula.tipo = "numero"
        celula.valor = numero[1]
        return
    if numero[0] == "ambiguo":
        celula.tipo = "numero_pendente"
        celula.bruto = bruto
        celula.valor = bruto
        return


def _resolver_datas(matriz, bloqueadas, alertas, ambiguidades, regiao) -> None:
    for coluna in range(len(matriz[0]) if matriz else []):
        pendentes = [
            matriz[indice][coluna]
            for indice, _linha in enumerate(matriz)
            if indice not in bloqueadas and matriz[indice][coluna].tipo == "data_pendente"
        ]
        if not pendentes:
            continue
        if any(item.ordem_data == "ambigua" for item in pendentes):
            alertas.append("data_ambigua")
            ambiguidades.append({"codigo": "data_ambigua", "regiao": regiao, "coluna": coluna + 1})
        for item in pendentes:
            if item.ordem_data == "ambigua" or item.valor is None:
                item.tipo = "texto"
                item.valor = item.bruto
            else:
                item.tipo = "data"


def _resolver_numeros_ambiguos(matriz, alertas, ambiguidades, regiao) -> None:
    for coluna in range(len(matriz[0]) if matriz else []):
        pendentes = [linha[coluna] for linha in matriz if linha[coluna].tipo == "numero_pendente"]
        if not pendentes:
            continue
        for celula in pendentes:
            if isinstance(celula.bruto, str):
                celula.valor = celula.bruto
            celula.tipo = "numero_pendente"
        alertas.append("numero_ambiguo")
        ambiguidades.append({"codigo": "numero_ambiguo", "regiao": regiao, "coluna": coluna + 1})


def _parse_numero(token: str, separador: str | None):
    if not token:
        return "nao", None
    if _excesso_numerico(token):
        if any(caractere not in "+-0123456789.," for caractere in token):
            return "nao", None
        raise IngestaoErro("limite_processamento")
    sinal = 1
    bruto = token
    if bruto[0] == "-":
        sinal = -1
        bruto = bruto[1:]
    elif bruto[0] == "+":
        bruto = bruto[1:]
    if not bruto or any(caractere not in "0123456789.," for caractere in bruto):
        return "nao", None
    if "," not in bruto and "." not in bruto:
        if len(bruto) > 1 and bruto[0] == "0":
            return "nao", None
        return "numero", _como_inteiro(bruto, sinal)
    if separador == ",":
        if "," in bruto:
            return "nao", None
        return _numero_com_separador(bruto, ".", sinal, milhar_ambiguo=True)
    if "," in bruto and "." in bruto:
        return _numero_misto(bruto, sinal)
    if "," in bruto:
        return _numero_com_separador(bruto, ",", sinal, milhar_ambiguo=True)
    return _numero_com_separador(bruto, ".", sinal, milhar_ambiguo=True)


def _numero_com_separador(bruto: str, separador: str, sinal: int, *, milhar_ambiguo: bool):
    partes = bruto.split(separador)
    if any(not parte.isdigit() for parte in partes):
        return "nao", None
    if len(partes) == 2:
        esquerda, direita = partes
        if not esquerda or not direita:
            return "nao", None
        if len(esquerda) > 1 and esquerda[0] == "0":
            return "nao", None
        if milhar_ambiguo and len(direita) == 3 and 1 <= len(esquerda) <= 3:
            return "ambiguo", None
        return "numero", _como_decimal(f"{esquerda}.{direita}", sinal)
    if any(len(parte) != 3 for parte in partes[1:]):
        return "nao", None
    if not partes[0] or len(partes[0]) > 3 or (len(partes[0]) > 1 and partes[0][0] == "0"):
        return "nao", None
    return "numero", _como_inteiro("".join(partes), sinal)


def _numero_misto(bruto: str, sinal: int):
    ultimo_ponto = bruto.rfind(".")
    ultima_virgula = bruto.rfind(",")
    decimal = "." if ultimo_ponto > ultima_virgula else ","
    milhar = "," if decimal == "." else "."
    esquerda, direita = bruto.rsplit(decimal, 1)
    if not direita.isdigit() or not esquerda:
        return "nao", None
    partes = esquerda.split(milhar)
    if any(not parte.isdigit() for parte in partes):
        return "nao", None
    if any(len(parte) != 3 for parte in partes[1:]):
        return "nao", None
    if len(partes[0]) > 3 or (len(partes[0]) > 1 and partes[0][0] == "0"):
        return "nao", None
    return "numero", _como_decimal(f"{''.join(partes)}.{direita}", sinal)


def _excesso_numerico(token: str) -> bool:
    return sum(caractere.isdigit() for caractere in token) > MAX_DIGITOS_NUMERO


def _como_inteiro(texto: str, sinal: int):
    try:
        return sinal * int(texto)
    except (ValueError, OverflowError) as exc:
        raise IngestaoErro("limite_processamento") from exc


def _como_decimal(texto: str, sinal: int):
    try:
        valor = float(texto)
    except (ValueError, OverflowError) as exc:
        raise IngestaoErro("limite_processamento") from exc
    if not math.isfinite(valor):
        raise IngestaoErro("limite_processamento")
    return sinal * valor


def _data_iso(token: str) -> str | None:
    if len(token) != 10 or not token[:4].isdigit() or not token[5:7].isdigit() or not token[8:].isdigit():
        return None
    if token[4] not in "-/" or token[7] != token[4]:
        return None
    try:
        return date(int(token[:4]), int(token[5:7]), int(token[8:])).isoformat()
    except ValueError:
        return None


def _data_curta(token: str):
    if len(token) != 10 or not token[:2].isdigit() or not token[3:5].isdigit() or not token[6:].isdigit():
        return None
    if token[2] not in "-/" or token[5] != token[2]:
        return None
    dia_ou_mes, mes_ou_dia, ano = int(token[:2]), int(token[3:5]), int(token[6:])
    if ano < 1900 or ano > 2100:
        return None
    dmy = _data_valida(ano, mes_ou_dia, dia_ou_mes)
    mdy = _data_valida(ano, dia_ou_mes, mes_ou_dia)
    if dmy and mdy:
        return ("ambigua", None)
    if dmy:
        return ("dmy", dmy)
    if mdy:
        return ("mdy", mdy)
    return None


def _data_valida(ano: int, mes: int, dia: int) -> str | None:
    try:
        return date(ano, mes, dia).isoformat()
    except ValueError:
        return None


def _pontuar(linha, largura, proxima, reinterpretar, separador) -> int:
    ocupadas = [celula for celula in linha if celula.tipo != "ausente"]
    if not ocupadas:
        return -10
    if len(ocupadas) == 1 and largura > 1:
        return -5
    if not all(_eh_rotulo(celula, reinterpretar, separador) for celula in ocupadas):
        return 0
    pontos = 3
    if len(ocupadas) == largura:
        pontos += 2
    rotulos = [_texto_celula(celula) for celula in ocupadas]
    if len(set(rotulos)) == len(rotulos):
        pontos += 2
    if proxima is not None and any(
        celula.tipo != "ausente" and not _eh_rotulo(celula, reinterpretar, separador) for celula in proxima
    ):
        pontos += 2
    return pontos


def _eh_rotulo(celula: Celula, reinterpretar: bool, separador: str | None) -> bool:
    if celula.tipo != "texto" or not isinstance(celula.valor, str):
        return False
    if not reinterpretar:
        return True
    token = celula.valor.strip()
    if _data_iso(token) is not None or _data_curta(token) is not None:
        return False
    return _parse_numero(token, separador)[0] != "numero"


def _rotulos(linha, reinterpretar, separador):
    saida = []
    for celula in linha:
        if _eh_rotulo(celula, reinterpretar, separador):
            saida.append(celula.valor)
        elif celula.tipo == "ausente":
            saida.append(None)
        else:
            saida.append(("outro", celula.tipo))
    return saida


def _rotulos_publicos(linha):
    saida = []
    for celula in linha:
        if celula.tipo == "texto" and isinstance(celula.valor, str):
            saida.append(celula.valor)
        else:
            saida.append(None)
    return saida


def _texto_celula(celula: Celula):
    return celula.valor


def _cabecalho_publico(deteccao: str, confianca: str, linha, candidatos: list[dict]) -> dict:
    """A decisão usa a lista inteira. O contrato guarda só uma amostra quando ela explode.

    Milhares de linhas com a mesma pontuação continuam na contagem e na
    classificação. A lista publicada não acompanha cada linha empatada.
    """
    total = len(candidatos)
    amostra = total > LIMITE_CANDIDATOS_PUBLICADOS
    publicados = _amostra_candidatos(candidatos) if amostra else candidatos
    return {
        "deteccao": deteccao,
        "confianca": confianca,
        "linha": linha,
        "candidatos": [_candidato_publico(item) for item in publicados],
        "candidatos_total": total,
        "candidatos_amostra": amostra,
    }


def _amostra_candidatos(candidatos: list[dict]) -> list[dict]:
    """Amostra estável: melhor, segundo, meio e fim da faixa empatada, depois os primeiros."""
    limite = LIMITE_CANDIDATOS_PUBLICADOS
    melhor = candidatos[0]["pontuacao"]
    faixa = [item for item in candidatos if melhor - item["pontuacao"] < MARGEM_CABECALHO]
    escolhidos = []
    vistos = set()

    def incluir(item) -> None:
        if item["indice"] in vistos or len(escolhidos) >= limite:
            return
        vistos.add(item["indice"])
        escolhidos.append(item)

    incluir(candidatos[0])
    if len(candidatos) > 1:
        incluir(candidatos[1])
    if faixa:
        incluir(faixa[len(faixa) // 2])
        incluir(faixa[-1])
    incluir(max(candidatos, key=lambda item: item["indice"]))
    for item in candidatos:
        incluir(item)
    escolhidos.sort(key=lambda item: (-item["pontuacao"], item["indice"]))
    return escolhidos


def _candidato_publico(candidato: dict) -> dict:
    return {
        "linha": candidato["linha"],
        "pontuacao": candidato["pontuacao"],
        "rotulos": candidato["rotulos"],
    }


def _grupos_laterais(colunas: list[int]) -> list[list[int]]:
    grupos = [[colunas[0]]]
    for coluna in colunas[1:]:
        if coluna - grupos[-1][-1] > 1:
            grupos.append([coluna])
        else:
            grupos[-1].append(coluna)
    return grupos


def _valor_json(celula: Celula) -> dict:
    if celula.tipo == "ausente":
        return {"tipo": "ausente"}
    if celula.tipo == "formula":
        corpo = {"tipo": "formula", "formula": celula.formula, "cache_estado": celula.cache_estado or "ausente"}
        if celula.cache_estado == "presente":
            corpo["valor"] = celula.valor
            corpo["tipo_cache"] = celula.tipo_cache
        return corpo
    if celula.tipo == "nao_suportado":
        return {"tipo": "nao_suportado"}
    return {"tipo": celula.tipo, "valor": celula.valor}


def _inspecionar_zip(conteudo: bytes, limites: LimitesIngestao) -> dict:
    if not conteudo.startswith(b"PK"):
        raise IngestaoErro("arquivo_corrompido")
    try:
        with zipfile.ZipFile(io.BytesIO(conteudo)) as arquivo:
            infos = arquivo.infolist()
            if not infos or len(infos) > MAX_ENTRADAS_ZIP:
                raise IngestaoErro("limite_descompressao")
            total = 0
            nomes = []
            for info in infos:
                nome = info.filename.replace("\\", "/")
                if nome.startswith("/") or ".." in nome.split("/"):
                    raise IngestaoErro("arquivo_indevido")
                if info.flag_bits & 0x1:
                    raise IngestaoErro("arquivo_indevido")
                if info.compress_size <= 0 or info.file_size < 0:
                    raise IngestaoErro("limite_descompressao")
                if info.file_size / info.compress_size > MAX_RAZAO_ZIP:
                    raise IngestaoErro("limite_descompressao")
                total += info.file_size
                if total > limites.max_descompactado or info.file_size > limites.max_descompactado:
                    raise IngestaoErro("limite_descompressao")
                nomes.append(nome)
            baixos = [nome.lower() for nome in nomes]
            if any(nome.endswith("vbaproject.bin") or "/embeddings/" in nome or "/activex/" in nome for nome in baixos):
                raise IngestaoErro("arquivo_indevido")
            if "[Content_Types].xml" not in nomes or "xl/workbook.xml" not in nomes:
                raise IngestaoErro("arquivo_corrompido")
            tipos = arquivo.read("[Content_Types].xml")
            if b"macroEnabled" in tipos:
                raise IngestaoErro("arquivo_indevido")
            desenhos = set()
            for nome in nomes:
                if not nome.startswith("xl/worksheets/sheet") or not nome.endswith(".xml"):
                    continue
                if b"<drawing" in arquivo.read(nome):
                    desenhos.add(nome)
            return {
                "contem_grafico": any(nome.startswith("xl/charts/") for nome in nomes),
                "desenhos": desenhos,
            }
    except IngestaoErro:
        raise
    except zipfile.BadZipFile as exc:
        raise IngestaoErro("arquivo_corrompido") from exc


def _estado_aba(estado: str) -> tuple[str, bool]:
    if estado == "hidden":
        return "oculta", True
    if estado == "veryHidden":
        return "muito_oculta", True
    return "visivel", False


def _texto_formula(valor) -> str | None:
    if isinstance(valor, str):
        texto = valor
    else:
        texto = getattr(valor, "text", None)
        if not isinstance(texto, str):
            return None
    texto = texto.strip()
    if not texto or texto == "=":
        return None
    if not texto.startswith("="):
        texto = "=" + texto
    if len(texto) > MAX_TEXTO_CELULA:
        raise IngestaoErro("limite_processamento")
    return texto


def _texto_rico(valor) -> str | None:
    if type(valor).__name__ != "CellRichText":
        return None
    partes = []
    for parte in valor:
        if isinstance(parte, str):
            partes.append(parte)
        else:
            trecho = getattr(parte, "text", None)
            if not isinstance(trecho, str):
                return None
            partes.append(trecho)
    texto = "".join(partes)
    if len(texto) > MAX_TEXTO_CELULA:
        raise IngestaoErro("limite_processamento")
    return texto


def _parece_binario(conteudo: bytes) -> bool:
    if conteudo.startswith((b"PK", b"%PDF", b"\xd0\xcf\x11\xe0")) or b"\x00" in conteudo:
        return True
    amostra = conteudo[:8192]
    controles = sum(1 for byte in amostra if byte < 9 or 13 < byte < 32)
    return controles > max(8, len(amostra) // 10)


def _decodificar(conteudo: bytes) -> tuple[str, str]:
    if conteudo.startswith(b"\xef\xbb\xbf"):
        try:
            return conteudo[3:].decode("utf-8"), "utf-8-bom"
        except UnicodeDecodeError as exc:
            raise IngestaoErro("codificacao_nao_suportada") from exc
    try:
        return conteudo.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        try:
            return conteudo.decode("cp1252"), "cp1252"
        except UnicodeDecodeError as exc:
            raise IngestaoErro("codificacao_nao_suportada") from exc


def _ler_requisicao(*, campos_form: frozenset[str] | None = None) -> tuple[list[ArquivoBruto], str | None, LimitesIngestao]:
    limites = resolver_limites()
    if request.content_length is not None and request.content_length > limites.max_bytes_total + _SOBRECARGA_MULTIPART:
        raise IngestaoErro("limite_bytes")
    if request.is_json:
        raise IngestaoErro("payload_invalido")
    permitidos = campos_form if campos_form is not None else frozenset({"chave"})
    if set(request.form.keys()) - permitidos:
        raise IngestaoErro("payload_invalido")
    if set(request.files.keys()) != {"arquivos"}:
        raise IngestaoErro("payload_invalido")
    storages = request.files.getlist("arquivos")
    if not storages or len(storages) > limites.max_arquivos:
        raise IngestaoErro("limite_arquivos" if len(storages) > limites.max_arquivos else "payload_invalido")
    chave = request.form.get("chave")
    if "chave" in request.form and not _chave_valida(chave or ""):
        raise IngestaoErro("payload_invalido")
    arquivos = []
    total = 0
    for storage in storages:
        nome = storage.filename or ""
        ext = _extensao(nome)
        teto = limites.max_bytes_csv if ext == ".csv" else limites.max_bytes_xlsx
        conteudo = _ler_limitado(storage.stream, teto)
        total += len(conteudo)
        if total > limites.max_bytes_total:
            raise IngestaoErro("limite_bytes")
        arquivos.append(ArquivoBruto(nome=nome, mime=storage.mimetype, conteudo=conteudo))
    return arquivos, chave, limites


def _ler_limitado(stream, teto: int) -> bytes:
    blocos = []
    restante = teto + 1
    while restante > 0:
        pedaco = stream.read(min(65_536, restante))
        if not pedaco:
            break
        blocos.append(pedaco)
        restante -= len(pedaco)
    conteudo = b"".join(blocos)
    if len(conteudo) > teto or stream.read(1):
        raise IngestaoErro("limite_bytes")
    return conteudo


def _metadados(limites: LimitesIngestao) -> dict:
    return {
        "versao_preparacao": VERSAO_PREPARACAO,
        "leitor": "agente_bi_ingestao",
        "envio_externo": False,
        "ia": False,
        "cabecalho_margem": MARGEM_CABECALHO,
        "cabecalho_minimo": MINIMO_CABECALHO,
        "politica_persistencia": "contrato_scrum_252",
        "limites": {
            "arquivos": limites.max_arquivos,
            "bytes_xlsx": limites.max_bytes_xlsx,
            "bytes_csv": limites.max_bytes_csv,
            "bytes_total": limites.max_bytes_total,
            "celulas": limites.max_celulas,
            "segundos": int(limites.max_segundos),
            "descompactado": limites.max_descompactado,
            "json_bytes": MAX_JSON_BYTES,
            "lista": MAX_LISTA,
            "texto": MAX_TEXTO,
            "digitos_numero": MAX_DIGITOS_NUMERO,
        },
        "lacunas": "celula_ausente_nao_preenchida",
    }


def _fora_do_contrato(valor) -> bool:
    try:
        bruto = _material_bytes(valor)
    except (TypeError, ValueError):
        return True
    if len(bruto) > MAX_JSON_BYTES:
        return True
    return _forma_recusada(valor, 0)


def _normalizar_zip(nome: str) -> str:
    return (nome or "").replace("\\", "/").lstrip("/")


def _dados_cabem_na_segmentacao(dados) -> bool:
    """A lista raiz pode passar de um documento. O miolo continua no contrato."""
    return not _forma_recusada(dados, 0, lista_aberta=isinstance(dados, list))


def _forma_recusada(valor, profundidade: int, *, lista_aberta: bool = False) -> bool:
    if profundidade > MAX_PROFUNDIDADE:
        return True
    if isinstance(valor, dict):
        if len(valor) > MAX_CHAVES:
            return True
        return any(
            not isinstance(chave, str) or len(chave) > 80 or _forma_recusada(item, profundidade + 1)
            for chave, item in valor.items()
        )
    if isinstance(valor, list):
        if not lista_aberta and len(valor) > MAX_LISTA:
            return True
        return any(_forma_recusada(item, profundidade + 1) for item in valor)
    if valor is None or isinstance(valor, bool):
        return False
    if isinstance(valor, int):
        return False
    if isinstance(valor, float):
        return not math.isfinite(valor)
    if isinstance(valor, str):
        return len(valor) > MAX_TEXTO
    return True


def _material_bytes(*valores) -> bytes:
    import json

    return json.dumps(valores if len(valores) != 1 else valores[0], ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8")


def _sha256_bytes(conteudo: bytes) -> str:
    return hashlib.sha256(conteudo).hexdigest()


def _chave_de(arquivos: list[ArquivoBruto]) -> str:
    digest = hashlib.sha256()
    for arquivo in arquivos:
        digest.update(hashlib.sha256(arquivo.conteudo).digest())
    return "ing-" + digest.hexdigest()[:16]


def _chave_valida(valor: str) -> bool:
    if not isinstance(valor, str) or not valor or len(valor) > 80:
        return False
    if not valor[0].isalnum():
        return False
    return all(caractere.isalnum() or caractere in "._-" for caractere in valor[1:])


def _teto(nome: str, padrao: int, maximo: int) -> int:
    bruto = os.getenv(nome)
    if bruto is None or not str(bruto).strip() or not str(bruto).strip().isdigit():
        return padrao
    valor = int(str(bruto).strip())
    if valor < 1:
        return padrao
    return min(valor, maximo)


def _extensao(nome: str) -> str:
    bruto = (nome or "").strip().lower()
    if "." not in bruto:
        return ""
    return "." + bruto.rsplit(".", 1)[-1]


def _unicos(itens) -> list[str]:
    saida = []
    for item in itens:
        if item not in saida:
            saida.append(item)
    return saida

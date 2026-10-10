"""Extração esparsa de XLSX via openpyxl. Fórmulas e cache ficam em campos distintos.

``Worksheet.iter_rows`` materializa o retângulo min/max. A malha ocupada fica em
``_cells``; o slice só percorre essa malha.
"""
from __future__ import annotations

import io
import zipfile
from datetime import date, datetime
from decimal import Decimal
from typing import Any
from urllib.parse import unquote
from xml.etree.ElementTree import ParseError

import openpyxl
from defusedxml import ElementTree as DefusedET
from defusedxml.common import EntitiesForbidden
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

from app.services.document_semantic_ir.structure import (
    BLOCK_TABLE,
    BLOCK_TEXT,
    KIND_SHEET,
    LIMITATION_CACHED,
    LIMITATION_CELL_CAP,
    LIMITATION_GAP_CAP,
    LIMITATION_MERGE_CAP,
    LIMITATION_RESOURCE,
    LIMITATION_TEXT_CAP,
    MAX_MERGE_AREA_ONE,
    MAX_MERGE_AREA_TOTAL,
    REL_MERGED_OVER,
    REL_SPANS_WIDTH,
    DocumentSemanticIrError,
    _add_rel,
    assign_block_ids,
    coverage,
    range_cell_area,
    relate_blocks,
    render_document_scalar,
    union_bbox,
)

MAX_STRUCTURAL_CELLS = 8000
MAX_MERGES = 2000
MAX_SERIALIZED_TEXT = 4000
MAX_INDEX_SPAN = 5000
MAX_SHEETS = 50
PREFLIGHT_MAX_MERGES = 2000
MAX_XML_UNCOMPRESSED = 4 * 1024 * 1024
MAX_ZIP_DECLARED_UNCOMPRESSED = 32 * 1024 * 1024
_WORKBOOK_PART = "xl/workbook.xml"
_WORKBOOK_RELS_PART = "xl/_rels/workbook.xml.rels"
_WORKSHEET_REL_TYPES = frozenset(
    {
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet",
        "http://purl.oclc.org/ooxml/officeDocument/relationships/worksheet",
    }
)
_WORKSHEET_ROOT_TAGS = frozenset(
    {
        "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}worksheet",
        "{http://purl.oclc.org/ooxml/spreadsheetml/main}worksheet",
    }
)


class _StructuralBudget:
    """Orçamento acumulado do documento inteiro, não de cada aba."""

    def __init__(self) -> None:
        self.cells = 0
        self.merges = 0
        self.text_chars = 0
        self.row_span = 0
        self.col_span = 0
        self.gap_items = 0
        self.stop_sheets = False
        self.text_closed = False


def extract_xlsx_sheets(data: bytes) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    wb_formula = None
    wb_cached = None
    try:
        limitations = _preflight_limitations(data)
        if limitations:
            return [], coverage(
                pages_or_sheets=0,
                tables_detected=0,
                visual_ununderstood_pages=[],
                limitations=limitations,
            )
        bio = io.BytesIO(data)
        wb_formula = openpyxl.load_workbook(bio, data_only=False, read_only=False)
        bio.seek(0)
        wb_cached = openpyxl.load_workbook(bio, data_only=True, read_only=False)
        if len(wb_formula.worksheets) != len(wb_cached.worksheets):
            raise DocumentSemanticIrError("xlsx_workbook_mismatch")
        limitations = []
        sheets: list[dict[str, Any]] = []
        budget = _StructuralBudget()
        for index, (sheet, cached_sheet) in enumerate(zip(wb_formula.worksheets, wb_cached.worksheets), start=1):
            if budget.stop_sheets:
                break
            built, missing_cached = _extract_sheet(
                sheet,
                cached_sheet,
                sheet_id=f"s{index}",
                limitations=limitations,
                budget=budget,
            )
            if built is None:
                break
            sheets.append(built)
            if missing_cached:
                limitations.append(LIMITATION_CACHED)
            if budget.text_closed:
                budget.stop_sheets = True
        tables_detected = sum(1 for sheet in sheets for block in sheet["blocks"] if block.get("type") == BLOCK_TABLE)
        return sheets, coverage(
            pages_or_sheets=len(sheets),
            tables_detected=tables_detected,
            visual_ununderstood_pages=[],
            limitations=limitations,
        )
    except DocumentSemanticIrError:
        raise
    except Exception as exc:
        raise DocumentSemanticIrError("xlsx_unreadable") from exc
    finally:
        if wb_formula is not None:
            wb_formula.close()
        if wb_cached is not None:
            wb_cached.close()


def _extract_sheet(
    sheet: Any,
    cached_sheet: Any,
    *,
    sheet_id: str,
    limitations: list[str],
    budget: _StructuralBudget,
) -> tuple[dict[str, Any] | None, bool]:
    sheet_name = str(sheet.title)
    if not _reserve_document_text(sheet_name, limitations, budget):
        return None, False
    hidden_rows = _hidden_indexes(sheet.row_dimensions)
    hidden_columns, hidden_spans = _hidden_column_ranges(sheet.column_dimensions)
    hidden_row_set = set(hidden_rows)
    occupied: dict[tuple[int, int], dict[str, Any]] = {}
    missing_cached = False
    for cell in _populated_cells(sheet):
        if budget.text_closed or budget.stop_sheets:
            break
        if not _is_occupied(cell):
            continue
        if budget.cells >= MAX_STRUCTURAL_CELLS:
            _append_limitation(limitations, LIMITATION_CELL_CAP)
            budget.stop_sheets = True
            break
        budget.cells += 1
        raw_formula = _formula(cell)
        formula = _bounded_text(raw_formula, limitations, budget) if raw_formula else None
        cached_value = None
        include_cached = False
        if raw_formula:
            cached_value = _cached_value(cached_sheet, int(cell.row), int(cell.column))
            if isinstance(cached_value, str) and cached_value.startswith("="):
                cached_value = None
            if cached_value is None:
                missing_cached = True
                include_cached = True
            else:
                cached_text = _document_text(cached_value)
                if _reserve_document_text(cached_text, limitations, budget):
                    include_cached = True
                else:
                    cached_value = None
        stored_value = None if raw_formula else _json_value(cell.value)
        if not raw_formula:
            value_text = _document_text(stored_value)
            if not _reserve_document_text(value_text, limitations, budget):
                stored_value = None
        payload: dict[str, Any] = {
            "row": int(cell.row),
            "column": int(cell.column),
            "coordinate": str(cell.coordinate),
            "value": stored_value,
            "value_type": "formula" if raw_formula else _value_type(cell),
            "bbox": {
                "x0": int(cell.column),
                "y0": int(cell.row),
                "x1": int(cell.column),
                "y1": int(cell.row),
            },
            "relationships": [],
        }
        if raw_formula:
            if formula:
                payload["formula"] = formula
            if include_cached:
                payload["cached_value"] = cached_value
        comment = _bounded_text(_comment_text(cell), limitations, budget)
        if comment:
            payload["comment"] = comment
        style_ref = _bounded_style(_style_ref(cell), limitations, budget)
        if style_ref:
            payload["style_ref"] = style_ref
        if cell.row in hidden_row_set or _column_in_spans(int(cell.column), hidden_spans):
            payload["hidden"] = True
        occupied[(int(cell.row), int(cell.column))] = payload
        if budget.text_closed:
            break

    blocks = _blocks_from_occupied(occupied)
    assign_block_ids(sheet_id, blocks)
    relate_blocks(blocks, edge_mode="inclusive")
    merges = _merges(sheet, limitations, budget)
    _attach_merges(blocks, merges)
    gaps, boundaries = _extent(list(occupied), limitations, budget)
    return (
        {
            "kind": KIND_SHEET,
            "sheet_id": sheet_id,
            "sheet_name": sheet_name,
            "visibility": str(sheet.sheet_state or "visible"),
            "coordinate_space": "cell_index_inclusive",
            "blocks": blocks,
            "merges": merges,
            "hidden_rows": hidden_rows,
            "hidden_columns": hidden_columns,
            "gaps": gaps,
            "boundaries": boundaries,
        },
        missing_cached,
    )


def _blocks_from_occupied(occupied: dict[tuple[int, int], dict[str, Any]]) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    for cluster in _clusters(list(occupied)):
        cells = [occupied[point] for point in cluster]
        rows = {cell["row"] for cell in cells}
        cols = {cell["column"] for cell in cells}
        block_type = BLOCK_TABLE if len(rows) >= 2 and len(cols) >= 2 else BLOCK_TEXT
        blocks.append(
            {
                "type": block_type,
                "bbox": union_bbox([cell["bbox"] for cell in cells]),
                "cells": cells,
                "provenance": {"extractor": "openpyxl", "method": "sparse_cells"},
                "relationships": [],
            }
        )
    return blocks


def _clusters(points: list[tuple[int, int]]) -> list[list[tuple[int, int]]]:
    pending = set(points)
    clusters: list[list[tuple[int, int]]] = []
    while pending:
        start = min(pending)
        pending.remove(start)
        stack = [start]
        group: list[tuple[int, int]] = []
        while stack:
            row, col = stack.pop()
            group.append((row, col))
            for neighbor in ((row + 1, col), (row - 1, col), (row, col + 1), (row, col - 1)):
                if neighbor in pending:
                    pending.remove(neighbor)
                    stack.append(neighbor)
        clusters.append(group)
    return clusters


def _extent(
    points: list[tuple[int, int]],
    limitations: list[str],
    budget: _StructuralBudget,
) -> tuple[list[dict[str, Any]], dict[str, int] | None]:
    if not points:
        return [], None
    rows = [row for row, _ in points]
    cols = [col for _, col in points]
    min_row, max_row = min(rows), max(rows)
    min_col, max_col = min(cols), max(cols)
    row_set = set(rows)
    col_set = set(cols)
    boundaries = {
        "min_row": min_row,
        "max_row": max_row,
        "min_column": min_col,
        "max_column": max_col,
    }
    row_span = max_row - min_row
    col_span = max_col - min_col
    row_over = budget.row_span + row_span > MAX_INDEX_SPAN
    col_over = budget.col_span + col_span > MAX_INDEX_SPAN
    if row_over or col_over:
        _append_limitation(limitations, LIMITATION_GAP_CAP)
        budget.stop_sheets = True
    budget.row_span += row_span
    budget.col_span += col_span
    gaps = _axis_gaps("row", min_row, max_row, row_set, row_over, limitations, budget)
    gaps.extend(_axis_gaps("column", min_col, max_col, col_set, col_over, limitations, budget))
    return gaps, boundaries


def _axis_gaps(
    axis: str,
    start: int,
    end: int,
    occupied: set[int],
    over: bool,
    limitations: list[str],
    budget: _StructuralBudget,
) -> list[dict[str, Any]]:
    if over:
        return [{"axis": axis, "start": start, "end": end, "collapsed": True}]
    items = [{"axis": axis, "index": index} for index in range(start, end + 1) if index not in occupied]
    if budget.gap_items + len(items) > MAX_INDEX_SPAN:
        _append_limitation(limitations, LIMITATION_GAP_CAP)
        budget.stop_sheets = True
        return [{"axis": axis, "start": start, "end": end, "collapsed": True}]
    budget.gap_items += len(items)
    return items


def _merges(sheet: Any, limitations: list[str], budget: _StructuralBudget) -> list[dict[str, str]]:
    ranges = sorted(sheet.merged_cells.ranges, key=lambda merged: str(merged))
    merges: list[dict[str, str]] = []
    for merged in ranges:
        if budget.merges >= MAX_MERGES:
            _append_limitation(limitations, LIMITATION_MERGE_CAP)
            budget.stop_sheets = True
            break
        budget.merges += 1
        anchor_col, anchor_row, _, _ = range_boundaries(str(merged))
        anchor = f"{get_column_letter(anchor_col)}{anchor_row}"
        merges.append({"range": str(merged), "anchor": anchor})
    merges.sort(key=lambda item: (item["anchor"], item["range"]))
    return merges


def _attach_merges(blocks: list[dict[str, Any]], merges: list[dict[str, str]]) -> None:
    by_coordinate: dict[str, tuple[dict[str, Any], dict[str, Any]]] = {}
    for block in blocks:
        for cell in block.get("cells") or []:
            coordinate = cell.get("coordinate")
            if isinstance(coordinate, str):
                by_coordinate[coordinate] = (block, cell)
    for merge in merges:
        found = by_coordinate.get(merge["anchor"])
        if found is None:
            continue
        block, cell = found
        _add_rel(block, REL_MERGED_OVER, cell["id"], range=merge["range"], anchor=merge["anchor"])
        _add_rel(cell, REL_MERGED_OVER, cell["id"], range=merge["range"], anchor=merge["anchor"])
        min_col, _, max_col, _ = range_boundaries(merge["range"])
        width = int(max_col) - int(min_col) + 1
        if width > 1:
            _add_rel(block, REL_SPANS_WIDTH, cell["id"], columns=width, range=merge["range"])


def _is_occupied(cell: Any) -> bool:
    if isinstance(cell, MergedCell):
        return False
    if _comment_text(cell):
        return True
    value = cell.value
    if isinstance(value, str) and value.startswith("="):
        return True
    return value is not None


def _formula(cell: Any) -> str | None:
    value = cell.value
    if isinstance(value, str) and value.startswith("="):
        return value
    return None


def _comment_text(cell: Any) -> str | None:
    comment = getattr(cell, "comment", None)
    if comment is None:
        return None
    text = getattr(comment, "text", None)
    if text is None:
        text = getattr(comment, "content", None)
    if text is None:
        return None
    cleaned = str(text).strip()
    return cleaned or None


def _style_ref(cell: Any) -> dict[str, Any] | None:
    ref: dict[str, Any] = {}
    number_format = getattr(cell, "number_format", None)
    if number_format and number_format != "General":
        ref["number_format"] = str(number_format)
    font = getattr(cell, "font", None)
    if font is not None and getattr(font, "bold", False):
        ref["bold"] = True
    if font is not None and getattr(font, "italic", False):
        ref["italic"] = True
    fill = getattr(cell, "fill", None)
    pattern = getattr(fill, "patternType", None) if fill is not None else None
    if pattern and pattern != "none":
        ref["fill_pattern"] = str(pattern)
    return ref or None


def _value_type(cell: Any) -> str:
    if getattr(cell, "data_type", None) == "e":
        return "error"
    value = cell.value
    if value is None:
        return "empty"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float, Decimal)):
        return "number"
    if isinstance(value, (datetime, date)):
        return "datetime"
    if isinstance(value, str):
        return "string"
    return "unknown"


def _json_value(value: Any) -> Any:
    if isinstance(value, bool) or value is None or isinstance(value, (int, float, str)):
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _hidden_indexes(dimensions: Any) -> list[int]:
    found: list[int] = []
    for index, dim in dimensions.items():
        if getattr(dim, "hidden", False):
            found.append(int(index))
    return sorted(set(found))


def _populated_cells(sheet: Any) -> list[Any]:
    cells = getattr(sheet, "_cells", None)
    if not isinstance(cells, dict):
        raise DocumentSemanticIrError("xlsx_cell_index_unavailable")
    return [cells[key] for key in cells]


def _cached_value(cached_sheet: Any, row: int, column: int) -> Any:
    cells = getattr(cached_sheet, "_cells", None)
    if not isinstance(cells, dict):
        raise DocumentSemanticIrError("xlsx_cell_index_unavailable")
    found = cells.get((row, column))
    if found is None:
        return None
    return _json_value(found.value)


def _document_text(value: Any) -> str | None:
    """Texto documental que pode chegar ao payload. Não inclui id, coordenada nem enum."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, str):
        return value
    rendered = render_document_scalar(value)
    if rendered is not None:
        return rendered
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return None


def _reserve_document_text(text: str | None, limitations: list[str], budget: _StructuralBudget) -> bool:
    """Reserva o texto inteiro no orçamento global. Não trunca."""
    if text is None:
        return not budget.text_closed
    if budget.text_closed or budget.text_chars + len(text) > MAX_SERIALIZED_TEXT:
        _append_limitation(limitations, LIMITATION_TEXT_CAP)
        budget.text_closed = True
        budget.stop_sheets = True
        return False
    budget.text_chars += len(text)
    return True


def _bounded_text(text: str | None, limitations: list[str], budget: _StructuralBudget) -> str | None:
    if text is None:
        return None
    if not _reserve_document_text(text, limitations, budget):
        return None
    return text


def _bounded_style(
    style_ref: dict[str, Any] | None,
    limitations: list[str],
    budget: _StructuralBudget,
) -> dict[str, Any] | None:
    if not style_ref:
        return None
    for key in ("number_format", "fill_pattern"):
        raw = style_ref.get(key)
        if not isinstance(raw, str):
            continue
        if not _reserve_document_text(raw, limitations, budget):
            style_ref.pop(key, None)
    return style_ref or None


def _preflight_limitations(data: bytes) -> list[str]:
    """Gate leve no ZIP/XML. Não chama load_workbook e não descomprime além do orçamento.

    As worksheets vêm das entradas ``<sheet>`` do workbook. Cada entrada conta no
    orçamento, mesmo quando várias apontam para o mesmo part.
    """
    budget = _logical_sheet_budget(data)
    if budget["limitation"]:
        return [str(budget["limitation"])]
    limitations: list[str] = []
    if budget["sheets"] > MAX_SHEETS or budget["xml_bytes"] > MAX_XML_UNCOMPRESSED:
        _append_limitation(limitations, LIMITATION_RESOURCE)
        return limitations
    if budget["cells"] > MAX_STRUCTURAL_CELLS:
        _append_limitation(limitations, LIMITATION_CELL_CAP)
        return limitations
    if budget["merge_area"] > MAX_MERGE_AREA_TOTAL or budget["merges"] > PREFLIGHT_MAX_MERGES:
        _append_limitation(limitations, LIMITATION_MERGE_CAP)
        return limitations
    return limitations


def _logical_sheet_budget(data: bytes) -> dict[str, Any]:
    """Custo lógico por referência de sheet. O parse do part é cacheado; o custo não.

    Não chama load_workbook. Estrutura inconsistente levanta ``xlsx_unreadable``.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = [info for info in archive.infolist() if not info.is_dir()]
            declared = sum(max(info.file_size, 0) for info in infos)
            if declared > MAX_ZIP_DECLARED_UNCOMPRESSED:
                raise DocumentSemanticIrError("xlsx_unreadable")
            by_name: dict[str, zipfile.ZipInfo] = {}
            for info in infos:
                by_name.setdefault(info.filename, info)
            sheets, early = _referenced_worksheets(archive, by_name)
            if early:
                return _sheet_budget(limitation=early[0])
            workbook_info = by_name[_WORKBOOK_PART]
            xml_bytes = workbook_info.file_size + sum(info.file_size for info in sheets)
            if any(info.file_size > MAX_XML_UNCOMPRESSED for info in sheets) or xml_bytes > MAX_XML_UNCOMPRESSED:
                return _sheet_budget(sheets=len(sheets), xml_bytes=xml_bytes, limitation=LIMITATION_RESOURCE)
            part_analysis_cache: dict[str, dict[str, Any]] = {}
            cells = 0
            merges = 0
            merge_area = 0
            for info in sheets:
                analysis = part_analysis_cache.get(info.filename)
                if analysis is None:
                    analysis = _analyze_worksheet_part(archive, info)
                    if analysis["limitation"]:
                        return _sheet_budget(
                            sheets=len(sheets),
                            cells=cells,
                            merges=merges,
                            merge_area=merge_area,
                            xml_bytes=xml_bytes,
                            limitation=str(analysis["limitation"]),
                        )
                    part_analysis_cache[info.filename] = analysis
                cells += int(analysis["cells"])
                merges += int(analysis["merges"])
                merge_area += int(analysis["merge_area"])
            return _sheet_budget(
                sheets=len(sheets),
                cells=cells,
                merges=merges,
                merge_area=merge_area,
                xml_bytes=xml_bytes,
            )
    except DocumentSemanticIrError:
        raise
    except (zipfile.BadZipFile, EntitiesForbidden, ParseError, OSError, ValueError) as exc:
        raise DocumentSemanticIrError("xlsx_unreadable") from exc


def _sheet_budget(
    *,
    sheets: int = 0,
    cells: int = 0,
    merges: int = 0,
    merge_area: int = 0,
    xml_bytes: int = 0,
    limitation: str | None = None,
) -> dict[str, Any]:
    return {
        "sheets": sheets,
        "cells": cells,
        "merges": merges,
        "merge_area": merge_area,
        "xml_bytes": xml_bytes,
        "limitation": limitation,
    }


def _analyze_worksheet_part(archive: zipfile.ZipFile, info: zipfile.ZipInfo) -> dict[str, Any]:
    """Mede um part uma vez. Estouro do part sozinho vira limitação e não entra no cache."""
    cells = 0
    merges = 0
    merge_area = 0
    saw_root = False
    with archive.open(info.filename) as raw:
        for event, elem in DefusedET.iterparse(raw, events=("start", "end")):
            if event == "start":
                if not saw_root:
                    saw_root = True
                    if elem.tag not in _WORKSHEET_ROOT_TAGS:
                        raise DocumentSemanticIrError("xlsx_unreadable")
                continue
            name = _local_name(elem.tag)
            if name == "c":
                cells += 1
                if cells > MAX_STRUCTURAL_CELLS:
                    elem.clear()
                    return _part_analysis(cells, merges, merge_area, info.file_size, LIMITATION_CELL_CAP)
            elif name == "mergeCell":
                area = range_cell_area(str(elem.attrib.get("ref") or ""))
                if (
                    area is None
                    or area > MAX_MERGE_AREA_ONE
                    or merge_area + area > MAX_MERGE_AREA_TOTAL
                ):
                    elem.clear()
                    return _part_analysis(cells, merges, merge_area, info.file_size, LIMITATION_MERGE_CAP)
                merge_area += area
                merges += 1
                if merges > PREFLIGHT_MAX_MERGES:
                    elem.clear()
                    return _part_analysis(cells, merges, merge_area, info.file_size, LIMITATION_MERGE_CAP)
            elem.clear()
    if not saw_root:
        raise DocumentSemanticIrError("xlsx_unreadable")
    return _part_analysis(cells, merges, merge_area, info.file_size, None)


def _part_analysis(
    cells: int,
    merges: int,
    merge_area: int,
    xml_size: int,
    limitation: str | None,
) -> dict[str, Any]:
    return {
        "cells": cells,
        "merges": merges,
        "merge_area": merge_area,
        "xml_size": xml_size,
        "limitation": limitation,
    }


def _referenced_worksheets(
    archive: zipfile.ZipFile,
    by_name: dict[str, zipfile.ZipInfo],
) -> tuple[list[zipfile.ZipInfo], list[str]]:
    """Resolve cada ``<sheet>`` do workbook. Não deduplica o part de destino.

    Relationship que não seja worksheet interno — chartsheet, styles, drawing,
    theme, externalLink ou qualquer outro tipo — fecha como ``xlsx_unreadable``.
    Relationship órfão, sem ``<sheet>``, fica fora da IR.
    """
    workbook_info = by_name.get(_WORKBOOK_PART)
    rels_info = by_name.get(_WORKBOOK_RELS_PART)
    if workbook_info is None or rels_info is None:
        raise DocumentSemanticIrError("xlsx_unreadable")
    if workbook_info.file_size > MAX_XML_UNCOMPRESSED or rels_info.file_size > MAX_XML_UNCOMPRESSED:
        return [], [LIMITATION_RESOURCE]
    workbook = _parse_xml_part(archive, _WORKBOOK_PART)
    rels = _parse_xml_part(archive, _WORKBOOK_RELS_PART)
    relationships: dict[str, Any] = {}
    for elem in rels.iter():
        if _local_name(elem.tag) != "Relationship":
            continue
        rel_id = elem.attrib.get("Id")
        if not rel_id or rel_id in relationships:
            raise DocumentSemanticIrError("xlsx_unreadable")
        relationships[rel_id] = elem
    sheets: list[zipfile.ZipInfo] = []
    for elem in workbook.iter():
        if _local_name(elem.tag) != "sheet":
            continue
        rel_id = _namespaced_attr(elem, "id")
        if not rel_id or not str(rel_id).strip():
            raise DocumentSemanticIrError("xlsx_unreadable")
        rel = relationships.get(rel_id)
        if rel is None:
            raise DocumentSemanticIrError("xlsx_unreadable")
        mode = str(rel.attrib.get("TargetMode") or "Internal").strip().lower()
        if mode != "internal":
            raise DocumentSemanticIrError("xlsx_unreadable")
        rel_type = str(rel.attrib.get("Type") or "").strip()
        if rel_type not in _WORKSHEET_REL_TYPES:
            raise DocumentSemanticIrError("xlsx_unreadable")
        target = _resolve_internal_target(str(rel.attrib.get("Target") or ""))
        if target is None:
            raise DocumentSemanticIrError("xlsx_unreadable")
        info = by_name.get(target)
        if info is None:
            raise DocumentSemanticIrError("xlsx_unreadable")
        sheets.append(info)
    return sheets, []


def _parse_xml_part(archive: zipfile.ZipFile, name: str) -> Any:
    with archive.open(name) as raw:
        return DefusedET.parse(raw).getroot()


def _local_name(tag: str) -> str:
    return str(tag).rsplit("}", 1)[-1]


def _namespaced_attr(elem: Any, local: str) -> str | None:
    for key, value in elem.attrib.items():
        if _local_name(key) == local:
            return str(value)
    return None


def _resolve_internal_target(target: str) -> str | None:
    """Normaliza um target OOXML interno. Rejeita externo, travessia e path fora de xl/."""
    if not isinstance(target, str) or not target or target != target.strip():
        return None
    if any(char in target for char in ("\\", "\x00", "\r", "\n", "?", "#")):
        return None
    if target.startswith("//"):
        return None
    scheme = target.split("/", 1)[0]
    if ":" in scheme:
        return None
    decoded = unquote(target)
    if any(char in decoded for char in ("\\", "\x00", "?", "#")):
        return None
    absolute = decoded.startswith("/")
    body = decoded[1:] if absolute else decoded
    segments = body.split("/")
    if any(segment in {"", ".", ".."} for segment in segments):
        return None
    parts = segments if absolute else ["xl", *segments]
    if not parts or parts[0] != "xl" or any(part in {".", ".."} for part in parts):
        return None
    resolved = "/".join(parts)
    if not resolved.endswith(".xml") or resolved.startswith("/"):
        return None
    return resolved


def _append_limitation(limitations: list[str], code: str) -> None:
    if code not in limitations:
        limitations.append(code)


def _hidden_column_ranges(dimensions: Any) -> tuple[list[str], list[tuple[int, int]]]:
    found: list[tuple[int, int, str]] = []
    for letter, dim in dimensions.items():
        if not getattr(dim, "hidden", False):
            continue
        start = getattr(dim, "min", None)
        end = getattr(dim, "max", None)
        if not start or not end:
            try:
                start = end = column_index_from_string(str(getattr(dim, "index", None) or letter))
            except ValueError:
                continue
        start_i, end_i = int(start), int(end)
        if end_i < start_i:
            start_i, end_i = end_i, start_i
        start_label = get_column_letter(start_i)
        end_label = get_column_letter(end_i)
        label = start_label if start_i == end_i else f"{start_label}:{end_label}"
        found.append((start_i, end_i, label))
    found.sort(key=lambda item: (item[0], item[1], item[2]))
    return [item[2] for item in found], [(item[0], item[1]) for item in found]


def _column_in_spans(column: int, spans: list[tuple[int, int]]) -> bool:
    return any(start <= column <= end for start, end in spans)

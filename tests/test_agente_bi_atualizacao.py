"""SCRUM-264: atualização compatível, recálculo determinístico e bloqueio fechado."""
from __future__ import annotations

import hashlib
import io
import json
import logging
from datetime import date
from pathlib import Path

import pytest

from app.extensions import db
from app.models import (
    AgenteBiAnalise,
    AgenteBiDashboard,
    AgenteBiDashboardDataset,
    AgenteBiDataset,
    AgenteBiDatasetBloco,
    IaChamadaTentativa,
    IaConsumoEvento,
)
from app.services.agente_bi_atualizacao_service import AtualizacaoErro, atualizar_dashboard
from app.services.agente_bi_dashboard_service import consultar_dashboard, montar_dashboard
from app.services.agente_bi_definir_service import definir_selecao
from app.services.agente_bi_ingestao_service import ArquivoBruto, _fora_do_contrato, ingerir
from app.services.agente_bi_perfil_semantico_service import perfilar_dataset
from app.services.agente_bi_persistencia_service import AgenteBiErro, escopo_do_usuario, obter_analise, obter_dataset
from tests.conftest import seed_usuario
from tests.test_agente_bi_analisar import SENTINELA, _movimento
from tests.test_agente_bi_calculo import _calcular, _jornada_de
from tests.test_agente_bi_dashboard import _painel_soma, _pronto
from tests.test_agente_bi_ingestao import _cliente as _cliente_ingestao
from tests.test_agente_bi_ingestao import _entrar, _par, _xlsx

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(autouse=True)
def _sem_anthropic(monkeypatch):
    def proibido(*_args, **_kwargs):
        raise AssertionError("anthropic")

    monkeypatch.setattr(
        "app.services.agente_bi_anthropic_service.chamar_anthropic_governado",
        proibido,
    )


def _arquivos(conteudo: bytes, nome="movimento.xlsx"):
    return [ArquivoBruto(nome=nome, mime=None, conteudo=conteudo)]


def _painel(email, slug):
    usuario, franquia, escopo, definicao, calculo = _pronto(email, slug)
    painel = montar_dashboard(escopo, {"analise_id": calculo["calculo"]["id"]})
    return usuario, franquia, escopo, definicao, calculo, painel


def _soma(consulta) -> str:
    return _painel_soma(consulta["graficos"])["resultado"]


def _planilha(linhas_movimento, linhas_faixas=None):
    return _xlsx(
        {
            "movimento": linhas_movimento,
            "faixas": linhas_faixas
            or [
                ["alfa", "beta"],
                [SENTINELA, 1],
                ["K200", 2],
                ["K300", 3],
            ],
        }
    )


def _compativel():
    return _planilha(
        [
            ["alfa", "beta", "quando", "qtd", "valor"],
            [SENTINELA, "norte", date(2024, 3, 1), 10, 10.5],
            ["K200", "sul", date(2024, 3, 2), 4, 4.25],
            [SENTINELA, "norte", date(2024, 3, 3), 10, 10.5],
            ["K300", "leste", date(2024, 3, 4), 4, 4.25],
        ]
    )


def _nomes_parecidos():
    return _planilha(
        [
            ["alfa", "beta", "quando", "quantidade", "valor_total"],
            [SENTINELA, "norte", date(2024, 3, 1), 10, 10.5],
            ["K200", "sul", date(2024, 3, 2), 4, 4.25],
            [SENTINELA, "norte", date(2024, 3, 3), 10, 10.5],
        ]
    )


def _colunas_trocadas():
    return _planilha(
        [
            ["alfa", "beta", "quando", "valor", "qtd"],
            [SENTINELA, "norte", date(2024, 3, 1), 10.5, 10],
            ["K200", "sul", date(2024, 3, 2), 4.25, 4],
            [SENTINELA, "norte", date(2024, 3, 3), 10.5, 10],
        ]
    )


def _tipo_incompativel():
    return _planilha(
        [
            ["alfa", "beta", "quando", "qtd", "valor"],
            [SENTINELA, "norte", date(2024, 3, 1), "dez", 10.5],
            ["K200", "sul", date(2024, 3, 2), "quatro", 4.25],
            [SENTINELA, "norte", date(2024, 3, 3), "dez", 10.5],
        ]
    )


def _atualizar(escopo, painel_id, conteudo, modo):
    return atualizar_dashboard(escopo, painel_id, modo=modo, arquivos=_arquivos(conteudo))


def _movimento_distinto():
    """Mesmas linhas do movimento, com bytes diferentes. Não é o mesmo lote."""
    from openpyxl import Workbook

    livro = Workbook()
    livro.properties.creator = "lote-distinto"
    livro.properties.title = "lote-distinto"
    abas = (
        (
            "movimento",
            [
                ["alfa", "beta", "quando", "qtd", "valor"],
                [SENTINELA, "norte", date(2024, 3, 1), 10, 10.5],
                ["K200", "sul", date(2024, 3, 2), 4, 4.25],
                [SENTINELA, "norte", date(2024, 3, 3), 10, 10.5],
            ],
        ),
        (
            "faixas",
            [
                ["alfa", "beta"],
                [SENTINELA, 1],
                ["K200", 2],
                ["K300", 3],
            ],
        ),
    )
    primeira = True
    for nome, linhas in abas:
        aba = livro.active if primeira else livro.create_sheet(nome)
        primeira = False
        aba.title = nome
        for i, linha in enumerate(linhas, start=1):
            for j, valor in enumerate(linha, start=1):
                aba.cell(i, j, valor)
    destino = io.BytesIO()
    livro.save(destino)
    return destino.getvalue()


def _sha(conteudo: bytes) -> str:
    return hashlib.sha256(conteudo).hexdigest()


def _artefatos(escopo):
    filtro = {
        "conta_id": escopo.conta_id,
        "franquia_id": escopo.franquia_id,
        "usuario_id": escopo.usuario_id,
    }
    return (
        AgenteBiDataset.query.filter_by(**filtro).count(),
        AgenteBiAnalise.query.filter_by(**filtro).count(),
        AgenteBiDashboard.query.filter_by(**filtro).count(),
        AgenteBiDashboardDataset.query.filter_by(**filtro).count(),
        AgenteBiDatasetBloco.query.filter_by(**filtro).count(),
    )


def _contagens(escopo):
    filtro = {
        "conta_id": escopo.conta_id,
        "franquia_id": escopo.franquia_id,
        "usuario_id": escopo.usuario_id,
    }
    return (
        AgenteBiDataset.query.filter_by(**filtro).count(),
        AgenteBiAnalise.query.filter_by(**filtro).count(),
        AgenteBiDashboard.query.filter_by(**filtro).count(),
    )


def test_substituicao_compativel_recalcula_preserva_definicao_e_historico(ctx):
    _usuario, _franquia, escopo, definicao, calculo, painel = _painel(
        "atualizacao.substituir@test.com",
        "atualizacao-substituir",
    )
    antes_tentativas = IaChamadaTentativa.query.count()
    antes_consumo = IaConsumoEvento.query.count()
    antigo_id = painel["dashboard"]["id"]
    dataset_id = calculo["calculo"]["resultado"]["dataset_id"]
    definicao_id = definicao["definicao"]["id"]
    definicoes = obter_analise(escopo, definicao_id)["resultado"]["definicoes"]
    impressao_antiga = obter_dataset(escopo, dataset_id)["impressao"]

    resultado = _atualizar(escopo, antigo_id, _compativel(), "substituir")

    assert resultado["chamada_anthropic"] is False
    assert resultado["ia_calculo"] is False
    assert resultado["recalculado"] is True
    assert resultado["definicao_preservada"] is True
    assert resultado["definicao_origem_id"] == definicao_id
    assert resultado["definicao_id"] != definicao_id
    assert resultado["dataset_anterior_id"] == dataset_id
    assert resultado["dataset_id"] != dataset_id
    assert resultado["versao"] == 2
    assert resultado["impressao"] != impressao_antiga
    assert _soma(resultado["consulta"]) == "29.5"
    assert _painel_soma(resultado["consulta"]["graficos"])["elegiveis"] == 4
    assert resultado["consulta"]["chamada_anthropic"] is False
    assert resultado["consulta"]["versao"] == 2

    nova_definicao = obter_analise(escopo, resultado["definicao_id"])
    antiga_definicao = obter_analise(escopo, definicao_id)
    assert nova_definicao["resultado"]["definicoes"] == definicoes
    assert antiga_definicao["resultado"]["definicoes"] == definicoes
    assert nova_definicao["resultado"]["dataset_versao"] == 2
    assert antiga_definicao["resultado"]["dataset_versao"] == 1
    assert nova_definicao["resultado"]["chamada_anthropic"] is False
    assert obter_analise(escopo, calculo["calculo"]["id"])["resultado"]["dataset_versao"] == 1

    anterior = obter_dataset(escopo, dataset_id)
    assert anterior["versao"] == 1
    assert anterior["impressao"] == impressao_antiga
    novo = obter_dataset(escopo, resultado["dataset_id"])
    assert novo["versao_anterior_id"] == dataset_id
    assert novo["chave"] == anterior["chave"]
    assert novo["usuario_id"] == escopo.usuario_id
    assert novo["conta_id"] == escopo.conta_id
    assert novo["franquia_id"] == escopo.franquia_id

    db.session.remove()
    assert _soma(consultar_dashboard(escopo, antigo_id)["consulta"]) == "25.25"
    assert consultar_dashboard(escopo, antigo_id)["consulta"]["versao"] == 1
    assert IaChamadaTentativa.query.count() == antes_tentativas
    assert IaConsumoEvento.query.count() == antes_consumo


def test_acrescimo_compativel_soma_as_linhas_novas(ctx):
    _usuario, _franquia, escopo, _definicao, calculo, painel = _painel(
        "atualizacao.acrescentar@test.com",
        "atualizacao-acrescentar",
    )
    antes = _contagens(escopo)
    resultado = _atualizar(escopo, painel["dashboard"]["id"], _movimento(), "acrescentar")
    assert resultado["recalculado"] is True
    assert resultado["modo"] == "acrescentar"
    assert resultado["chamada_anthropic"] is False
    assert _soma(resultado["consulta"]) == "50.5"
    assert _painel_soma(resultado["consulta"]["graficos"])["elegiveis"] == 6
    assert resultado["versao"] == 2
    assert resultado["dataset_anterior_id"] == calculo["calculo"]["resultado"]["dataset_id"]
    datasets, analises, paineis = _contagens(escopo)
    assert datasets == antes[0] + 1
    assert analises > antes[1]
    assert paineis == antes[2] + 1
    assert _soma(consultar_dashboard(escopo, painel["dashboard"]["id"])["consulta"]) == "25.25"


def test_reenvio_do_mesmo_lote_preserva_o_acrescimo(ctx):
    _usuario, _franquia, escopo, _definicao, _calculo, painel = _painel(
        "atualizacao.reenvio@test.com",
        "atualizacao-reenvio",
    )
    original = painel["dashboard"]["id"]
    assert _soma(painel["consulta"]) == "25.25"
    assert _painel_soma(painel["consulta"]["graficos"])["elegiveis"] == 3

    primeiro_arquivo = _movimento()
    primeiro = _atualizar(escopo, original, primeiro_arquivo, "acrescentar")
    assert _soma(primeiro["consulta"]) == "50.5"
    assert _painel_soma(primeiro["consulta"]["graficos"])["elegiveis"] == 6
    assert primeiro["versao"] == 2
    gravado = obter_dataset(escopo, primeiro["dataset_id"])
    assert gravado["metadados_preparacao"]["lotes_aplicados"] == [_sha(primeiro_arquivo)]
    assert _fora_do_contrato(gravado["metadados_preparacao"]) is False
    perfil = perfilar_dataset(escopo, primeiro["dataset_id"])
    definicao = obter_analise(escopo, primeiro["definicao_id"])
    assert definicao["resultado"]["vinculo_preparacao"] == perfil["vinculo_preparacao"]
    assert _soma(consultar_dashboard(escopo, original)["consulta"]) == "25.25"

    antes = _artefatos(escopo)
    impressao = primeiro["impressao"]
    with pytest.raises(AtualizacaoErro) as erro:
        _atualizar(escopo, primeiro["dashboard"]["id"], primeiro_arquivo, "acrescentar")
    assert erro.value.codigo == "lote_ja_aplicado"
    assert erro.value.status == 422
    assert _artefatos(escopo) == antes
    assert obter_dataset(escopo, primeiro["dataset_id"])["impressao"] == impressao
    assert obter_dataset(escopo, primeiro["dataset_id"])["versao"] == 2
    preservado = consultar_dashboard(escopo, primeiro["dashboard"]["id"])
    assert _soma(preservado["consulta"]) == "50.5"
    assert _painel_soma(preservado["consulta"]["graficos"])["elegiveis"] == 6
    assert preservado["consulta"]["versao"] == 2

    outro_arquivo = _movimento_distinto()
    assert _sha(outro_arquivo) != _sha(primeiro_arquivo)
    segundo = _atualizar(escopo, primeiro["dashboard"]["id"], outro_arquivo, "acrescentar")
    assert segundo["recalculado"] is True
    assert _soma(segundo["consulta"]) == "75.75"
    assert _painel_soma(segundo["consulta"]["graficos"])["elegiveis"] == 9
    assert segundo["versao"] == 3
    linhagem = obter_dataset(escopo, segundo["dataset_id"])["metadados_preparacao"]["lotes_aplicados"]
    assert linhagem == [_sha(primeiro_arquivo), _sha(outro_arquivo)]
    assert _fora_do_contrato(obter_dataset(escopo, segundo["dataset_id"])["metadados_preparacao"]) is False
    assert _soma(consultar_dashboard(escopo, primeiro["dashboard"]["id"])["consulta"]) == "50.5"

    antes_repetido = _artefatos(escopo)
    with pytest.raises(AtualizacaoErro) as repetido:
        _atualizar(escopo, segundo["dashboard"]["id"], outro_arquivo, "acrescentar")
    assert repetido.value.codigo == "lote_ja_aplicado"
    assert _artefatos(escopo) == antes_repetido
    assert _soma(consultar_dashboard(escopo, segundo["dashboard"]["id"])["consulta"]) == "75.75"


def test_mesmo_arquivo_reutiliza_a_versao_sem_nova_revisao(ctx):
    conteudo = _movimento()
    _usuario, _franquia, escopo, _perfil, pacote = _jornada_de(
        "atualizacao.igual@test.com",
        "atualizacao-igual",
        conteudo,
    )
    definicao = definir_selecao(escopo, {"analise_id": pacote["id"]})
    calculo = _calcular(escopo, definicao["definicao"]["id"])
    painel = montar_dashboard(escopo, {"analise_id": calculo["calculo"]["id"]})
    antes = _contagens(escopo)
    antes_tentativas = IaChamadaTentativa.query.count()
    resultado = _atualizar(escopo, painel["dashboard"]["id"], conteudo, "substituir")
    assert resultado["recalculado"] is False
    assert resultado["reutilizado"] is True
    assert resultado["dashboard"]["id"] == painel["dashboard"]["id"]
    assert resultado["versao"] == 1
    assert _soma(resultado["consulta"]) == "25.25"
    assert _contagens(escopo) == antes
    assert IaChamadaTentativa.query.count() == antes_tentativas


def test_tipo_incompativel_nao_grava_versao_nem_dashboard(ctx):
    _usuario, _franquia, escopo, _definicao, _calculo, painel = _painel(
        "atualizacao.tipo@test.com",
        "atualizacao-tipo",
    )
    antes = _contagens(escopo)
    with pytest.raises(AtualizacaoErro) as erro:
        _atualizar(escopo, painel["dashboard"]["id"], _tipo_incompativel(), "substituir")
    assert erro.value.codigo == "atualizacao_incompativel"
    assert erro.value.status == 422
    assert _contagens(escopo) == antes
    assert _soma(consultar_dashboard(escopo, painel["dashboard"]["id"])["consulta"]) == "25.25"


def test_troca_de_colunas_com_nomes_parecidos_nao_remapeia(ctx):
    _usuario, _franquia, escopo, _definicao, _calculo, painel = _painel(
        "atualizacao.troca@test.com",
        "atualizacao-troca",
    )
    antes = _contagens(escopo)
    with pytest.raises(AtualizacaoErro) as erro:
        _atualizar(escopo, painel["dashboard"]["id"], _colunas_trocadas(), "substituir")
    assert erro.value.codigo == "mudanca_significado"
    assert _contagens(escopo) == antes
    assert _soma(consultar_dashboard(escopo, painel["dashboard"]["id"])["consulta"]) == "25.25"


def test_rotulos_parecidos_bloqueiam_por_ambiguidade(ctx):
    _usuario, _franquia, escopo, _definicao, _calculo, painel = _painel(
        "atualizacao.rotulo@test.com",
        "atualizacao-rotulo",
    )
    antes = _contagens(escopo)
    with pytest.raises(AtualizacaoErro) as erro:
        _atualizar(escopo, painel["dashboard"]["id"], _nomes_parecidos(), "acrescentar")
    assert erro.value.codigo == "ambiguidade_relevante"
    assert _contagens(escopo) == antes


def test_isolamento_entre_usuarios_e_contas(ctx):
    usuario, franquia, escopo, _definicao, _calculo, painel = _painel(
        "atualizacao.dono@test.com",
        "atualizacao-dono",
    )
    _outra_conta, _outra_franquia, outro = _par("atualizacao-outro", "atualizacao.outro@test.com")
    outro_escopo = escopo_do_usuario(outro)
    colega = seed_usuario(franquia.id, usuario.conta_id, email="atualizacao.colega@test.com")
    colega_escopo = escopo_do_usuario(colega)
    painel_id = painel["dashboard"]["id"]
    for alheio in (outro_escopo, colega_escopo):
        with pytest.raises(AgenteBiErro) as erro:
            _atualizar(alheio, painel_id, _compativel(), "substituir")
        assert erro.value.codigo == "nao_encontrado"
        assert erro.value.status == 404
    resultado = _atualizar(escopo, painel_id, _compativel(), "substituir")
    for alheio in (outro_escopo, colega_escopo):
        with pytest.raises(AgenteBiErro) as erro:
            consultar_dashboard(alheio, resultado["dashboard"]["id"])
        assert erro.value.codigo == "nao_encontrado"


def test_reabre_o_dashboard_atualizado_sem_recalcular_nem_anthropic(ctx, monkeypatch):
    _usuario, _franquia, escopo, _definicao, _calculo, painel = _painel(
        "atualizacao.reabre@test.com",
        "atualizacao-reabre",
    )
    resultado = _atualizar(escopo, painel["dashboard"]["id"], _compativel(), "substituir")
    esperado = resultado["consulta"]["graficos"]
    antes = AgenteBiAnalise.query.count()

    def proibido(*_args, **_kwargs):
        raise AssertionError("recalculo")

    monkeypatch.setattr("app.services.agente_bi_dashboard_service.calcular_definicao", proibido)
    monkeypatch.setattr("app.services.agente_bi_atualizacao_service.calcular_definicao", proibido)
    db.session.remove()
    reaberto = consultar_dashboard(escopo, resultado["dashboard"]["id"])
    assert reaberto["consulta"]["graficos"] == esperado
    assert reaberto["consulta"]["chamada_anthropic"] is False
    assert reaberto["consulta"]["ia_calculo"] is False
    assert _soma(reaberto["consulta"]) == "29.5"
    assert AgenteBiAnalise.query.count() == antes


def test_preparacao_em_memoria_nao_grava_dataset(ctx):
    _usuario, _franquia, escopo = _painel("atualizacao.memoria@test.com", "atualizacao-memoria")[:3]
    antes = AgenteBiDataset.query.count()
    preparada = ingerir(escopo, _arquivos(_movimento()), chave="lote-memoria", persistir=False)
    assert preparada["persistido"] is False
    assert preparada["dataset"] is None
    assert preparada["preparacao"]["estrutura"]["contrato"] == "preparar.v1"
    assert AgenteBiDataset.query.count() == antes


def test_api_atualiza_reabre_e_devolve_encaminhamento(app, monkeypatch):
    client = _cliente_ingestao(app, monkeypatch)
    with app.app_context():
        _usuario, _franquia, escopo, _definicao, _calculo, painel = _painel(
            "atualizacao.api@test.com",
            "atualizacao-api",
        )
        usuario_id = escopo.usuario_id
        painel_id = painel["dashboard"]["id"]
        antes_tentativas = IaChamadaTentativa.query.count()
        antes_consumo = IaConsumoEvento.query.count()
    assert client.post(f"/api/agente-bi/dashboards/{painel_id}/atualizacao").status_code == 401
    _entrar(client, usuario_id)
    bloqueado = client.post(
        f"/api/agente-bi/dashboards/{painel_id}/atualizacao",
        data={
            "modo": "substituir",
            "arquivos": (
                io.BytesIO(_nomes_parecidos()),
                "movimento.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
        content_type="multipart/form-data",
    )
    assert bloqueado.status_code == 422
    assert bloqueado.get_json()["codigo"] == "ambiguidade_relevante"
    assert bloqueado.get_json()["encaminhamento"] == "reinterpretacao_autorizada"
    resposta = client.post(
        f"/api/agente-bi/dashboards/{painel_id}/atualizacao",
        data={
            "modo": "substituir",
            "arquivos": (
                io.BytesIO(_compativel()),
                "movimento.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
        content_type="multipart/form-data",
    )
    assert resposta.status_code == 201
    corpo = resposta.get_json()
    assert corpo["chamada_anthropic"] is False
    assert _soma(corpo["consulta"]) == "29.5"
    reaberto = client.get(f"/api/agente-bi/dashboards/{corpo['dashboard']['id']}/consulta")
    assert reaberto.status_code == 200
    assert _soma(reaberto.get_json()["consulta"]) == "29.5"
    with app.app_context():
        assert IaChamadaTentativa.query.count() == antes_tentativas
        assert IaConsumoEvento.query.count() == antes_consumo


def test_rota_permanece_fechada_sem_a_flag_de_ingestao(app, monkeypatch):
    client = _cliente_ingestao(app, monkeypatch, ingestao=False)
    with app.app_context():
        _usuario, _franquia, escopo, _definicao, _calculo, painel = _painel(
            "atualizacao.flag@test.com",
            "atualizacao-flag",
        )
        usuario_id = escopo.usuario_id
        painel_id = painel["dashboard"]["id"]
    _entrar(client, usuario_id)
    resposta = client.post(
        f"/api/agente-bi/dashboards/{painel_id}/atualizacao",
        data={
            "modo": "substituir",
            "arquivos": (
                io.BytesIO(_movimento()),
                "movimento.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
        content_type="multipart/form-data",
    )
    assert resposta.status_code == 404


def _painel_filtrado(email, slug):
    from app.services.agente_bi_definir_service import definir
    from app.services.agente_bi_evolucao_service import evoluir_dashboard
    from tests.test_agente_bi_calculo import _coluna, _gravar, _indicador, _ref
    from tests.test_agente_bi_definir import _jornada

    _usuario, _franquia, escopo, perfil, pacote, _aprovada = _jornada(email, slug)
    valor = _coluna(perfil, dimensao_principal="valor", aba=0)
    pacote_um = _indicador(pacote, "soma", leitura="indicador_de_medida")
    pacote_um["graficos"][0]["medida"] = _ref(valor)
    documento = definir(perfil, pacote_um)
    salva = _gravar(escopo, pacote["id"], documento)
    calculo = _calcular(escopo, salva["id"])
    painel = montar_dashboard(escopo, {"analise_id": calculo["calculo"]["id"]})
    assert _soma(painel["consulta"]) == "25.25"
    filtrado = evoluir_dashboard(escopo, painel["dashboard"]["id"], {"texto": "filtre beta igual a sul"})
    assert filtrado["situacao"] == "aplicada", filtrado.get("esclarecimento")
    assert filtrado["chamada_anthropic"] is False
    assert set(filtrado["parametros"].values()) == {"sul"}
    assert _soma(filtrado["consulta"]) == "4.25"
    return escopo, painel, filtrado


def test_atualizacao_filtrada_encaminha_parametros_e_preserva_revisoes(ctx, caplog):
    escopo, painel, filtrado = _painel_filtrado("atualizacao.filtro@test.com", "atualizacao-filtro")
    caplog.set_level(logging.INFO)
    antes_tentativas = IaChamadaTentativa.query.count()
    antes_consumo = IaConsumoEvento.query.count()
    antes = _contagens(escopo)
    antigo_id = painel["dashboard"]["id"]
    filtrado_id = filtrado["dashboard"]["id"]
    parametros = filtrado["parametros"]

    resultado = atualizar_dashboard(
        escopo,
        filtrado_id,
        modo="substituir",
        arquivos=_arquivos(_compativel()),
        parametros=parametros,
    )

    assert resultado["contrato"] == "atualizacao.v1"
    assert resultado["recalculado"] is True
    assert resultado["chamada_anthropic"] is False
    assert resultado["ia_calculo"] is False
    assert "parametros" not in resultado
    assert _soma(resultado["consulta"]) == "4.25"
    assert _painel_soma(resultado["consulta"]["graficos"])["elegiveis"] == 1
    assert all(item.get("motivo") != "informar_parametro" for item in resultado["consulta"]["graficos"])
    assert resultado["versao"] == painel["consulta"]["versao"] + 1
    nova = obter_analise(escopo, resultado["definicao_id"])["resultado"]
    assert any(item.get("filtros") for item in nova["definicoes"])
    assert "sul" not in json.dumps(nova, ensure_ascii=False)
    assert "sul" not in db.session.get(AgenteBiDashboard, resultado["dashboard"]["id"]).definicoes_json
    assert "sul" not in "\n".join(
        registro.getMessage()
        for registro in caplog.records
        if registro.name.startswith("app.services.agente_bi_")
    )
    datasets, _analises, paineis = _contagens(escopo)
    assert datasets == antes[0] + 1
    assert paineis == antes[2] + 1
    db.session.remove()
    assert _soma(consultar_dashboard(escopo, antigo_id)["consulta"]) == "25.25"
    assert _soma(consultar_dashboard(escopo, filtrado_id)["consulta"]) == "4.25"
    assert consultar_dashboard(escopo, resultado["dashboard"]["id"])["consulta"]["versao"] == resultado["versao"]
    assert IaChamadaTentativa.query.count() == antes_tentativas
    assert IaConsumoEvento.query.count() == antes_consumo


def test_atualizacao_filtrada_sem_parametros_recusa_publicacao(ctx):
    escopo, painel, filtrado = _painel_filtrado("atualizacao.sem.param@test.com", "atualizacao-sem-param")
    antes_tentativas = IaChamadaTentativa.query.count()
    antes = _contagens(escopo)
    filtrado_id = filtrado["dashboard"]["id"]

    with pytest.raises(AtualizacaoErro) as ausente:
        _atualizar(escopo, filtrado_id, _compativel(), "substituir")
    assert ausente.value.codigo == "informar_parametro"
    assert ausente.value.status == 422

    with pytest.raises(AtualizacaoErro) as estranho:
        atualizar_dashboard(
            escopo,
            filtrado_id,
            modo="substituir",
            arquivos=_arquivos(_compativel()),
            parametros={"nao_declarado": "sul"},
        )
    assert estranho.value.codigo == "parametro_inconsistente"

    assert _contagens(escopo) == antes
    assert _soma(consultar_dashboard(escopo, filtrado_id)["consulta"]) == "4.25"
    assert _soma(consultar_dashboard(escopo, painel["dashboard"]["id"])["consulta"]) == "25.25"
    assert IaChamadaTentativa.query.count() == antes_tentativas
    publicados = [
        item
        for item in consultar_dashboard(escopo, filtrado_id)["consulta"]["graficos"]
        if item.get("motivo") == "informar_parametro" and item.get("status") == "bloqueado"
    ]
    assert publicados == []


def test_api_atualizacao_filtrada_falha_fechada_e_recalcula_com_parametro(app, monkeypatch):
    client = _cliente_ingestao(app, monkeypatch)
    with app.app_context():
        escopo, _painel, filtrado = _painel_filtrado("atualizacao.api.filtro@test.com", "atualizacao-api-filtro")
        usuario_id = escopo.usuario_id
        filtrado_id = filtrado["dashboard"]["id"]
        parametros = filtrado["parametros"]
        antes = _contagens(escopo)
        antes_tentativas = IaChamadaTentativa.query.count()
    _entrar(client, usuario_id)

    def envio(parametros=None):
        dados = {
            "modo": "substituir",
            "arquivos": (
                io.BytesIO(_compativel()),
                "movimento.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        }
        if parametros is not None:
            dados["parametros"] = parametros
        return dados

    recusa = client.post(
        f"/api/agente-bi/dashboards/{filtrado_id}/atualizacao",
        data=envio(),
        content_type="multipart/form-data",
    )
    assert recusa.status_code == 422
    assert recusa.get_json()["codigo"] == "informar_parametro"
    assert "sul" not in recusa.get_data(as_text=True)
    invalido = client.post(
        f"/api/agente-bi/dashboards/{filtrado_id}/atualizacao",
        data=envio("{"),
        content_type="multipart/form-data",
    )
    assert invalido.status_code == 422
    assert invalido.get_json()["codigo"] == "parametro_inconsistente"
    with app.app_context():
        assert _contagens(escopo) == antes
        assert _soma(consultar_dashboard(escopo, filtrado_id)["consulta"]) == "4.25"
    resposta = client.post(
        f"/api/agente-bi/dashboards/{filtrado_id}/atualizacao",
        data=envio(json.dumps(parametros)),
        content_type="multipart/form-data",
    )
    assert resposta.status_code == 201
    corpo = resposta.get_json()
    assert corpo["contrato"] == "atualizacao.v1"
    assert corpo["recalculado"] is True
    assert corpo["chamada_anthropic"] is False
    assert "parametros" not in corpo
    assert _soma(corpo["consulta"]) == "4.25"
    assert '"sul"' not in resposta.get_data(as_text=True)
    with app.app_context():
        assert _soma(consultar_dashboard(escopo, filtrado_id)["consulta"]) == "4.25"
        assert IaChamadaTentativa.query.count() == antes_tentativas


def test_fontes_nao_chamam_modelo_nem_leem_imagem():
    servico = (ROOT / "app" / "services" / "agente_bi_atualizacao_service.py").read_text(encoding="utf-8")
    ingestao = (ROOT / "app" / "services" / "agente_bi_ingestao_service.py").read_text(encoding="utf-8")
    script = (ROOT / "app" / "static" / "js" / "agente_bi_dashboard.js").read_text(encoding="utf-8")
    assert "chamar_anthropic" not in servico
    assert "chamar_anthropic" not in script
    assert "def _classificar(" not in servico
    assert "def _classificar(" not in script
    assert not any(linha.strip() == "import re" for linha in servico.splitlines())
    assert not any(linha.strip() == "import re" for linha in script.splitlines())
    assert "PIL" not in servico
    assert "ocr" not in servico.lower()
    assert "persistir: bool = True" in ingestao
    assert "innerHTML" not in script

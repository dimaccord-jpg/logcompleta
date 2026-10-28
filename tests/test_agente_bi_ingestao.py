"""SCRUM-254: ingestão determinística de XLSX/CSV, sem IA e sem upload público."""
from __future__ import annotations

import io
import json
import logging
import zipfile
from datetime import date
from pathlib import Path

import flask_login.utils
import pytest
from openpyxl import Workbook

from app.agente_bi_api_routes import agente_bi_api_bp
from app.agente_bi_flags import agente_bi_ingestao_habilitada
from app.extensions import db, login_manager
from app.infra import get_user_by_id
from app.models import (
    AgenteBiAnalise,
    AgenteBiDashboard,
    AgenteBiDataset,
    AgenteBiDatasetBloco,
    IaChamadaTentativa,
)
from app.services.agente_bi_ingestao_service import (
    LIMITE_CANDIDATOS_PUBLICADOS,
    TETO_BYTES,
    TETO_CELULAS,
    ArquivoBruto,
    IngestaoErro,
    ingerir,
    resolver_limites,
)
from app.services.agente_bi_persistencia_service import (
    MAX_JSON_BYTES,
    MAX_LISTA,
    AgenteBiErro,
    _recompor_segmentado,
    calcular_impressao_dataset,
    criar_dashboard,
    criar_dataset,
    escopo_do_usuario,
    obter_dataset,
)
from tests.agente_bi_fonte import fonte_sem_docstring_de_modulo
from tests.conftest import seed_conta_franquia_cliente, seed_usuario
from tests.test_agente_bi_fundacao import _analise, _dashboard, _dataset

ROOT = Path(__file__).resolve().parents[1]


def _par(slug, email):
    conta, franquia = seed_conta_franquia_cliente(slug=slug)
    usuario = seed_usuario(franquia.id, conta.id, email=email)
    return conta, franquia, usuario


def _arquivo(conteudo: bytes, nome: str) -> ArquivoBruto:
    return ArquivoBruto(nome=nome, mime=None, conteudo=conteudo)


def _ingerir(usuario, conteudo: bytes, nome: str, *, chave="lote-sintetico", **kwargs):
    return ingerir(
        escopo_do_usuario(usuario),
        [_arquivo(conteudo, nome)],
        chave=chave,
        **kwargs,
    )


def _estrutura(resultado):
    return resultado["preparacao"]["estrutura"]


def _abas(resultado):
    return _estrutura(resultado)["abas"]


def _regiao(resultado, aba=0, regiao=0):
    return next(
        item
        for item in _estrutura(resultado)["regioes"]
        if item["aba"] == aba and item["indice"] == regiao
    )


def _dados(resultado, *, aba=None, regiao=None):
    linhas = resultado["preparacao"]["dados_tipados"]
    if aba is not None:
        linhas = [item for item in linhas if item["aba"] == aba]
    if regiao is not None:
        linhas = [item for item in linhas if item["regiao"] == regiao]
    return linhas


def _valor(celula):
    if celula["tipo"] == "ausente":
        return None
    return celula.get("valor")


def _xlsx(planilhas: dict[str, list[list]], ocultas=()):
    livro = Workbook()
    primeira = True
    for nome, linhas in planilhas.items():
        aba = livro.active if primeira else livro.create_sheet(nome)
        primeira = False
        aba.title = nome
        if nome in ocultas:
            aba.sheet_state = "hidden"
        for i, linha in enumerate(linhas, start=1):
            for j, valor in enumerate(linha, start=1):
                if valor is not None:
                    aba.cell(i, j, valor)
    destino = io.BytesIO()
    livro.save(destino)
    return destino.getvalue()


def _xlsx_formulas() -> bytes:
    livro = Workbook()
    aba = livro.active
    aba.title = "Calculo"
    aba["A1"] = "item"
    aba["B1"] = "total"
    aba["A2"] = "caixa"
    aba["B2"] = "=1+2"
    aba["A3"] = "fita"
    aba["B3"] = "=2+2"
    aba["B4"] = "=A2"
    aba["B5"] = "=A2"
    destino = io.BytesIO()
    livro.save(destino)
    bruto = destino.getvalue()
    with zipfile.ZipFile(io.BytesIO(bruto)) as arquivo:
        nomes = arquivo.namelist()
        pecas = {nome: arquivo.read(nome) for nome in nomes}
    xml = pecas["xl/worksheets/sheet1.xml"]
    xml = xml.replace(b"<f>1+2</f><v></v>", b"<f>1+2</f><v>3</v>", 1)
    pecas["xl/worksheets/sheet1.xml"] = xml
    saida = io.BytesIO()
    with zipfile.ZipFile(saida, "w", compression=zipfile.ZIP_DEFLATED) as arquivo:
        for nome, peca in pecas.items():
            arquivo.writestr(nome, peca)
    return saida.getvalue()


def _cliente(app, monkeypatch, *, api=True, ingestao=True, permitido=True):
    app.config["SECRET_KEY"] = "teste-agente-bi-ingestao"
    app.config["TESTING"] = True
    login_manager.init_app(app)

    @login_manager.user_loader
    def _load_user(user_id):
        return get_user_by_id(user_id)

    if "agente_bi_api" not in app.blueprints:
        app.register_blueprint(agente_bi_api_bp)

    if api:
        monkeypatch.setenv("AGENTE_BI_API_ENABLED", "1")
    else:
        monkeypatch.delenv("AGENTE_BI_API_ENABLED", raising=False)
    if ingestao:
        monkeypatch.setenv("AGENTE_BI_INGESTAO_ENABLED", "1")
    else:
        monkeypatch.delenv("AGENTE_BI_INGESTAO_ENABLED", raising=False)
    monkeypatch.setattr(
        "app.agente_bi_api_routes.avaliar_autorizacao_operacao_por_franquia",
        lambda _usuario: {"permitido": permitido},
    )
    return app.test_client()


def _entrar(client, usuario_id: int) -> None:
    with client.session_transaction() as sess:
        sess["_user_id"] = str(usuario_id)
        sess["_fresh"] = True


def _postar(client, conteudo: bytes, nome: str, *, chave="lote-sintetico"):
    tipo = "text/csv" if nome.lower().endswith(".csv") else (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    return client.post(
        "/api/agente-bi/ingestao",
        data={"chave": chave, "arquivos": (io.BytesIO(conteudo), nome, tipo)},
        content_type="multipart/form-data",
    )


def test_flag_de_ingestao_desligada_por_padrao(monkeypatch):
    monkeypatch.delenv("AGENTE_BI_INGESTAO_ENABLED", raising=False)
    assert agente_bi_ingestao_habilitada() is False
    monkeypatch.setenv("AGENTE_BI_INGESTAO_ENABLED", "talvez")
    assert agente_bi_ingestao_habilitada() is False
    monkeypatch.setenv("AGENTE_BI_INGESTAO_ENABLED", "on")
    assert agente_bi_ingestao_habilitada() is True


def test_limites_efetivos_respeitam_o_ambiente_e_o_teto(app, monkeypatch):
    with app.app_context():
        limites = resolver_limites()
    assert limites.max_bytes_xlsx == 5 * 1024 * 1024
    assert limites.max_bytes_csv == 2 * 1024 * 1024
    assert limites.max_bytes_total == 15 * 1024 * 1024
    assert limites.max_arquivos == 3
    assert limites.max_celulas == TETO_CELULAS
    assert limites.max_bytes_xlsx <= TETO_BYTES

    monkeypatch.setenv("AGENTE_BI_INGESTAO_MAX_BYTES", str(50 * 1024 * 1024))
    monkeypatch.setenv("AGENTE_BI_INGESTAO_MAX_CELULAS", "900000")
    monkeypatch.setenv("AGENTE_BI_INGESTAO_MAX_ARQUIVOS", "9")
    from types import SimpleNamespace

    altos = SimpleNamespace(
        excel_max_bytes=50 * 1024 * 1024,
        csv_max_bytes=50 * 1024 * 1024,
        session_max_bytes=80 * 1024 * 1024,
    )
    travado = resolver_limites(altos)
    assert travado.max_bytes_xlsx == TETO_BYTES
    assert travado.max_celulas == TETO_CELULAS
    assert travado.max_arquivos == 3


def test_csv_convencional_tipa_sem_exigir_nome_de_coluna(ctx):
    _conta, _franquia, usuario = _par("ing-csv", "ing.csv@test.com")
    texto = "origem,destino,peso\nnorte,sul,12\nleste,oeste,4\n"
    resultado = _ingerir(usuario, texto.encode("utf-8"), "sintetico.csv")
    regiao = _regiao(resultado)
    assert regiao["cabecalho_deteccao"] == "inequivoca"
    assert regiao["cabecalho_linha"] == 1
    assert regiao["candidatos_amostra"] is False
    assert regiao["candidatos_total"] == 1
    assert regiao["candidatos"] == [
        {"linha": 1, "pontuacao": 9, "rotulos": ["origem", "destino", "peso"]}
    ]
    assert regiao["tipos_fisicos"] == ["texto", "texto", "numero"]
    assert regiao["linhas_elegiveis"] == 2
    linhas = _dados(resultado)
    assert [_valor(item) for item in linhas[0]["valores"]] == ["norte", "sul", 12]
    assert linhas[1]["valores"][2]["valor"] == 4
    assert resultado["persistido"] is True
    assert resultado["dataset"]["versao"] == 1
    assert resultado["preparacao"]["metadados_preparacao"]["ia"] is False
    assert resultado["preparacao"]["metadados_preparacao"]["envio_externo"] is False


def test_csv_com_cabecalho_deslocado_e_separador_ponto_e_virgula(ctx):
    _conta, _franquia, usuario = _par("ing-deslocado", "ing.deslocado@test.com")
    texto = "Relatorio interno\norigem;destino;peso\nnorte;sul;10\n"
    resultado = _ingerir(usuario, texto.encode("utf-8"), "deslocado.csv")
    aba = _abas(resultado)[0]
    assert aba["separador"] == "ponto_e_virgula"
    regiao = _regiao(resultado)
    assert regiao["cabecalho_deteccao"] == "inequivoca"
    assert regiao["cabecalho_linha"] == 2
    assert regiao["preambulo"][0]["linha"] == 1
    assert regiao["preambulo"][0]["rotulos"][0] == "Relatorio interno"
    assert len(_dados(resultado)) == 1
    assert _dados(resultado)[0]["linha"] == 3
    assert _dados(resultado)[0]["valores"][2]["valor"] == 10


def test_xlsx_com_varias_abas_e_data(ctx):
    _conta, _franquia, usuario = _par("ing-abas", "ing.abas@test.com")
    conteudo = _xlsx(
        {
            "movimento": [
                ["origem", "destino", "dia"],
                ["norte", "sul", date(2024, 3, 15)],
            ],
            "faixas": [
                ["codigo", "limite"],
                ["n1", 10],
            ],
        }
    )
    resultado = _ingerir(usuario, conteudo, "varias.xlsx")
    abas = _abas(resultado)
    assert [aba["nome"] for aba in abas] == ["movimento", "faixas"]
    assert all(aba["oculta"] is False for aba in abas)
    assert _regiao(resultado, 0)["tipos_fisicos"][2] == "data"
    assert _dados(resultado, aba=0)[0]["valores"][2] == {"tipo": "data", "valor": "2024-03-15"}
    assert _dados(resultado, aba=1)[0]["valores"][0]["valor"] == "n1"


def test_aba_oculta_entra_na_preparacao(ctx):
    _conta, _franquia, usuario = _par("ing-oculta", "ing.oculta@test.com")
    conteudo = _xlsx(
        {
            "movimento": [["origem", "destino"], ["norte", "sul"]],
            "apoio": [["codigo", "limite"], ["n1", 10]],
        },
        ocultas={"apoio"},
    )
    resultado = _ingerir(usuario, conteudo, "oculta.xlsx")
    abas = _abas(resultado)
    oculta = abas[1]
    assert oculta["nome"] == "apoio"
    assert oculta["oculta"] is True
    assert oculta["indicio_auxiliar"] is True
    assert "aba_oculta" in oculta["alertas"]
    assert _dados(resultado, aba=1)[0]["valores"][1]["valor"] == 10


def test_formula_com_cache_e_sem_cache_nao_vira_registro_vazio(ctx):
    _conta, _franquia, usuario = _par("ing-formula", "ing.formula@test.com")
    resultado = _ingerir(usuario, _xlsx_formulas(), "formulas.xlsx")
    regiao = _regiao(resultado)
    assert regiao["formulas_com_cache"] >= 1
    assert regiao["formulas_sem_cache"] >= 1
    assert regiao["linhas_pre_preenchidas"] == 2
    assert "linha_pre_preenchida" in regiao["alertas"]
    assert "formula_sem_cache" in regiao["alertas"]
    linhas = _dados(resultado)
    assert [item["linha"] for item in linhas] == [2, 3]
    formula = linhas[0]["valores"][1]
    assert formula["tipo"] == "formula"
    assert formula["formula"] == "=1+2"
    assert formula["cache_estado"] == "presente"
    assert formula["valor"] == 3
    sem_cache = linhas[1]["valores"][1]
    assert sem_cache["formula"] == "=2+2"
    assert sem_cache["cache_estado"] == "ausente"
    assert "valor" not in sem_cache
    fora = [item for item in regiao["formulas_observadas"] if item["registro"] is False]
    assert {item["linha"] for item in fora} == {4, 5}
    assert all(item["cache_estado"] == "ausente" for item in fora)


def test_arquivo_sem_cabecalho_inequivoco_nao_escolhe_interpretacao(ctx):
    _conta, _franquia, usuario = _par("ing-sem-cab", "ing.semcab@test.com")
    numeros = _ingerir(usuario, b"1,2\n3,4\n", "numeros.csv", chave="lote-numeros")
    regiao = _regiao(numeros)
    assert regiao["cabecalho_deteccao"] == "ausente"
    assert regiao["cabecalho_linha"] is None
    assert len(_dados(numeros)) == 2
    assert _dados(numeros)[0]["valores"][0]["valor"] == 1

    ambiguo = _ingerir(
        usuario,
        "grupo,familia\norigem,destino\nnorte,sul\n".encode("utf-8"),
        "ambiguo.csv",
        chave="lote-ambiguo",
    )
    cabecalho = _regiao(ambiguo)
    assert cabecalho["cabecalho_deteccao"] == "ambigua"
    assert cabecalho["cabecalho_linha"] is None
    assert cabecalho["candidatos_amostra"] is False
    assert cabecalho["candidatos_total"] == len(cabecalho["candidatos"]) == 3
    assert len(_dados(ambiguo)) == 3
    assert any(item["codigo"] == "cabecalho_ambiguo" for item in ambiguo["preparacao"]["estrutura"]["ambiguidades"])


def test_regioes_empilhadas_e_laterais_ambiguas(ctx):
    _conta, _franquia, usuario = _par("ing-regioes", "ing.regioes@test.com")
    empilhadas = _xlsx(
        {
            "blocos": [
                ["origem", "destino", "peso"],
                ["norte", "sul", 10],
                [None, None, None],
                ["codigo", "limite"],
                ["n1", 3],
            ]
        }
    )
    resultado = _ingerir(usuario, empilhadas, "blocos.xlsx", chave="lote-blocos")
    regioes = [item for item in _estrutura(resultado)["regioes"] if item["aba"] == 0]
    assert len(regioes) == 2
    assert regioes[0]["cabecalho_linha"] == 1
    assert regioes[1]["linha_inicio"] == 4
    assert "regioes_multiplas" in _abas(resultado)[0]["alertas"]

    laterais = _xlsx(
        {
            "lado": [
                ["origem", "destino", None, "codigo", "limite"],
                ["norte", "sul", None, "n1", 3],
            ]
        }
    )
    lado = _ingerir(usuario, laterais, "lado.xlsx", chave="lote-lado")
    regiao = _regiao(lado)
    assert len(_estrutura(lado)["regioes"]) == 1
    assert regiao["divisao_escolha"] is None
    assert len(regiao["divisao_candidatos"]) == 2
    assert any(item["codigo"] == "divisao_lateral_ambigua" for item in lado["preparacao"]["estrutura"]["ambiguidades"])
    assert _dados(lado)[0]["valores"][2]["tipo"] == "ausente"


def test_cabecalho_repetido_nao_entra_como_registro(ctx):
    _conta, _franquia, usuario = _par("ing-repete", "ing.repete@test.com")
    texto = "origem,destino,peso\nnorte,sul,10\norigem,destino,peso\nleste,oeste,4\n"
    resultado = _ingerir(usuario, texto.encode("utf-8"), "repete.csv")
    regiao = _regiao(resultado)
    assert regiao["linhas_cabecalho_repetido"] == 1
    assert regiao["linhas_elegiveis"] == 2
    assert [item["linha"] for item in _dados(resultado)] == [2, 4]


def test_arquivo_corrompido_indevido_e_macro_nao_sao_executados(ctx):
    _conta, _franquia, usuario = _par("ing-ruim", "ing.ruim@test.com")
    with pytest.raises(IngestaoErro) as corrompido:
        _ingerir(usuario, b"PK\x03\x04isto-nao-e-zip", "quebrado.xlsx")
    assert corrompido.value.codigo == "arquivo_corrompido"

    with pytest.raises(IngestaoErro) as pdf:
        _ingerir(usuario, b"%PDF-1.4 falso", "falso.csv")
    assert pdf.value.codigo == "arquivo_indevido"

    with pytest.raises(IngestaoErro) as extensao:
        _ingerir(usuario, b"origem,destino\n", "notas.exe")
    assert extensao.value.codigo == "arquivo_indevido"

    base = _xlsx({"movimento": [["origem"], ["norte"]]})
    saida = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base)) as origem, zipfile.ZipFile(saida, "w") as destino:
        for info in origem.infolist():
            destino.writestr(info.filename, origem.read(info.filename))
        destino.writestr("xl/vbaProject.bin", b"macro")
    with pytest.raises(IngestaoErro) as macro:
        _ingerir(usuario, saida.getvalue(), "macro.xlsx")
    assert macro.value.codigo == "arquivo_indevido"
    assert AgenteBiDataset.query.count() == 0


def test_limites_de_celulas_bytes_e_descompressao(ctx, monkeypatch):
    _conta, _franquia, usuario = _par("ing-limite", "ing.limite@test.com")
    monkeypatch.setenv("AGENTE_BI_INGESTAO_MAX_CELULAS", "3")
    with pytest.raises(IngestaoErro) as celulas:
        _ingerir(usuario, b"origem,destino\nnorte,sul\n", "largo.csv")
    assert celulas.value.codigo == "limite_celulas"

    monkeypatch.delenv("AGENTE_BI_INGESTAO_MAX_CELULAS", raising=False)
    monkeypatch.setenv("AGENTE_BI_INGESTAO_MAX_BYTES", "40")
    with pytest.raises(IngestaoErro) as tamanho:
        _ingerir(usuario, b"origem,destino,observacao\nnorte,sul,caixa-grande\n", "grande.csv", chave="lote-grande")
    assert tamanho.value.codigo == "limite_bytes"

    monkeypatch.delenv("AGENTE_BI_INGESTAO_MAX_BYTES", raising=False)
    monkeypatch.delenv("AGENTE_BI_INGESTAO_MAX_CELULAS", raising=False)
    monkeypatch.setenv("AGENTE_BI_INGESTAO_MAX_DESCOMPACTADO", "500")
    with pytest.raises(IngestaoErro) as zip_bomba:
        _ingerir(usuario, _xlsx({"movimento": [["origem"], ["norte"]]}), "cheio.xlsx", chave="lote-zip")
    assert zip_bomba.value.codigo == "limite_descompressao"
    assert AgenteBiDataset.query.count() == 0


def test_limite_de_tempo_e_de_arquivos(ctx):
    _conta, _franquia, usuario = _par("ing-tempo", "ing.tempo@test.com")
    chamadas = {"n": 0}

    def relogio():
        chamadas["n"] += 1
        return 0 if chamadas["n"] < 4 else 50

    with pytest.raises(IngestaoErro) as tempo:
        _ingerir(usuario, b"origem,destino\nnorte,sul\n", "tempo.csv", relogio=relogio)
    assert tempo.value.codigo == "limite_tempo"

    with pytest.raises(IngestaoErro) as arquivos:
        ingerir(
            escopo_do_usuario(usuario),
            [_arquivo(b"a\n1\n", f"arq{i}.csv") for i in range(4)],
            chave="lote-quatro",
        )
    assert arquivos.value.codigo == "limite_arquivos"


def test_isolamento_entre_usuarios_e_versao_imutavel(app, monkeypatch):
    client = _cliente(app, monkeypatch)
    with app.app_context():
        conta, franquia, dono = _par("ing-iso", "ing.iso.dono@test.com")
        colega = seed_usuario(franquia.id, conta.id, email="ing.iso.colega@test.com")
        dono_id = dono.id
        colega_id = colega.id
        reserva = franquia.reserva_pendente
        consumo = franquia.consumo_acumulado
        primeiro = _ingerir(dono, b"origem,destino,peso\nnorte,sul,8\n", "a.csv", chave="lote-iso")
        dataset_id = primeiro["dataset"]["id"]
        segundo = _ingerir(dono, b"origem,destino,peso\nleste,oeste,9\n", "b.csv", chave="lote-iso")
        repetido = _ingerir(dono, b"origem,destino,peso\nnorte,sul,8\n", "a.csv", chave="lote-iso")
        assert primeiro["dataset"]["versao"] == 1
        assert segundo["dataset"]["versao"] == 2
        assert segundo["dataset"]["versao_anterior_id"] == dataset_id
        assert repetido["reutilizado"] is True
        assert repetido["dataset"]["id"] == dataset_id
        assert AgenteBiDataset.query.filter_by(usuario_id=dono_id).count() == 2
        original = db.session.get(AgenteBiDataset, dataset_id).estrutura_json
        painel = criar_dashboard(escopo_do_usuario(dono), _dashboard(primeiro["dataset"], estado_publicacao="publicado"))
        definicao = db.session.get(AgenteBiDashboard, painel["id"]).definicoes_json
        db.session.refresh(franquia)
        assert franquia.reserva_pendente == reserva
        assert franquia.consumo_acumulado == consumo
        assert IaChamadaTentativa.query.count() == 0
    _entrar(client, colega_id)
    leitura = client.get(f"/api/agente-bi/datasets/{dataset_id}")
    assert leitura.status_code == 404
    _entrar(client, dono_id)
    proprio = client.get(f"/api/agente-bi/datasets/{dataset_id}")
    assert proprio.status_code == 200
    assert proprio.get_json()["dataset"]["dados_tipados"][0]["valores"][0]["valor"] == "norte"
    with app.app_context():
        assert db.session.get(AgenteBiDataset, dataset_id).estrutura_json == original
        assert db.session.get(AgenteBiDashboard, painel["id"]).definicoes_json == definicao
        assert AgenteBiAnalise.query.count() == 0


def test_api_recusa_flag_auth_e_nao_abre_o_upload_publico(app, monkeypatch):
    client = _cliente(app, monkeypatch, api=False, ingestao=True)
    assert client.post("/api/agente-bi/ingestao").status_code == 404

    monkeypatch.setenv("AGENTE_BI_API_ENABLED", "1")
    monkeypatch.delenv("AGENTE_BI_INGESTAO_ENABLED", raising=False)
    assert _postar(client, b"origem\nnorte\n", "a.csv").status_code == 404

    monkeypatch.setenv("AGENTE_BI_INGESTAO_ENABLED", "1")
    anonimo = _postar(client, b"origem\nnorte\n", "a.csv")
    assert anonimo.status_code == 401

    with app.app_context():
        _conta, _franquia, usuario = _par("ing-gate", "ing.gate@test.com")
        usuario_id = usuario.id
    monkeypatch.setattr(
        "app.agente_bi_api_routes.avaliar_autorizacao_operacao_por_franquia",
        lambda _usuario: {"permitido": False},
    )
    _entrar(client, usuario_id)
    bloqueado = _postar(client, b"origem\nnorte\n", "a.csv")
    assert bloqueado.status_code == 403

    monkeypatch.setattr(
        "app.agente_bi_api_routes.avaliar_autorizacao_operacao_por_franquia",
        lambda _usuario: {"permitido": True},
    )
    legado = client.post(
        "/api/agente-bi/datasets",
        data={"arquivo": (io.BytesIO(b"origem\nnorte\n"), "a.csv")},
        content_type="multipart/form-data",
    )
    assert legado.status_code == 422
    html = (ROOT / "app" / "templates" / "agente_bi.html").read_text(encoding="utf-8")
    script = (ROOT / "app" / "static" / "js" / "agente_bi.js").read_text(encoding="utf-8")
    assert "/api/agente-bi/ingestao" not in html
    assert 'type="file"' not in html
    assert "Para começar, envie seus dados logísticos pelo botão +." in script
    assert "uploadItem.disabled = true" in script or "if (!uploadItem.disabled) uploadItem.disabled = true;" in script


def test_persistencia_recusada_quando_a_governanca_nao_autoriza_intacto(ctx, caplog):
    _conta, _franquia, usuario = _par("ing-gov", "ing.gov@test.com")
    documento = "123.456.789-09"
    caplog.set_level(logging.INFO)
    resultado = _ingerir(
        usuario,
        f"codigo,faixa,qtd\n{documento},ate10,1\n".encode("utf-8"),
        "restrito.csv",
    )
    assert resultado["persistido"] is False
    assert resultado["motivo"] == "conteudo_nao_permitido"
    assert resultado["dependencia"] == "SCRUM-252"
    assert resultado["dataset"] is None
    assert AgenteBiDataset.query.count() == 0
    assert resultado["preparacao"]["dados_tipados"][0]["valores"][1]["valor"] == "ate10"
    assert documento not in caplog.text
    assert "norte" not in caplog.text


def test_log_nao_grava_valor_de_celula(ctx, caplog):
    _conta, _franquia, usuario = _par("ing-log", "ing.log@test.com")
    caplog.set_level(logging.INFO)
    _ingerir(usuario, b"origem,destino,peso\nalfa,beta,8\n", "log.csv")
    assert "alfa" not in caplog.text
    assert "beta" not in caplog.text


def test_modulos_nao_chamam_ia_rede_nem_cobranca():
    for relativo in (
        "app/agente_bi_flags.py",
        "app/agente_bi_api_routes.py",
        "app/services/agente_bi_ingestao_service.py",
    ):
        fonte = fonte_sem_docstring_de_modulo(ROOT / relativo)
        assert "import " in fonte
        assert "def _classificar(" not in fonte
        assert not any(
            linha.strip().startswith("_classificar(") or " _classificar(" in f" {linha}"
            for linha in fonte.splitlines()
        )
        if relativo == "app/agente_bi_api_routes.py":
            _assert_re_so_valida_codigo_interno(fonte)
        else:
            assert not any(linha.strip() == "import re" for linha in fonte.splitlines())
        fonte = fonte.lower()
        for termo in (
            "anthropic",
            "openai",
            "gemini",
            "httpx",
            "urlopen",
            "subprocess",
            "thread(",
            "cleide_upload",
        ):
            assert termo not in fonte
    fonte_ingestao = (ROOT / "app" / "services" / "agente_bi_ingestao_service.py").read_text(encoding="utf-8").lower()
    assert "cpf" not in fonte_ingestao
    assert "criar_dataset" in fonte_ingestao


def _assert_re_so_valida_codigo_interno(fonte: str) -> None:
    """`import re` nas rotas só valida código interno. Não classifica intenção humana."""
    linhas = [linha.strip() for linha in fonte.splitlines()]
    assert "import re" in linhas
    compilacoes = [
        linha
        for linha in linhas
        if "re.compile" in linha
        or "re.search" in linha
        or "re.match" in linha
        or "re.findall" in linha
        or "re.sub" in linha
        or "re.fullmatch(" in linha
    ]
    assert compilacoes == ['_CODIGO_PUBLICO_ERRO = re.compile(r"^[a-z0-9_]{1,40}$")']
    assert any("_CODIGO_PUBLICO_ERRO.fullmatch(" in linha for linha in linhas)


def test_preparacao_reutiliza_dataset_sem_releitura_obrigatoria(ctx):
    _conta, _franquia, usuario = _par("ing-reuso", "ing.reuso@test.com")
    resultado = _ingerir(usuario, b"origem,destino,peso\nnorte,sul,8\n", "reuso.csv", chave="lote-reuso")
    gravado = db.session.get(AgenteBiDataset, resultado["dataset"]["id"])
    assert gravado.origem == "preparada"
    assert gravado.dados_tipados_json
    from app.services.agente_bi_persistencia_service import obter_dataset

    corpo = obter_dataset(escopo_do_usuario(usuario), resultado["dataset"]["id"])
    assert corpo["dados_tipados"] == resultado["preparacao"]["dados_tipados"]
    assert corpo["estrutura"]["regioes"][0]["linhas_elegiveis"] == 1
    assert corpo["metadados_preparacao"]["versao_preparacao"] == "preparar.v1"


def _literais_numericos(resultado, literais):
    achados = []
    for registro in _dados(resultado):
        for celula in registro["valores"]:
            if celula.get("valor") in literais:
                achados.append((celula["valor"], celula["tipo"]))
    return achados


def _csv_codex() -> bytes:
    """CSV sintético de 1.058.901 bytes e 10 mil registros de dados.

    A coluna do meio é numérica para o cabeçalho permanecer inequívoco.
    """
    alvo = 1_058_901
    cabecalho = b"categoria,medida,nota\n"
    base = [f"cat-{indice:05d},{indice},obs-{indice:05d}" for indice in range(10_000)]
    nucleo = sum(len(item) + 1 for item in base)
    falta = alvo - len(cabecalho) - nucleo
    assert falta >= 0
    quociente, resto = divmod(falta, len(base))
    linhas = []
    for indice, item in enumerate(base):
        complemento = quociente + (1 if indice < resto else 0)
        linhas.append(item + ("x" * complemento) + "\n")
    bruto = cabecalho + "".join(linhas).encode("ascii")
    assert len(bruto) == alvo
    assert max(len(item) for item in linhas) < 4_000
    return bruto


def _csv_textual_um_mega() -> bytes:
    """Caso A: 1.058.901 bytes, cabeçalho e 10 mil linhas, tudo textual."""
    alvo = 1_058_901
    cabecalho = b"categoria,medida,nota\n"
    base = [f"cat-{indice:05d},med-{indice:05d},obs-{indice:05d}" for indice in range(10_000)]
    nucleo = sum(len(item) + 1 for item in base)
    falta = alvo - len(cabecalho) - nucleo
    assert falta >= 0
    quociente, resto = divmod(falta, len(base))
    linhas = []
    for indice, item in enumerate(base):
        complemento = quociente + (1 if indice < resto else 0)
        linhas.append(item + ("x" * complemento) + "\n")
    bruto = cabecalho + "".join(linhas).encode("ascii")
    assert len(bruto) == alvo
    assert max(len(item) for item in linhas) < 4_000
    return bruto


def _csv_campos_de_cem() -> bytes:
    """Caso B: 1.110.015 bytes e campos de dados com 100 caracteres."""
    alvo = 1_110_015
    cabecalho = b"categoria,medida_textual\n"
    assert len(cabecalho) == 25
    tamanho = 100
    largura = tamanho + 1 + tamanho + 1
    quantidade = (alvo - len(cabecalho)) // largura
    assert len(cabecalho) + quantidade * largura == alvo
    linhas = []
    for indice in range(quantidade):
        esquerda = f"{indice:05d}".ljust(tamanho, "a")
        direita = f"{indice:05d}".ljust(tamanho, "b")
        linhas.append(f"{esquerda},{direita}\n")
    bruto = cabecalho + "".join(linhas).encode("ascii")
    assert len(bruto) == alvo
    return bruto


def _canonico_bytes(valor) -> int:
    return len(json.dumps(valor, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8"))


def _payload_segmentado(escopo_dados):
    corpo = {
        "chave": "atomo-lista",
        "origem": "sintetica",
        "origem_ref": "sintetico-atomo",
        "versao": 1,
        "estrutura": {"colunas": ["i"]},
        "perfil": {"linhas": len(escopo_dados)},
        "dados_tipados": escopo_dados,
        "metadados_preparacao": {"versao_preparacao": "preparar.v1"},
    }
    corpo["impressao"] = calcular_impressao_dataset(
        chave=corpo["chave"],
        origem=corpo["origem"],
        origem_ref=corpo["origem_ref"],
        versao=1,
        estrutura=corpo["estrutura"],
        perfil=corpo["perfil"],
        dados_tipados=escopo_dados,
        metadados_preparacao=corpo["metadados_preparacao"],
    )
    return corpo


def test_preambulo_numerico_permanece_fora_do_registro_analitico(ctx):
    _conta, _franquia, usuario = _par("ing-pre", "ing.pre@test.com")
    grade = [[1, 2], ["categoria", "medida"], ["a", 3]]
    resultado = _ingerir(usuario, _xlsx({"movimento": grade}), "preambulo.xlsx", chave="lote-pre")
    regiao = _regiao(resultado)
    assert regiao["cabecalho_deteccao"] == "inequivoca"
    assert regiao["cabecalho_linha"] == 2
    assert [item["linha"] for item in _dados(resultado)] == [3]
    assert [_valor(celula) for celula in _dados(resultado)[0]["valores"]] == ["a", 3]
    preambulo = [item for item in regiao["celulas_preservadas"] if item["papel"] == "preambulo"]
    assert [(item["linha"], item["coluna"], item["celula"]) for item in preambulo] == [
        (1, 1, {"tipo": "numero", "valor": 1}),
        (1, 2, {"tipo": "numero", "valor": 2}),
    ]
    cabecalho = [item for item in regiao["celulas_preservadas"] if item["papel"] == "cabecalho"]
    assert [item["celula"]["valor"] for item in cabecalho] == ["categoria", "medida"]
    assert {item["linha"] for item in cabecalho} == {2}
    observados = [1, 2, "categoria", "medida", "a", 3]
    achados = [item["celula"].get("valor") for item in regiao["celulas_preservadas"]]
    achados.extend(
        celula.get("valor")
        for registro in _dados(resultado)
        for celula in registro["valores"]
        if celula["tipo"] != "ausente"
    )
    for valor in observados:
        assert achados.count(valor) >= 1
    assert resultado["preparacao"]["metadados_preparacao"]["lacunas"] == "celula_ausente_nao_preenchida"
    corpo = obter_dataset(escopo_do_usuario(usuario), resultado["dataset"]["id"])
    assert corpo["estrutura"]["regioes"][0]["celulas_preservadas"] == regiao["celulas_preservadas"]
    assert [celula.get("valor") for celula in corpo["dados_tipados"][0]["valores"]] == ["a", 3]


def test_lacuna_observada_nao_recebe_valor_inventado(ctx):
    _conta, _franquia, usuario = _par("ing-lacuna", "ing.lacuna@test.com")
    resultado = _ingerir(usuario, _xlsx({"movimento": [[1, None, 2]]}), "lacuna.xlsx", chave="lote-lacuna")
    valores = _dados(resultado)[0]["valores"]
    assert valores[0]["valor"] == 1
    assert valores[1] == {"tipo": "ausente"}
    assert "valor" not in valores[1]
    assert valores[2]["valor"] == 2


def test_numeros_equivalentes_ficam_pendentes_sem_depender_da_ordem(ctx):
    _conta, _franquia, usuario = _par("ing-ponto", "ing.ponto@test.com")
    literais = {"1.234", "2.345", "3.456"}
    direto = _ingerir(
        usuario,
        "categoria,medida\nalfa,1.234\nbeta,2.345\ngama,3.456\n".encode("utf-8"),
        "direto.csv",
        chave="lote-ponto",
    )
    inverso = _ingerir(
        usuario,
        "categoria,medida\ngama,3.456\nbeta,2.345\nalfa,1.234\n".encode("utf-8"),
        "inverso.csv",
        chave="lote-ponto-inv",
    )
    ida = _literais_numericos(direto, literais)
    volta = _literais_numericos(inverso, literais)
    assert ida == [("1.234", "numero_pendente"), ("2.345", "numero_pendente"), ("3.456", "numero_pendente")]
    assert volta == [("3.456", "numero_pendente"), ("2.345", "numero_pendente"), ("1.234", "numero_pendente")]
    assert all(isinstance(valor, str) for valor, _tipo in ida + volta)
    assert any(item["codigo"] == "numero_ambiguo" for item in direto["preparacao"]["estrutura"]["ambiguidades"])


def test_linhas_elegiveis_contam_registros_e_nao_celulas(ctx):
    _conta, _franquia, usuario = _par("ing-largura", "ing.largura@test.com")
    escopo = escopo_do_usuario(usuario)
    resultado = ingerir(
        escopo,
        [
            _arquivo(b"a,b\n1,2\n3,4\n", "estreito.csv"),
            _arquivo(b"c,d,e,f\n5,6,7,8\n9,10,11,12\n", "largo.csv"),
        ],
        chave="lote-largura",
    )
    assert resultado["preparacao"]["perfil"]["linhas_elegiveis"] == 4
    assert [len(item["valores"]) for item in _dados(resultado)] == [2, 2, 4, 4]
    irregular = _ingerir(
        usuario,
        b"a,b,c\n1,2,3\n4,,6\n7,8,\n10,11,12\n",
        "irregular.csv",
        chave="lote-irregular",
    )
    assert irregular["preparacao"]["perfil"]["linhas_elegiveis"] == 4
    assert irregular["preparacao"]["dados_tipados"][1]["valores"][1] == {"tipo": "ausente"}


def test_codificacao_invalida_e_numero_enorme_respondem_422(app, monkeypatch):
    client = _cliente(app, monkeypatch)
    with app.app_context():
        _conta, _franquia, usuario = _par("ing-http", "ing.http@test.com")
        usuario_id = usuario.id
    _entrar(client, usuario_id)
    codificacao = _postar(client, b"categoria,medida\n1,\x81\n", "ruim.csv", chave="lote-codec")
    assert codificacao.status_code == 422
    assert codificacao.get_json() == {"ok": False, "codigo": "codificacao_nao_suportada"}
    assert b"Traceback" not in codificacao.data
    assert b"\x81" not in codificacao.data
    digitos = b"9" * 5000
    enorme = _postar(client, b"categoria,medida\n1," + digitos + b"\n", "enorme.csv", chave="lote-digitos")
    assert enorme.status_code == 422
    assert enorme.get_json() == {"ok": False, "codigo": "limite_processamento"}
    assert b"Traceback" not in enorme.data
    assert b"9" * 80 not in enorme.data
    with app.app_context():
        assert AgenteBiDataset.query.count() == 0


def test_timeout_antes_da_confirmacao_nao_persiste(ctx):
    _conta, _franquia, usuario = _par("ing-prazo", "ing.prazo@test.com")

    class Relogio:
        def __init__(self):
            self.chamadas = 0
            self.limite = 10**9

        def __call__(self):
            self.chamadas += 1
            if self.chamadas >= self.limite:
                return 10**6
            return 0

    relogio = Relogio()
    primeiro = _ingerir(
        usuario,
        b"origem,destino\nnorte,sul\n",
        "a.csv",
        chave="lote-prazo",
        relogio=relogio,
    )
    assert primeiro["persistido"] is True
    seguinte = Relogio()
    seguinte.limite = relogio.chamadas
    with pytest.raises(IngestaoErro) as exc:
        _ingerir(
            usuario,
            b"origem,destino\nnorte,sul\n",
            "a.csv",
            chave="lote-prazo-2",
            relogio=seguinte,
        )
    assert exc.value.codigo == "limite_tempo"
    assert AgenteBiDataset.query.count() == 1
    assert AgenteBiDatasetBloco.query.count() == 0


def test_atomicidade_quando_a_confirmacao_falha(ctx, monkeypatch):
    _conta, _franquia, usuario = _par("ing-atomo", "ing.atomo@test.com")
    dados = [{"i": indice} for indice in range(MAX_LISTA + 1)]
    payload = _payload_segmentado(dados)

    def falha():
        raise RuntimeError("falha intermediaria")

    monkeypatch.setattr(db.session, "commit", falha)
    with pytest.raises(RuntimeError, match="falha intermediaria"):
        criar_dataset(escopo_do_usuario(usuario), payload)
    assert AgenteBiDataset.query.count() == 0
    assert AgenteBiDatasetBloco.query.count() == 0


def test_governanca_bloqueia_dataset_segmentado_sem_gravar_parte(ctx):
    _conta, _franquia, usuario = _par("ing-gov-bloco", "ing.gov.bloco@test.com")
    dados = [{"i": indice} for indice in range(MAX_LISTA + 1)]
    dados[-1] = {"i": "123.456.789-09"}
    payload = _payload_segmentado(dados)
    with pytest.raises(AgenteBiErro) as exc:
        criar_dataset(escopo_do_usuario(usuario), payload)
    assert exc.value.codigo == "conteudo_nao_permitido"
    assert AgenteBiDataset.query.count() == 0
    assert AgenteBiDatasetBloco.query.count() == 0


def test_csv_de_um_megabyte_persiste_recupera_e_versiona(ctx):
    assert MAX_JSON_BYTES == 1_500_000
    assert MAX_LISTA == 20_000
    conta, franquia, usuario = _par("ing-mega", "ing.mega@test.com")
    reserva = franquia.reserva_pendente
    consumo = franquia.consumo_acumulado
    bruto = _csv_codex()
    resultado = _ingerir(usuario, bruto, "sintetico.csv", chave="lote-mega")
    dados = resultado["preparacao"]["dados_tipados"]
    assert resultado["persistido"] is True, resultado["motivo"]
    assert resultado["dataset"]["versao"] == 1
    assert len(dados) == 10_000
    assert resultado["preparacao"]["perfil"]["linhas_elegiveis"] == 10_000
    canonico = json.dumps(dados, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    assert len(canonico) > MAX_JSON_BYTES
    escopo = escopo_do_usuario(usuario)
    corpo = obter_dataset(escopo, resultado["dataset"]["id"])
    assert corpo["dados_tipados"] == dados
    assert corpo["dados_tipados"][0]["valores"][0]["valor"].startswith("cat-")
    assert corpo["dados_tipados"][-1]["linha"] == 10_001
    impressao = calcular_impressao_dataset(
        chave=corpo["chave"],
        origem=corpo["origem"],
        origem_ref=corpo["origem_ref"],
        versao=corpo["versao"],
        estrutura=corpo["estrutura"],
        perfil=corpo["perfil"],
        dados_tipados=corpo["dados_tipados"],
        metadados_preparacao=corpo["metadados_preparacao"],
    )
    assert impressao == corpo["impressao"]
    gravado = db.session.get(AgenteBiDataset, resultado["dataset"]["id"])
    manifesto = json.loads(gravado.dados_tipados_json)
    assert manifesto["contrato"] == "dados.segmentado.v1"
    assert manifesto["registros"] == 10_000
    blocos = (
        AgenteBiDatasetBloco.query.filter_by(dataset_id=gravado.id, usuario_id=usuario.id)
        .order_by(AgenteBiDatasetBloco.indice.asc())
        .all()
    )
    assert len(blocos) == len(manifesto["blocos"]) >= 2
    assert all(len(bloco.conteudo_json.encode("utf-8")) <= MAX_JSON_BYTES for bloco in blocos)
    assert all(bloco.registros <= MAX_LISTA for bloco in blocos)
    repetido = _ingerir(usuario, bruto, "sintetico.csv", chave="lote-mega")
    assert repetido["reutilizado"] is True
    assert repetido["dataset"]["id"] == gravado.id
    assert AgenteBiDatasetBloco.query.filter_by(dataset_id=gravado.id).count() == len(blocos)
    segundo = _ingerir(usuario, b"origem,destino\nleste,oeste\n", "outro.csv", chave="lote-mega")
    assert segundo["dataset"]["versao"] == 2
    assert segundo["dataset"]["versao_anterior_id"] == gravado.id
    assert obter_dataset(escopo, gravado.id)["dados_tipados"] == dados
    colega = seed_usuario(franquia.id, conta.id, email="ing.mega.colega@test.com")
    with pytest.raises(AgenteBiErro) as alheio:
        obter_dataset(escopo_do_usuario(colega), gravado.id)
    assert alheio.value.codigo == "nao_encontrado"
    assert AgenteBiDatasetBloco.query.filter_by(usuario_id=colega.id).count() == 0
    bloco = blocos[0]
    bloco.conteudo_json = "[]"
    with pytest.raises(ValueError, match="imutavel"):
        db.session.commit()
    db.session.rollback()
    db.session.refresh(franquia)
    assert franquia.reserva_pendente == reserva
    assert franquia.consumo_acumulado == consumo
    assert IaChamadaTentativa.query.count() == 0


def test_vinte_mil_e_um_registros_persistem_no_mesmo_dataset(ctx):
    _conta, _franquia, usuario = _par("ing-vinte", "ing.vinte@test.com")
    linhas = ["id,valor"]
    linhas.extend(f"r{indice},1" for indice in range(20_001))
    bruto = ("\n".join(linhas) + "\n").encode("ascii")
    assert len(bruto) < 2 * 1024 * 1024
    resultado = _ingerir(usuario, bruto, "vinte.csv", chave="lote-vinte")
    assert resultado["persistido"] is True
    assert resultado["preparacao"]["perfil"]["linhas_elegiveis"] == 20_001
    assert len(resultado["preparacao"]["dados_tipados"]) == 20_001
    corpo = obter_dataset(escopo_do_usuario(usuario), resultado["dataset"]["id"])
    assert corpo["dados_tipados"] == resultado["preparacao"]["dados_tipados"]
    assert corpo["dados_tipados"][0]["valores"][0]["valor"] == "r0"
    assert corpo["dados_tipados"][-1]["valores"][0]["valor"] == "r20000"
    assert AgenteBiDataset.query.count() == 1
    assert AgenteBiDatasetBloco.query.count() >= 2


def test_duzentas_mil_celulas_persistem_dentro_do_teto(ctx, monkeypatch):
    monkeypatch.setenv("AGENTE_BI_INGESTAO_MAX_SEGUNDOS", "30")
    _conta, _franquia, usuario = _par("ing-celulas", "ing.celulas@test.com")
    bruto = b"1,2\n" * 100_000
    assert len(bruto) < 2 * 1024 * 1024
    resultado = _ingerir(usuario, bruto, "celulas.csv", chave="lote-celulas")
    assert resultado["persistido"] is True
    assert resultado["preparacao"]["perfil"]["celulas"] == 200_000
    assert resultado["preparacao"]["perfil"]["linhas_elegiveis"] == 100_000
    corpo = obter_dataset(escopo_do_usuario(usuario), resultado["dataset"]["id"])
    assert len(corpo["dados_tipados"]) == 100_000
    assert corpo["dados_tipados"][0]["valores"][0]["valor"] == 1
    assert corpo["dados_tipados"][-1]["valores"][1]["valor"] == 2
    limites = corpo["metadados_preparacao"]["limites"]
    assert limites["celulas"] == 200_000
    assert limites["json_bytes"] == MAX_JSON_BYTES
    assert limites["lista"] == MAX_LISTA
    assert all(
        len(bloco.conteudo_json.encode("utf-8")) <= MAX_JSON_BYTES
        for bloco in AgenteBiDatasetBloco.query.all()
    )


def test_csv_textual_nao_materializa_milhares_de_candidatos(ctx):
    conta, franquia, usuario = _par("ing-texto", "ing.texto@test.com")
    reserva = franquia.reserva_pendente
    consumo = franquia.consumo_acumulado
    bruto = _csv_textual_um_mega()
    resultado = _ingerir(usuario, bruto, "massa.csv", chave="lote-textual")
    _afirmar_massa_textual(
        resultado,
        usuario,
        bruto,
        chave="lote-textual",
        linhas_arquivo=10_001,
        conta=conta,
        franquia=franquia,
        reserva=reserva,
        consumo=consumo,
    )
    primeiro = resultado["preparacao"]["dados_tipados"][1]["valores"]
    assert primeiro[0]["valor"].startswith("cat-")
    assert primeiro[1]["valor"].startswith("med-")
    assert {item["valores"][0]["valor"] for item in resultado["preparacao"]["dados_tipados"][1:]} == {
        f"cat-{indice:05d}" for indice in range(10_000)
    }


def test_csv_de_campos_longos_permanece_amostra_e_persiste(ctx):
    conta, franquia, usuario = _par("ing-cem", "ing.cem@test.com")
    reserva = franquia.reserva_pendente
    consumo = franquia.consumo_acumulado
    bruto = _csv_campos_de_cem()
    resultado = _ingerir(usuario, bruto, "massa.csv", chave="lote-cem")
    _afirmar_massa_textual(
        resultado,
        usuario,
        bruto,
        chave="lote-cem",
        linhas_arquivo=5_496,
        conta=conta,
        franquia=franquia,
        reserva=reserva,
        consumo=consumo,
    )
    campo = resultado["preparacao"]["dados_tipados"][1]["valores"][0]["valor"]
    assert len(campo) == 100
    assert campo.startswith("00000")


def _afirmar_massa_textual(
    resultado,
    usuario,
    bruto,
    *,
    chave,
    linhas_arquivo,
    conta,
    franquia,
    reserva,
    consumo,
):
    dados = resultado["preparacao"]["dados_tipados"]
    regiao = _regiao(resultado)
    assert resultado["persistido"] is True, resultado["motivo"]
    assert resultado["dataset"]["versao"] == 1
    assert regiao["cabecalho_deteccao"] == "ambigua"
    assert regiao["cabecalho_linha"] is None
    assert regiao["candidatos_amostra"] is True
    assert regiao["candidatos_total"] == linhas_arquivo
    assert len(regiao["candidatos"]) == LIMITE_CANDIDATOS_PUBLICADOS
    linhas_amostra = [item["linha"] for item in regiao["candidatos"]]
    assert linhas_amostra[0] == 1
    assert linhas_arquivo in linhas_amostra
    assert _canonico_bytes(resultado["preparacao"]["estrutura"]) <= MAX_JSON_BYTES
    assert len(dados) == linhas_arquivo
    assert resultado["preparacao"]["perfil"]["linhas_elegiveis"] == linhas_arquivo
    assert regiao["celulas_preservadas"] == []
    assert all(len(item["valores"]) >= 1 for item in dados)
    escopo = escopo_do_usuario(usuario)
    corpo = obter_dataset(escopo, resultado["dataset"]["id"])
    assert corpo["dados_tipados"] == dados
    assert corpo["estrutura"]["regioes"][0]["candidatos_total"] == linhas_arquivo
    assert len(corpo["estrutura"]["regioes"][0]["candidatos"]) == LIMITE_CANDIDATOS_PUBLICADOS
    gravado = db.session.get(AgenteBiDataset, resultado["dataset"]["id"])
    repetido = _ingerir(usuario, bruto, "massa.csv", chave=chave)
    assert repetido["reutilizado"] is True
    assert repetido["dataset"]["id"] == gravado.id
    segundo = _ingerir(usuario, b"origem,destino\nleste,oeste\n", "outro.csv", chave=chave)
    assert segundo["dataset"]["versao"] == 2
    assert segundo["dataset"]["versao_anterior_id"] == gravado.id
    assert obter_dataset(escopo, gravado.id)["dados_tipados"] == dados
    colega = seed_usuario(franquia.id, conta.id, email=f"{chave}@colega.test")
    with pytest.raises(AgenteBiErro) as alheio:
        obter_dataset(escopo_do_usuario(colega), gravado.id)
    assert alheio.value.codigo == "nao_encontrado"
    db.session.refresh(franquia)
    assert franquia.reserva_pendente == reserva
    assert franquia.consumo_acumulado == consumo
    assert IaChamadaTentativa.query.count() == 0


def test_dataset_simples_nao_e_lido_como_manifesto(ctx):
    _conta, _franquia, usuario = _par("ing-inline", "ing.inline@test.com")
    resultado = _ingerir(usuario, b"origem,destino\nnorte,10\n", "inline.csv", chave="lote-inline")
    assert resultado["persistido"] is True
    assert resultado["preparacao"]["estrutura"]["regioes"][0]["cabecalho_deteccao"] == "inequivoca"
    assert AgenteBiDatasetBloco.query.count() == 0
    corpo = obter_dataset(escopo_do_usuario(usuario), resultado["dataset"]["id"])
    assert isinstance(corpo["dados_tipados"], list)
    assert corpo["dados_tipados"] == resultado["preparacao"]["dados_tipados"]
    assert corpo["dados_tipados"][0]["valores"][0]["valor"] == "norte"
    gravado = db.session.get(AgenteBiDataset, resultado["dataset"]["id"])
    assert json.loads(gravado.dados_tipados_json) == corpo["dados_tipados"]


def test_bloco_ausente_nao_devolve_o_manifesto(ctx):
    _conta, _franquia, usuario = _par("ing-buraco", "ing.buraco@test.com")
    dados = [{"n": indice, "t": "a" * 1000} for indice in range(4_500)]
    escopo = escopo_do_usuario(usuario)
    criado = criar_dataset(escopo, _payload_segmentado(dados))
    assert obter_dataset(escopo, criado["id"])["dados_tipados"] == dados
    blocos = (
        AgenteBiDatasetBloco.query.filter_by(dataset_id=criado["id"])
        .order_by(AgenteBiDatasetBloco.indice.asc())
        .all()
    )
    assert len(blocos) >= 3
    indices = [bloco.indice for bloco in blocos]
    gravado = db.session.get(AgenteBiDataset, criado["id"])
    assert json.loads(gravado.dados_tipados_json)["contrato"] == "dados.segmentado.v1"

    def falha() -> None:
        db.session.expire_all()
        with pytest.raises(AgenteBiErro) as exc:
            obtido = obter_dataset(escopo, criado["id"])
            raise AssertionError(obtido["dados_tipados"])
        assert exc.value.codigo == "integridade_dados"
        assert str(exc.value) == "integridade_dados"

    def retirar(indice: int) -> None:
        alvo = AgenteBiDatasetBloco.query.filter_by(dataset_id=criado["id"], indice=indice).one()
        copia = AgenteBiDatasetBloco(
            dataset_id=alvo.dataset_id,
            conta_id=alvo.conta_id,
            franquia_id=alvo.franquia_id,
            usuario_id=alvo.usuario_id,
            indice=alvo.indice,
            registros=alvo.registros,
            sha256=alvo.sha256,
            conteudo_json=alvo.conteudo_json,
            created_at=alvo.created_at,
        )
        db.session.delete(alvo)
        db.session.commit()
        falha()
        db.session.add(copia)
        db.session.commit()
        assert obter_dataset(escopo, criado["id"])["dados_tipados"] == dados

    retirar(indices[0])
    retirar(indices[len(indices) // 2])
    retirar(indices[-1])
    AgenteBiDatasetBloco.query.filter_by(dataset_id=criado["id"]).delete()
    db.session.commit()
    falha()
    assert json.loads(db.session.get(AgenteBiDataset, criado["id"]).dados_tipados_json)["armazenamento"] == "segmentado"


def test_recomposicao_recusa_duplicata_e_ordem_trocada():
    from types import SimpleNamespace

    linha = SimpleNamespace(conta_id=1, franquia_id=1, usuario_id=1)
    manifesto = {
        "armazenamento": "segmentado",
        "contrato": "dados.segmentado.v1",
        "registros": 2,
        "sha256": "ab" * 32,
        "blocos": [
            {"indice": 0, "sha256": "aa"},
            {"indice": 1, "sha256": "bb"},
        ],
    }

    def bloco(indice, ident):
        return SimpleNamespace(
            id=ident,
            indice=indice,
            sha256="aa",
            conteudo_json="[]",
            registros=0,
            conta_id=1,
            franquia_id=1,
            usuario_id=1,
        )

    with pytest.raises(AgenteBiErro) as duplicado:
        _recompor_segmentado(manifesto, [bloco(0, 1), bloco(0, 2)], linha)
    assert duplicado.value.codigo == "integridade_dados"
    trocado = {
        **manifesto,
        "blocos": [
            {"indice": 1, "sha256": "bb"},
            {"indice": 0, "sha256": "aa"},
        ],
    }
    with pytest.raises(AgenteBiErro) as fora:
        _recompor_segmentado(trocado, [bloco(1, 2), bloco(0, 1)], linha)
    assert fora.value.codigo == "integridade_dados"
    with pytest.raises(AgenteBiErro) as vazio:
        _recompor_segmentado(manifesto, [], linha)
    assert vazio.value.codigo == "integridade_dados"


def test_migration_dos_blocos_casa_com_o_modelo_e_reverte(ctx):
    import importlib.util

    from sqlalchemy import create_engine, inspect, text

    from tests.test_agente_bi_fundacao import _run

    def carregar(nome):
        caminho = ROOT / "migrations" / "versions" / nome
        spec = importlib.util.spec_from_file_location(nome, caminho)
        modulo = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        spec.loader.exec_module(modulo)
        return modulo

    base = carregar("j8k9l0m1n2o3_agente_bi_persistencia_inicial.py")
    nova = carregar("k9l0m1n2o3p4_agente_bi_dados_segmentados.py")
    assert nova.down_revision == "j8k9l0m1n2o3"
    engine = create_engine("sqlite://")
    with engine.begin() as conn:
        conn.execute(text("CREATE TABLE conta (id INTEGER PRIMARY KEY)"))
        conn.execute(text("CREATE TABLE franquia (id INTEGER PRIMARY KEY)"))
        conn.execute(text('CREATE TABLE "user" (id INTEGER PRIMARY KEY)'))
        conn.execute(text("CREATE TABLE ia_chamada_tentativa (id INTEGER PRIMARY KEY)"))
    _run(engine, base, base.upgrade)
    _run(engine, nova, nova.upgrade)
    colunas = {coluna["name"] for coluna in inspect(engine).get_columns("agente_bi_dataset_bloco")}
    assert colunas == {coluna.name for coluna in AgenteBiDatasetBloco.__table__.columns}
    with engine.connect() as conn:
        gatilhos = {row[0] for row in conn.execute(text("SELECT name FROM sqlite_master WHERE type='trigger'"))}
    assert "trg_agente_bi_dataset_bloco_imutavel" in gatilhos
    _run(engine, nova, nova.downgrade)
    assert "agente_bi_dataset_bloco" not in inspect(engine).get_table_names()
    assert "agente_bi_dataset" in inspect(engine).get_table_names()


def test_postgres_descartavel_cria_e_remove_os_blocos():
    import importlib.util
    import os

    from sqlalchemy import create_engine, inspect, text

    from tests.test_agente_bi_fundacao import _run

    uri = os.environ.get("AGENTE_BI_POSTGRES_DESCARTAVEL", "").strip()
    if "scrum252_descartavel" not in uri or any(
        marca in uri for marca in ("homolog", "producao", "production")
    ):
        pytest.skip("PostgreSQL descartavel nao provisionado nesta execucao")

    def carregar(nome):
        caminho = ROOT / "migrations" / "versions" / nome
        spec = importlib.util.spec_from_file_location(nome + "_pg", caminho)
        modulo = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        spec.loader.exec_module(modulo)
        return modulo

    engine = create_engine(uri)
    with engine.begin() as conn:
        conn.execute(text("CREATE TABLE IF NOT EXISTS conta (id INTEGER PRIMARY KEY)"))
        conn.execute(text("CREATE TABLE IF NOT EXISTS franquia (id INTEGER PRIMARY KEY)"))
        conn.execute(text('CREATE TABLE IF NOT EXISTS "user" (id INTEGER PRIMARY KEY)'))
        conn.execute(
            text(
                "CREATE TABLE IF NOT EXISTS ia_chamada_tentativa ("
                "id INTEGER PRIMARY KEY, usuario_id INTEGER, conta_id INTEGER, franquia_id INTEGER)"
            )
        )
    base = carregar("j8k9l0m1n2o3_agente_bi_persistencia_inicial.py")
    nova = carregar("k9l0m1n2o3p4_agente_bi_dados_segmentados.py")
    _run(engine, base, base.upgrade, batch=False)
    _run(engine, nova, nova.upgrade, batch=False)
    with engine.connect() as conn:
        gatilhos = {row[0] for row in conn.execute(text("SELECT tgname FROM pg_trigger WHERE NOT tgisinternal"))}
    assert "trg_agente_bi_dataset_bloco_imutavel" in gatilhos
    assert "agente_bi_dataset_bloco" in inspect(engine).get_table_names()
    _run(engine, nova, nova.downgrade, batch=False)
    _run(engine, base, base.downgrade, batch=False)
    assert "agente_bi_dataset_bloco" not in inspect(engine).get_table_names()
    assert "agente_bi_dataset" not in inspect(engine).get_table_names()
    engine.dispose()

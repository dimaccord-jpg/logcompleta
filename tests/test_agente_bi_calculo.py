"""SCRUM-261: executor determinístico, gabaritos numéricos e memória de cálculo."""
from __future__ import annotations

import copy
import json
import logging
from datetime import date
from decimal import ROUND_HALF_EVEN, Context, Decimal
from pathlib import Path

import pytest
from sqlalchemy import func

from app.extensions import db
from app.models import AgenteBiAnalise, IaChamadaTentativa
from app.services.agente_bi_calculo_service import (
    CONTRATO_CALCULO,
    CONTRATO_CONTEXTO_CALCULO,
    CalculoErro,
    calcular_definicao,
)
from app.services.agente_bi_definir_service import definir, definir_selecao, provedor_simulado
from app.services.agente_bi_perfil_semantico_service import perfilar_dataset
from app.services.agente_bi_persistencia_service import criar_analise, escopo_do_usuario, obter_analise
from tests.conftest import seed_usuario
from tests.test_agente_bi_analisar import SENTINELA, _movimento
from tests.test_agente_bi_definir import _jornada
from tests.test_agente_bi_ingestao import _cliente as _cliente_ingestao
from tests.test_agente_bi_ingestao import _entrar, _ingerir, _par, _xlsx

PRECISAO = Context(prec=28, rounding=ROUND_HALF_EVEN)
ROOT = Path(__file__).resolve().parents[1]


def _texto(valor: Decimal) -> str:
    texto = format(valor, "f")
    if "." in texto:
        texto = texto.rstrip("0").rstrip(".")
    return texto or "0"


def _dividir(numerador, denominador) -> str:
    return _texto(PRECISAO.divide(Decimal(numerador), Decimal(denominador)))


def _coluna(perfil, **filtro):
    return next(coluna for coluna in perfil["colunas"] if all(coluna[chave] == valor for chave, valor in filtro.items()))


def _ref(coluna) -> dict:
    return {
        "arquivo": coluna["arquivo"],
        "aba": coluna["aba"],
        "regiao": coluna["regiao"],
        "indice": coluna["indice"],
    }


def _gravar(escopo, selecao_id, documento):
    origem = obter_analise(escopo, selecao_id)
    atual = (
        db.session.query(func.max(AgenteBiAnalise.rodada))
        .filter_by(conta_id=escopo.conta_id, franquia_id=escopo.franquia_id, usuario_id=escopo.usuario_id, chave=origem["chave"])
        .scalar()
    )
    return criar_analise(
        escopo,
        {
            "chave": origem["chave"],
            "rodada": 1 if atual is None else int(atual) + 1,
            "dataset_id": origem["dataset_id"],
            "versao": origem["versao"],
            "impressao": origem["impressao"],
            "etapa": "definicao",
            "estado": "concluida",
            "contexto": {"contrato": "contexto.definir.v1", "definicao": True},
            "regras_negocio": {"dsl": "dsl.analitica.v1", "execucao": "nao_executada"},
            "esclarecimentos": {"bloqueios": []},
            "contrato_versao": "definir.v1",
            "resultado": documento,
        },
    )


def _calcular(escopo, analise_id, **extras):
    corpo = {"analise_id": analise_id}
    corpo.update(extras)
    return calcular_definicao(escopo, corpo)


def _por_operacao(memoria, operacao, grafico="indicador"):
    return next(
        item
        for item in memoria["indicadores"]
        if item["operacao"] == operacao and item["grafico"] == grafico and item["status"] == "calculado"
    )


def _indicador(pacote, operacao, *, leitura=None, agrupamento=None, segunda=None, grafico="indicador"):
    copia = copy.deepcopy(pacote)
    candidatos = [
        grafico_item
        for grafico_item in copia["graficos"]
        if grafico_item["grafico"] == "indicador" and grafico_item["operacao"] == "soma"
    ]
    item = candidatos[0] if candidatos else next(grafico_item for grafico_item in copia["graficos"] if grafico_item["grafico"] == "indicador")
    item["operacao"] = operacao
    item["grafico"] = grafico
    item["leitura"] = leitura or item["leitura"]
    item["agrupamento"] = agrupamento
    item["segunda_medida"] = segunda
    copia["graficos"] = [item]
    return copia


def _filtrar(coluna, operador, nome, tipo, nome_fim=None):
    def provedor(contexto):
        base = provedor_simulado(contexto)
        item = base["definicoes"][0]
        item["filtros"] = [
            {
                "coluna": _ref(coluna) if "indice" in coluna else coluna,
                "operador": operador,
                "parametro": nome,
                "parametro_fim": nome_fim,
            }
        ]
        item["parametros"] = [{"nome": nome, "tipo": tipo, "obrigatorio": True}]
        if nome_fim:
            item["parametros"].append({"nome": nome_fim, "tipo": tipo, "obrigatorio": True})
        item["status"] = "definida"
        item["esclarecimento"] = None
        return base

    return provedor


def _planilha_partes():
    return _xlsx(
        {
            "movimento": [
                ["beta", "parte", "todo"],
                ["norte", 2, 4],
                ["norte", 2, 4],
                ["sul", 1, 2],
            ]
        }
    )


def _jornada_de(email, slug, conteudo):
    _conta, franquia, usuario = _par(slug, email)
    escopo = escopo_do_usuario(usuario)
    ingerido = _ingerir(usuario, conteudo, "movimento.xlsx", chave=f"lote-{slug}")
    perfil = perfilar_dataset(escopo, ingerido["dataset"]["id"])
    from app.services.agente_bi_analisar_service import analisar_dataset
    from app.services.agente_bi_descoberta_service import aprovar_selecao

    gravado = analisar_dataset(
        escopo,
        ingerido["dataset"]["id"],
        chave=ingerido["dataset"]["chave"],
        rodada=1,
        pedido={"modalidades": ["A"]},
    )
    ids = [
        item["id"]
        for item in gravado["resultado"]["sugestoes"]
        if item["status"] == "viavel" and item["grafico"] and item["juncao"] == "nao_executada"
    ][:6]
    assert ids
    aprovada = aprovar_selecao(
        escopo,
        {
            "analise_id": gravado["analise"]["id"],
            "sugestoes": ids,
            "perguntas": [{"id": item["id"], "confirmada": True} for item in gravado["resultado"]["perguntas"]],
            "confirmacao": True,
        },
    )
    pacote = dict(aprovada["selecao"]["resultado"])
    pacote["id"] = aprovada["selecao"]["id"]
    return usuario, franquia, escopo, perfil, pacote


def test_soma_contagem_media_e_distinta_batem_com_o_gabarito(ctx):
    _usuario, _franquia, escopo, perfil, pacote, _aprovada = _jornada("calculo.soma@test.com", "calculo-soma")
    gravada = definir_selecao(escopo, {"analise_id": pacote["id"]})
    memoria = _calcular(escopo, gravada["definicao"]["id"])["calculo"]["resultado"]
    assert memoria["ia_calculo"] is False
    assert memoria["chamada_anthropic"] is False
    assert memoria["executor"] == "executor.deterministico.v1"
    assert memoria["dataset_versao"] == perfil["dataset"]["versao"]
    assert memoria["dataset_impressao"] == perfil["dataset"]["impressao"]
    soma = _por_operacao(memoria, "soma")
    assert soma["resultado"] == "25.25"
    assert soma["elegiveis"] == 3
    assert soma["unidade"] == "valor"
    assert soma["juncao"] == "nao_aplicada"
    assert SENTINELA not in json.dumps(soma)

    valor = _coluna(perfil, dimensao_principal="valor", aba=0)
    alfa = _coluna(perfil, indice=0, aba=0)
    for operacao, leitura, esperado, coluna in (
        ("contagem", "contagem_de_registros", "3", valor),
        ("media", "indicador_de_medida", _dividir("25.25", "3"), valor),
        ("contagem_distinta", "contagem_de_registros", "2", alfa),
    ):
        pacote_op = _indicador(pacote, operacao, leitura=leitura)
        if operacao == "contagem_distinta":
            pacote_op["graficos"][0]["medida"] = _ref(alfa)
        documento = definir(perfil, pacote_op)
        salva = _gravar(escopo, pacote["id"], documento)
        indicador = _calcular(escopo, salva["id"])["calculo"]["resultado"]["indicadores"][0]
        assert indicador["resultado"] == esperado
        assert indicador["dataset_id"] == perfil["dataset"]["id"]
        assert SENTINELA not in json.dumps(indicador)


def test_count_if_razao_e_participacao(ctx):
    _usuario, _franquia, escopo, perfil, pacote = _jornada_de("calculo.partes@test.com", "calculo-partes", _planilha_partes())
    beta = _coluna(perfil, indice=0, aba=0)
    parte = _coluna(perfil, indice=1, aba=0)
    todo = _coluna(perfil, indice=2, aba=0)
    assert parte["dimensao_principal"] == todo["dimensao_principal"] == "quantidade"

    condicional = _indicador(pacote, "count_if", leitura="contagem_de_registros")
    condicional["graficos"][0]["medida"] = _ref(parte)
    documento = definir(perfil, condicional, provedor=_filtrar(beta, "igual", "faixa", "texto"))
    salva = _gravar(escopo, pacote["id"], documento)
    indicador = _calcular(escopo, salva["id"], parametros={"faixa": "norte"})["calculo"]["resultado"]["indicadores"][0]
    assert indicador["resultado"] == "2"
    assert indicador["elegiveis"] == 2
    assert any(item["motivo"] == "filtro" and item["quantidade"] == 1 for item in indicador["excluidos"])
    assert "norte" not in json.dumps(indicador["filtros"])

    razao = _indicador(
        pacote,
        "razao",
        leitura="indicador_de_medida",
        segunda=_ref(todo),
    )
    razao["graficos"][0]["medida"] = _ref(parte)
    documento = definir(perfil, razao)
    salva = _gravar(escopo, pacote["id"], documento)
    indicador = _calcular(escopo, salva["id"])["calculo"]["resultado"]["indicadores"][0]
    assert indicador["resultado"] == "0.5"
    assert indicador["unidade"] == "adimensional"
    assert indicador["elegiveis"] == 3

    participacao = _indicador(
        pacote,
        "participacao",
        leitura="medida_por_categoria",
        agrupamento=_ref(beta),
        grafico="barra",
    )
    participacao["graficos"][0]["medida"] = _ref(parte)
    documento = definir(perfil, participacao)
    salva = _gravar(escopo, pacote["id"], documento)
    indicador = _calcular(escopo, salva["id"])["calculo"]["resultado"]["indicadores"][0]
    assert indicador["total"] == "1"
    assert indicador["serie"] == [
        {"chave": "norte", "resultado": "0.8", "elegiveis": 2},
        {"chave": "sul", "resultado": "0.2", "elegiveis": 1},
    ]

    parte_filtrada = _indicador(pacote, "participacao", leitura="indicador_de_medida")
    parte_filtrada["graficos"][0]["medida"] = _ref(parte)
    documento = definir(perfil, parte_filtrada, provedor=_filtrar(beta, "igual", "faixa", "texto"))
    salva = _gravar(escopo, pacote["id"], documento)
    indicador = _calcular(escopo, salva["id"], parametros={"faixa": "norte"})["calculo"]["resultado"]["indicadores"][0]
    assert indicador["resultado"] == "0.8"


def test_filtros_de_data_e_dimensao(ctx):
    _usuario, _franquia, escopo, perfil, pacote, _aprovada = _jornada("calculo.filtro@test.com", "calculo-filtro")
    quando = next(coluna for coluna in perfil["colunas"] if coluna["aba"] == 0 and "data" in coluna["dimensoes"])
    beta = _coluna(perfil, indice=1, aba=0)
    valor = _coluna(perfil, dimensao_principal="valor", aba=0)

    por_data = _indicador(pacote, "soma", leitura="indicador_de_medida")
    por_data["graficos"][0]["medida"] = _ref(valor)
    documento = definir(perfil, por_data, provedor=_filtrar(quando, "entre", "inicio", "data", "fim"))
    salva = _gravar(escopo, pacote["id"], documento)
    indicador = _calcular(
        escopo,
        salva["id"],
        parametros={"inicio": "2024-03-01", "fim": "2024-03-02"},
    )["calculo"]["resultado"]["indicadores"][0]
    assert indicador["resultado"] == "14.75"
    assert indicador["elegiveis"] == 2
    assert any(item["motivo"] == "filtro" and item["quantidade"] == 1 for item in indicador["excluidos"])

    por_beta = _indicador(pacote, "soma", leitura="indicador_de_medida")
    por_beta["graficos"][0]["medida"] = _ref(valor)
    documento = definir(perfil, por_beta, provedor=_filtrar(beta, "igual", "faixa", "texto"))
    salva = _gravar(escopo, pacote["id"], documento)
    indicador = _calcular(escopo, salva["id"], parametros={"faixa": "sul"})["calculo"]["resultado"]["indicadores"][0]
    assert indicador["resultado"] == "4.25"
    assert indicador["elegiveis"] == 1
    assert date(2024, 3, 3).isoformat() not in json.dumps(indicador)


def test_data_e_unidade_ambiguas_bloqueiam_sem_fabricar_numero(ctx):
    _usuario, _franquia, escopo, perfil, pacote = _jornada_de(
        "calculo.ambiguo@test.com",
        "calculo-ambiguo",
        _xlsx(
            {
                "movimento": [
                    ["alfa", "quando", "valor"],
                    ["norte", date(2024, 1, 2), 10.5],
                    ["sul", "01/02/2024", 4.25],
                    ["norte", "02/01/2024", 10.5],
                ]
            }
        ),
    )
    quando = _coluna(perfil, indice=1, aba=0)
    valor = _coluna(perfil, indice=2, aba=0)
    assert quando["dominio_inconsistente"] is True
    pacote_data = _indicador(
        pacote,
        "soma",
        leitura="medida_por_categoria",
        agrupamento=_ref(quando),
        grafico="barra",
    )
    pacote_data["graficos"][0]["medida"] = _ref(valor)
    documento = definir(perfil, pacote_data)
    assert documento["definicoes"][0]["status"] == "bloqueada"
    salva = _gravar(escopo, pacote["id"], documento)
    indicador = _calcular(escopo, salva["id"])["calculo"]["resultado"]["indicadores"][0]
    assert indicador["status"] == "bloqueado"
    assert indicador["resultado"] is None
    assert "2024-01-02" not in json.dumps(indicador)
    assert "2024-02-01" not in json.dumps(indicador)

    _usuario_b, _franquia_b, escopo_b, perfil_b, pacote_b = _jornada_de(
        "calculo.misto@test.com",
        "calculo-misto",
        _xlsx(
            {
                "movimento": [
                    ["alfa", "valor", "nota"],
                    ["norte", 10.5, 10],
                    ["sul", 4.25, "dez"],
                    ["norte", 10.5, 3],
                ]
            }
        ),
    )
    medida = _coluna(perfil_b, indice=2, aba=0)
    assert medida["dominio_inconsistente"] is True
    pacote_misto = _indicador(pacote_b, "soma", leitura="indicador_de_medida")
    pacote_misto["graficos"][0]["medida"] = _ref(medida)
    documento = definir(perfil_b, pacote_misto)
    assert documento["definicoes"][0]["status"] == "bloqueada"
    salva = _gravar(escopo_b, pacote_b["id"], documento)
    indicador = _calcular(escopo_b, salva["id"])["calculo"]["resultado"]["indicadores"][0]
    assert indicador["resultado"] is None
    assert indicador["motivo"] == "resolver_tipo_misto"


def test_nulos_vazias_e_linhas_pre_preenchidas_nao_entram_na_conta(ctx):
    _usuario, _franquia, escopo, perfil, pacote = _jornada_de(
        "calculo.nulo@test.com",
        "calculo-nulo",
        _xlsx(
            {
                "movimento": [
                    ["alfa", "valor"],
                    ["a", 10],
                    ["b", None],
                    ["c", 5],
                    [None, None],
                ]
            }
        ),
    )
    valor = _coluna(perfil, indice=1, aba=0)
    pacote_soma = _indicador(pacote, "soma", leitura="indicador_de_medida")
    pacote_soma["graficos"][0]["medida"] = _ref(valor)
    documento = definir(perfil, pacote_soma)
    assert documento["definicoes"][0]["status"] == "definida"
    salva = _gravar(escopo, pacote["id"], documento)
    indicador = _calcular(escopo, salva["id"])["calculo"]["resultado"]["indicadores"][0]
    assert indicador["resultado"] == "15"
    assert indicador["elegiveis"] == 2
    assert any(item["motivo"] == "sem_valor" for item in indicador["excluidos"])
    assert "0" not in {item["resultado"] for item in indicador["serie"]}

    livro = _xlsx(
        {
            "movimento": [
                ["alfa", "valor"],
                ["norte", 10.5],
                ["sul", 4.25],
                ["norte", 10.5],
            ],
            "ajuste": [
                ["alfa", "valor"],
                ["a", 10],
                ["b", 20],
            ],
        }
    )
    from openpyxl import load_workbook
    import io

    planilha = load_workbook(io.BytesIO(livro))
    planilha["ajuste"].cell(4, 2, "=1+1")
    destino = io.BytesIO()
    planilha.save(destino)
    _usuario_p, _franquia_p, escopo_p, perfil_p, pacote_p = _jornada_de(
        "calculo.pre@test.com",
        "calculo-pre",
        destino.getvalue(),
    )
    regiao = next(item for item in perfil_p["qualidade"]["por_regiao"] if item["aba"] == 1)
    assert (regiao.get("linhas_pre_preenchidas") or 0) >= 1 or regiao["formulas_sem_cache"]["observadas"] >= 1
    medida = _coluna(perfil_p, indice=1, aba=1)
    pacote_pre = _indicador(pacote_p, "soma", leitura="indicador_de_medida")
    grafico = pacote_pre["graficos"][0]
    grafico["medida"] = _ref(medida)
    grafico["regiao"] = {"arquivo": medida["arquivo"], "aba": medida["aba"], "regiao": medida["regiao"]}
    documento = definir(perfil_p, pacote_pre)
    salva = _gravar(escopo_p, pacote_p["id"], documento)
    indicador = _calcular(escopo_p, salva["id"])["calculo"]["resultado"]["indicadores"][0]
    assert indicador["resultado"] is None
    assert indicador["motivo"] == "resolver_formula_sem_cache"
    assert indicador["resultado"] != "32"


def test_juncao_valida_preserva_soma_e_invalida_nao_multiplica(ctx):
    def montar(movimento, cadastro):
        return _xlsx({"movimento": movimento, "cadastro": cadastro})

    def executar(email, slug, movimento, cadastro):
        _usuario, _franquia, escopo, perfil, pacote = _jornada_de(email, slug, montar(movimento, cadastro))
        assert perfil["relacoes"]
        relacao = perfil["relacoes"][0]
        valor = _coluna(perfil, indice=1, aba=0)
        pacote_soma = _indicador(pacote, "soma", leitura="indicador_de_medida")
        pacote_soma["graficos"][0]["medida"] = _ref(valor)

        def provedor(contexto):
            base = provedor_simulado(contexto)
            item = base["definicoes"][0]
            item["relacionamentos"] = [{"indice": 0, "juncao": "nao_executada"}]
            item["status"] = "bloqueada"
            item["esclarecimento"] = "confirmar_relacao_candidata"
            return base

        documento = definir(perfil, pacote_soma, provedor=provedor)
        salva = _gravar(escopo, pacote["id"], documento)
        indicador = _calcular(escopo, salva["id"], relacoes_autorizadas=[0])["calculo"]["resultado"]["indicadores"][0]
        return relacao, indicador

    movimento = [
        ["codigo", "valor"],
        ["A", 10],
        ["A", 10],
        ["B", 5],
    ]
    _relacao, valida = executar(
        "calculo.join@test.com",
        "calculo-join",
        movimento,
        [["codigo", "nivel"], ["A", 1], ["B", 2]],
    )
    assert valida["status"] == "calculado"
    assert valida["resultado"] == "25"
    assert valida["elegiveis"] == 3
    assert valida["juncao"] == "muitos_para_um"
    assert "juncao_sem_duplicacao" in valida["avisos"]

    _relacao, invalida = executar(
        "calculo.joinbad@test.com",
        "calculo-join-bad",
        [["codigo", "valor"], ["A", 10], ["B", 10]],
        [["codigo", "nivel"], ["A", 1], ["B", 2], ["C", 3], ["C", 4]],
    )
    assert invalida["status"] == "bloqueado"
    assert invalida["resultado"] is None
    assert invalida["motivo"] == "juncao_ambigua"
    assert invalida["resultado"] != "65"


def test_isolamento_repeticao_e_releitura_da_memoria(ctx, caplog):
    usuario, franquia, escopo, _perfil, pacote, _aprovada = _jornada("calculo.iso@test.com", "calculo-iso")
    _outra, _outra_franquia, outro = _par("calculo-outro", "calculo.outro@test.com")
    colega = seed_usuario(franquia.id, usuario.conta_id, email="calculo.colega@test.com")
    outro_escopo = escopo_do_usuario(outro)
    colega_escopo = escopo_do_usuario(colega)
    gravada = definir_selecao(escopo, {"analise_id": pacote["id"]})
    with caplog.at_level(logging.INFO):
        primeira = _calcular(escopo, gravada["definicao"]["id"])
        segunda = _calcular(escopo, gravada["definicao"]["id"])
    assert primeira["calculo"]["resultado"]["indicadores"] == segunda["calculo"]["resultado"]["indicadores"]
    assert primeira["calculo"]["id"] != segunda["calculo"]["id"]
    relida = obter_analise(escopo, primeira["calculo"]["id"])
    assert relida["resultado"] == primeira["calculo"]["resultado"]
    assert relida["etapa"] == "execucao"
    assert relida["contrato_versao"] == CONTRATO_CALCULO
    assert IaChamadaTentativa.query.count() == 0
    assert SENTINELA not in caplog.text
    with pytest.raises(Exception) as erro:
        _calcular(outro_escopo, gravada["definicao"]["id"])
    assert getattr(erro.value, "codigo", "") == "nao_encontrado"
    with pytest.raises(Exception) as colega_erro:
        obter_analise(colega_escopo, primeira["calculo"]["id"])
    assert getattr(colega_erro.value, "codigo", "") == "nao_encontrado"


def test_limites_de_volume_e_agrupamento_nao_devolvem_parcial(ctx, monkeypatch):
    _usuario, _franquia, escopo, perfil, pacote, _aprovada = _jornada("calculo.limite@test.com", "calculo-limite")
    gravada = definir_selecao(escopo, {"analise_id": pacote["id"]})
    monkeypatch.setattr("app.services.agente_bi_calculo_service.MAX_LINHAS", 1)
    indicador = next(
        item
        for item in _calcular(escopo, gravada["definicao"]["id"])["calculo"]["resultado"]["indicadores"]
        if item["operacao"] == "soma" and item["grafico"] == "indicador"
    )
    assert indicador["status"] == "bloqueado"
    assert indicador["motivo"] == "limite_volume"
    assert indicador["resultado"] is None

    monkeypatch.setattr("app.services.agente_bi_calculo_service.MAX_LINHAS", 200_000)
    monkeypatch.setattr("app.services.agente_bi_calculo_service.MAX_GRUPOS", 1)
    agrupado = next(
        item
        for item in _calcular(escopo, gravada["definicao"]["id"])["calculo"]["resultado"]["indicadores"]
        if item["grafico"] == "barra"
    )
    assert agrupado["status"] == "bloqueado"
    assert agrupado["motivo"] == "limite_agrupamento"
    assert agrupado["serie"] == []
    assert agrupado["resultado"] is None


def test_api_rejeita_expressao_e_nao_chama_modelo(app, monkeypatch):
    client = _cliente_ingestao(app, monkeypatch)

    def proibido(*_args, **_kwargs):
        raise AssertionError("anthropic")

    monkeypatch.setattr("app.services.agente_bi_anthropic_service.chamar_anthropic_governado", proibido)
    with app.app_context():
        _conta, _franquia, usuario = _par("calculo-api", "calculo.api@test.com")
        _outra, _outra_franquia, outro = _par("calculo-api-outro", "calculo.api.outro@test.com")
        usuario_id = usuario.id
        outro_id = outro.id
    assert client.post("/api/agente-bi/calculo", json={"analise_id": 1}).status_code == 401
    _entrar(client, usuario_id)
    from tests.test_agente_bi_descoberta import _aprovar, _descobrir, _postar, _viaveis

    preparacao = _postar(client, _movimento(), "movimento.xlsx")
    assert preparacao.status_code == 201
    dataset_id = preparacao.get_json()["dataset"]["id"]
    descoberta = _descobrir(client, dataset_id, {"modalidades": ["A"]})
    aprovacao = _aprovar(
        client,
        descoberta["analise_id"],
        _viaveis(descoberta["resultado"])[:6],
        descoberta["resultado"],
    )
    selecao_id = aprovacao.get_json()["selecao"]["id"]
    definicao = client.post("/api/agente-bi/definicao", json={"analise_id": selecao_id})
    assert definicao.status_code == 201
    resposta = client.post("/api/agente-bi/calculo", json={"analise_id": definicao.get_json()["definicao"]["id"]})
    assert resposta.status_code == 201, resposta.get_data(as_text=True)
    corpo = resposta.get_json()
    assert corpo["calculo"]["contrato_versao"] == CONTRATO_CALCULO
    assert corpo["calculo"]["resultado"]["ia_calculo"] is False
    soma = _por_operacao(corpo["calculo"]["resultado"], "soma")
    assert soma["resultado"] == "25.25"
    reservado = client.post("/api/agente-bi/analises", json={"contrato_versao": CONTRATO_CALCULO})
    assert reservado.status_code == 422
    contexto = client.post("/api/agente-bi/analises", json={"contexto": {"contrato": CONTRATO_CONTEXTO_CALCULO}})
    assert contexto.status_code == 422
    _entrar(client, outro_id)
    alheio = client.post("/api/agente-bi/calculo", json={"analise_id": definicao.get_json()["definicao"]["id"]})
    assert alheio.status_code == 404
    fonte = (ROOT / "app" / "services" / "agente_bi_calculo_service.py").read_text(encoding="utf-8")
    rotas = (ROOT / "app" / "agente_bi_api_routes.py").read_text(encoding="utf-8")
    assert "chamar_anthropic" not in fonte
    assert "import anthropic" not in fonte
    assert "chamar_anthropic" not in rotas
    assert "eval(" not in rotas
    assert "exec(" not in rotas
    with app.app_context():
        with pytest.raises(CalculoErro):
            calcular_definicao(escopo_do_usuario(usuario), {"analise_id": 1, "sql": "select 1"})


def _barra(email, slug, linhas, operacao):
    _usuario, _franquia, escopo, perfil, pacote = _jornada_de(email, slug, _xlsx({"movimento": linhas}))
    grupo = _coluna(perfil, indice=0, aba=0)
    medida = _coluna(perfil, indice=1, aba=0)
    pedido = _indicador(
        pacote,
        operacao,
        leitura="medida_por_categoria",
        agrupamento=_ref(grupo),
        grafico="barra",
    )
    pedido["graficos"][0]["medida"] = _ref(medida)
    documento = definir(perfil, pedido)
    assert documento["definicoes"][0]["status"] == "definida"
    salva = _gravar(escopo, pacote["id"], documento)
    return _calcular(escopo, salva["id"])["calculo"]["resultado"]["indicadores"][0]


def _soma_serie(indicador) -> Decimal:
    return sum((Decimal(item["resultado"]) for item in indicador["serie"]), Decimal(0))


def test_total_agrupado_nao_soma_operacoes_nao_aditivas(ctx):
    media = _barra(
        "calculo.media.grupo@test.com",
        "calculo-media-grupo",
        [["grupo", "valor"], ["a", 10.5], ["a", 10.5], ["b", 31.5]],
        "media",
    )
    assert media["total"] == "17.5"
    assert media["resultado"] is None
    assert media["elegiveis"] == 3
    assert [item["resultado"] for item in media["serie"]] == ["10.5", "31.5"]
    assert _soma_serie(media) != Decimal(media["total"])
    assert all(set(item) == {"chave", "resultado", "elegiveis"} for item in media["serie"])

    distinta = _barra(
        "calculo.distinta.grupo@test.com",
        "calculo-distinta-grupo",
        [["grupo", "valor"], ["norte", 10.5], ["sul", 20.5], ["leste", 10.5]],
        "contagem_distinta",
    )
    assert distinta["total"] == "2"
    assert [item["resultado"] for item in distinta["serie"]] == ["1", "1", "1"]
    assert _soma_serie(distinta) != Decimal(distinta["total"])

    extremos = [["grupo", "valor"], ["a", 10.5], ["b", 20.5], ["c", 31.5], ["a", None]]
    minimo = _barra("calculo.minimo.grupo@test.com", "calculo-minimo-grupo", extremos, "minimo")
    maximo = _barra("calculo.maximo.grupo@test.com", "calculo-maximo-grupo", extremos, "maximo")
    assert minimo["total"] == "10.5"
    assert maximo["total"] == "31.5"
    assert [item["resultado"] for item in minimo["serie"]] == ["10.5", "20.5", "31.5"]
    assert [item["resultado"] for item in maximo["serie"]] == ["10.5", "20.5", "31.5"]
    assert _soma_serie(minimo) != Decimal(minimo["total"])
    assert _soma_serie(maximo) != Decimal(maximo["total"])


def test_agrupamento_nulo_entra_so_nas_exclusoes(ctx):
    indicador = _barra(
        "calculo.grupo.nulo@test.com",
        "calculo-grupo-nulo",
        [["grupo", "valor"], ["a", 10.5], ["b", 20.5], [None, 10.5]],
        "soma",
    )
    assert indicador["total"] == "31"
    assert indicador["elegiveis"] == 2
    assert sum(item["quantidade"] for item in indicador["excluidos"]) == 1
    assert indicador["excluidos"] == [{"motivo": "sem_valor", "quantidade": 1}]
    assert [item["chave"] for item in indicador["serie"]] == ["a", "b"]
    assert "" not in {item["chave"] for item in indicador["serie"]}
    assert _soma_serie(indicador) == Decimal(indicador["total"])

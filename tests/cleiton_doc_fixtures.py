"""Fixtures compartilhadas para testes documentais do Cleiton (Fase 3)."""
from __future__ import annotations

import io
import zipfile
from types import SimpleNamespace

import pytest

import app.cleiton_doc_store as store
from app.services import cleiton_doc_config_service as doc_cfg_svc


def patch_cleiton_doc_store(tmp_path, monkeypatch):
    monkeypatch.setattr(store, "get_cleiton_doc_tmp_dir", lambda: str(tmp_path))


def patch_cleiton_doc_cfg(monkeypatch, **overrides):
    base = doc_cfg_svc.get_cleiton_doc_config()
    cfg = SimpleNamespace(
        upload_enabled=overrides.get("upload_enabled", base.upload_enabled),
        max_files_per_session=overrides.get("max_files_per_session", base.max_files_per_session),
        session_max_bytes=overrides.get("session_max_bytes", base.session_max_bytes),
        upload_ttl_hours=overrides.get("upload_ttl_hours", base.upload_ttl_hours),
        cleanup_enabled=overrides.get("cleanup_enabled", base.cleanup_enabled),
        prompt_context_max_chars=overrides.get("prompt_context_max_chars", base.prompt_context_max_chars),
        prompt_max_files_considered=overrides.get("prompt_max_files_considered", base.prompt_max_files_considered),
        pdf_enabled=overrides.get("pdf_enabled", base.pdf_enabled),
        pdf_max_bytes=overrides.get("pdf_max_bytes", base.pdf_max_bytes),
        pdf_max_pages=overrides.get("pdf_max_pages", base.pdf_max_pages),
        pdf_max_chars=overrides.get("pdf_max_chars", base.pdf_max_chars),
        excel_enabled=overrides.get("excel_enabled", base.excel_enabled),
        excel_max_bytes=overrides.get("excel_max_bytes", base.excel_max_bytes),
        excel_max_rows=overrides.get("excel_max_rows", base.excel_max_rows),
        excel_max_columns=overrides.get("excel_max_columns", base.excel_max_columns),
        excel_max_chars=overrides.get("excel_max_chars", base.excel_max_chars),
        docx_enabled=overrides.get("docx_enabled", base.docx_enabled),
        docx_max_bytes=overrides.get("docx_max_bytes", base.docx_max_bytes),
        docx_max_paragraphs=overrides.get("docx_max_paragraphs", base.docx_max_paragraphs),
        docx_max_chars=overrides.get("docx_max_chars", base.docx_max_chars),
        txt_enabled=overrides.get("txt_enabled", base.txt_enabled),
        txt_max_bytes=overrides.get("txt_max_bytes", base.txt_max_bytes),
        txt_max_chars=overrides.get("txt_max_chars", base.txt_max_chars),
        xml_enabled=overrides.get("xml_enabled", base.xml_enabled),
        xml_max_bytes=overrides.get("xml_max_bytes", base.xml_max_bytes),
        xml_max_nodes=overrides.get("xml_max_nodes", base.xml_max_nodes),
        xml_max_depth=overrides.get("xml_max_depth", base.xml_max_depth),
        xml_max_chars=overrides.get("xml_max_chars", base.xml_max_chars),
        csv_enabled=overrides.get("csv_enabled", base.csv_enabled),
        csv_max_bytes=overrides.get("csv_max_bytes", base.csv_max_bytes),
        csv_max_rows=overrides.get("csv_max_rows", base.csv_max_rows),
        csv_max_columns=overrides.get("csv_max_columns", base.csv_max_columns),
        csv_max_chars=overrides.get("csv_max_chars", base.csv_max_chars),
    )
    monkeypatch.setattr("app.cleiton_doc_security.get_cleiton_doc_config", lambda: cfg)
    monkeypatch.setattr("app.cleiton_doc_prepare.get_cleiton_doc_config", lambda: cfg)
    monkeypatch.setattr("app.cleiton_doc_service.get_cleiton_doc_config", lambda: cfg)
    monkeypatch.setattr("app.julia_doc_context.get_cleiton_doc_config", lambda: cfg)
    monkeypatch.setattr("app.julia_documents_routes.get_cleiton_doc_config", lambda: cfg)
    return cfg


@pytest.fixture
def doc_cfg(monkeypatch, ctx):
    return patch_cleiton_doc_cfg(monkeypatch)


def patch_gemini_pdf_upload(monkeypatch, *, upload_state: str = "ACTIVE"):
    """Mock governado de upload Gemini Files API para testes de PDF."""
    from unittest.mock import MagicMock

    import app.cleiton_doc_gemini_files as gemini_files

    client = MagicMock()
    uploaded = type(
        "Uploaded",
        (),
        {
            "name": "files/test-pdf-mock",
            "uri": "https://generativelanguage.googleapis.com/v1beta/files/test-pdf-mock",
            "mime_type": "application/pdf",
            "state": upload_state,
        },
    )()
    client.files.upload.return_value = uploaded
    client.files.get.return_value = uploaded
    client.files.delete.return_value = None
    monkeypatch.setattr(gemini_files, "get_cleiton_gemini_client", lambda: client)
    return client


def make_minimal_pdf(*, pages: int = 1) -> bytes:
    chunks = [b"%PDF-1.4\n"]
    for _ in range(pages):
        chunks.append(b"1 0 obj\n<< /Type /Page >>\nendobj\n")
    chunks.append(b"%%EOF\n")
    return b"".join(chunks)


def make_invalid_pdf() -> bytes:
    return b"NOTPDF-content"


def make_txt(content: str, encoding: str = "utf-8") -> bytes:
    return content.encode(encoding)


def make_xml(content: str) -> bytes:
    return content.encode("utf-8")


def make_billion_laughs_xml() -> bytes:
    payload = """<?xml version="1.0"?>
<!DOCTYPE lolz [
 <!ENTITY lol "lol">
 <!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">
 <!ENTITY lol3 "&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;">
 <!ENTITY lol4 "&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;">
 <!ENTITY lol5 "&lol4;&lol4;&lol4;&lol4;&lol4;&lol4;&lol4;&lol4;&lol4;&lol4;">
 <!ENTITY lol6 "&lol5;&lol5;&lol5;&lol5;&lol5;&lol5;&lol5;&lol5;&lol5;&lol5;">
 <!ENTITY lol7 "&lol6;&lol6;&lol6;&lol6;&lol6;&lol6;&lol6;&lol6;&lol6;&lol6;">
 <!ENTITY lol8 "&lol7;&lol7;&lol7;&lol7;&lol7;&lol7;&lol7;&lol7;&lol7;&lol7;">
 <!ENTITY lol9 "&lol8;&lol8;&lol8;&lol8;&lol8;&lol8;&lol8;&lol8;&lol8;&lol8;">
]>
<lolz>&lol9;</lolz>"""
    return payload.encode("utf-8")


def make_nested_xml(depth: int) -> bytes:
    inner = "x"
    for idx in range(depth):
        inner = f"<n{idx}>{inner}</n{idx}>"
    return f'<?xml version="1.0"?><root>{inner}</root>'.encode("utf-8")


def make_many_nodes_xml(count: int) -> bytes:
    nodes = "".join(f"<item id='{idx}'/>" for idx in range(count))
    return f"<?xml version='1.0'?><root>{nodes}</root>".encode("utf-8")


def make_csv(rows: list[list[str]]) -> bytes:
    import csv

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    for row in rows:
        writer.writerow(row)
    return buffer.getvalue().encode("utf-8")


def make_xlsx(rows: list[list[str]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for row_idx, row in enumerate(rows, start=1):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def make_corrupted_xlsx() -> bytes:
    return b"PK\x03\x04corrupted-xlsx"


def make_docx(paragraphs: list[str]) -> bytes:
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def make_corrupted_docx() -> bytes:
    return b"PK\x03\x04corrupted-docx"


def make_zip_bomb_like() -> bytes:
    payload = b"0" * (1024 * 1024)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("sheet1.xml", payload)
    return buffer.getvalue()

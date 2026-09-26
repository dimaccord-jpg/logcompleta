"""Testes direcionados do Lote 3B SCRUM-148: task_* Cleide BI, Chat Cleide BI, Roberto BI, Roberto Chat."""
from __future__ import annotations

import io
from types import SimpleNamespace

import pytest

from app.cleide_analytics import AnalyticsProcessingError
from app.cleide_controlled_chat import run_cleide_controlled_chat
from app.funnel_event_service import (
    FUNNEL_EVENT_TASK_COMPLETED,
    FUNNEL_EVENT_TASK_FAILED,
    FUNNEL_EVENT_TASK_PREPARATION_COMPLETED,
    FUNNEL_EVENT_TASK_PREPARATION_STARTED,
    FUNNEL_EVENT_TASK_STARTED,
    FUNNEL_SOURCE_CLEIDE_BI,
    FUNNEL_SOURCE_CLEIDE_BI_CHAT,
    FUNNEL_SOURCE_ROBERTO_BI,
    FUNNEL_SOURCE_ROBERTO_CHAT,
    META_PIXEL_ALLOWED_EVENTS,
    TASK_TYPE_CLEIDE_BI,
    TASK_TYPE_CLEIDE_BI_CHAT,
    TASK_TYPE_ROBERTO_BI,
    TASK_TYPE_ROBERTO_CHAT,
    is_meta_pixel_allowed,
)
from app.models import FunnelEvent
from tests.conftest import seed_conta_franquia_cliente, seed_usuario


def _enable_session(app):
    app.config["SECRET_KEY"] = "test-secret"
    app.config["TESTING"] = True


def _auth_user(monkeypatch, *, email: str):
    conta, franquia = seed_conta_franquia_cliente(slug=f"lote3b-{email.split('@')[0]}")
    user = seed_usuario(franquia.id, conta.id, email=email)
    fake = SimpleNamespace(
        is_authenticated=True,
        id=user.id,
        conta_id=conta.id,
        franquia_id=franquia.id,
    )
    monkeypatch.setattr("flask_login.utils._get_user", lambda: fake)
    monkeypatch.setattr(
        "app.funnel_event_service._is_desktop_access_admin_test_mode",
        lambda: False,
    )
    return user


def _events_for(user_id: int, event_name: str):
    return (
        FunnelEvent.query.filter_by(user_id=user_id, event_name=event_name)
        .order_by(FunnelEvent.id.asc())
        .all()
    )


def _xlsx_bytes():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    headers = ("transportadora", "uf_origem", "uf_destino", "valor_frete", "peso", "data_emissao")
    row = ("XP", "SP", "RJ", 100, 10, "2026-01-01")
    for cidx, cell in enumerate(headers, start=1):
        ws.cell(row=1, column=cidx, value=cell)
    for cidx, cell in enumerate(row, start=1):
        ws.cell(row=2, column=cidx, value=cell)
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def _chat_context(*, total_docs=100):
    return {
        "chat_context_version": "cleide_chat_context.v1",
        "chat_ready_context": True,
        "safe_operational_context": {
            "schema_version": "cleide_contexto_operacional.v1",
            "session_scope": {"dataset_validado": True},
            "kpis": {
                "total_documentos": total_docs,
                "valor_total_frete": 12345.67,
                "peso_total": 54321.0,
                "ticket_medio_frete": 123.45,
                "percentual_fretes_zerados": 4.5,
                "periodo_dataset": {"inicio": "2026-01-01", "fim": "2026-01-31"},
            },
            "aggregate_tables": {
                "transportadora": [
                    {"chave": "XP", "quantidade": 80},
                    {"chave": "YZ", "quantidade": 20},
                ],
                "uf_origem": [{"chave": "SP", "quantidade": 70, "valor_total": 9000.0}],
                "uf_destino": [{"chave": "RJ", "quantidade": 60, "valor_total": 8000.0}],
                "temporal": [{"data": "2026-01", "quantidade": 100}],
            },
            "dataset_summary": {
                "linhas_processadas": 100,
                "invalid_numeric_rows": 0,
                "invalid_date_rows": 0,
                "negative_value_rows": 0,
            },
            "quality_flags": {
                "has_invalid_numeric": False,
                "has_invalid_date": False,
                "has_negative_values": False,
                "has_sparse_aggregates": False,
            },
            "filter_context": {
                "active_filters": {},
                "filter_mode": "row_level_intersection_backend",
                "kpi_scope": "filtered_session_intersection",
            },
            "semantic_limits": {
                "no_row_level_intersection": False,
                "multi_dimension_filters_are_approximate": False,
                "kpis_are_global_session_scope": False,
                "no_accusatory_financial_conclusion": True,
            },
            "language_policy": {
                "allowed_language": [
                    "concentração operacional",
                    "comportamento atípico",
                    "variação relevante",
                    "oportunidade de investigação",
                    "dados insuficientes",
                    "tendência operacional",
                    "participação relevante",
                    "ranking",
                    "uf",
                    "transportadora",
                    "frete",
                    "periodo",
                    "ticket",
                    "volume",
                    "operacional",
                ],
                "forbidden_language": [],
                "out_of_scope_language": [],
            },
        },
    }


def test_lote3b_task_events_not_meta_authorized():
    for name in (
        FUNNEL_EVENT_TASK_PREPARATION_STARTED,
        FUNNEL_EVENT_TASK_PREPARATION_COMPLETED,
        FUNNEL_EVENT_TASK_STARTED,
        FUNNEL_EVENT_TASK_COMPLETED,
        FUNNEL_EVENT_TASK_FAILED,
    ):
        assert name not in META_PIXEL_ALLOWED_EVENTS
        assert is_meta_pixel_allowed(name) is False


def test_cleide_bi_success_emits_task_completed(app, ctx, monkeypatch, tmp_path):
    from app.cleide_upload_pipeline import process_cleide_upload

    _enable_session(app)
    monkeypatch.setattr(
        "app.cleide_upload_pipeline.get_cleide_config",
        lambda: SimpleNamespace(
            upload_total_max=10000,
            upload_max_file_size_bytes=2 * 1024 * 1024,
            upload_ttl_minutes=30,
            csv_delimiter_default=",",
            structural_max_rows=10000,
            structural_max_columns=120,
            analytics_max_rows=10000,
            analytics_group_limit=25,
        ),
    )
    monkeypatch.setattr("app.cleide_upload_store.get_cleide_upload_tmp_dir", lambda: str(tmp_path))
    monkeypatch.setattr(
        "app.cleide_upload_pipeline.apropriar_billing_upload_cleide",
        lambda **_k: None,
    )

    with app.app_context():
        user = _auth_user(monkeypatch, email="lote3b-cleide-bi-ok@test.com")
        with app.test_request_context(
            "/api/cleide/upload",
            method="POST",
            data={"file": (io.BytesIO(_xlsx_bytes()), "base.xlsx")},
            content_type="multipart/form-data",
            headers={"X-Execution-ID": "exec-cleide-bi-ok"},
        ):
            resp, status = process_cleide_upload()
            assert status == 200
            body = resp.get_json()
            assert body["success"] is True
            assert body["analytics_ready"] is True
            assert body["dataset_validado"] is True

        completed = _events_for(user.id, FUNNEL_EVENT_TASK_COMPLETED)
        started = _events_for(user.id, FUNNEL_EVENT_TASK_STARTED)
        prep_started = _events_for(user.id, FUNNEL_EVENT_TASK_PREPARATION_STARTED)
        prep_completed = _events_for(user.id, FUNNEL_EVENT_TASK_PREPARATION_COMPLETED)
        assert len(prep_started) == 1
        assert len(prep_completed) == 1
        assert len(started) == 1
        assert len(completed) == 1
        assert completed[0].source == FUNNEL_SOURCE_CLEIDE_BI
        assert completed[0].metadata_json["task_type"] == TASK_TYPE_CLEIDE_BI
        assert completed[0].execution_id == "exec-cleide-bi-ok"

        # Idempotencia: mesmo execution_id + marco nao duplica.
        with app.test_request_context(
            "/api/cleide/upload",
            method="POST",
            data={"file": (io.BytesIO(_xlsx_bytes()), "base2.xlsx")},
            content_type="multipart/form-data",
            headers={"X-Execution-ID": "exec-cleide-bi-ok"},
        ):
            process_cleide_upload()
        assert len(_events_for(user.id, FUNNEL_EVENT_TASK_COMPLETED)) == 1


def test_cleide_bi_dataset_validado_false_emits_task_failed_only(app, ctx, monkeypatch, tmp_path):
    from app.cleide_upload_pipeline import process_cleide_upload

    _enable_session(app)
    monkeypatch.setattr(
        "app.cleide_upload_pipeline.get_cleide_config",
        lambda: SimpleNamespace(
            upload_total_max=10000,
            upload_max_file_size_bytes=2 * 1024 * 1024,
            upload_ttl_minutes=30,
            csv_delimiter_default=",",
            structural_max_rows=10000,
            structural_max_columns=120,
            analytics_max_rows=10000,
            analytics_group_limit=25,
        ),
    )
    monkeypatch.setattr("app.cleide_upload_store.get_cleide_upload_tmp_dir", lambda: str(tmp_path))
    monkeypatch.setattr(
        "app.cleide_upload_pipeline.apropriar_billing_upload_cleide",
        lambda **_k: None,
    )
    monkeypatch.setattr(
        "app.cleide_upload_pipeline.analyze_structural_layout",
        lambda **_k: {
            "dataset_validado": False,
            "dataset_tipo": "xlsx",
            "sheet_detectada": "Sheet",
            "linhas_detectadas": 1,
            "colunas_detectadas": ["transportadora"],
            "colunas_faltantes": ["uf_origem"],
            "aliases_resolvidos": {},
            "raw_headers": ["transportadora"],
            "canonical_headers": ["transportadora"],
            "normalized_columns": ["transportadora"],
            "colunas_duplicadas": [],
            "delimiter_detectado": None,
            "detected_encoding": None,
        },
    )

    with app.app_context():
        user = _auth_user(monkeypatch, email="lote3b-cleide-bi-invalid@test.com")
        with app.test_request_context(
            "/api/cleide/upload",
            method="POST",
            data={"file": (io.BytesIO(_xlsx_bytes()), "base.xlsx")},
            content_type="multipart/form-data",
            headers={"X-Execution-ID": "exec-cleide-bi-invalid"},
        ):
            resp, status = process_cleide_upload()
            # Resposta funcional original preservada (fluxo continua sem exception).
            assert status == 200
            body = resp.get_json()
            assert body["success"] is True
            assert body["dataset_validado"] is False

        assert len(_events_for(user.id, FUNNEL_EVENT_TASK_PREPARATION_STARTED)) == 1
        assert len(_events_for(user.id, FUNNEL_EVENT_TASK_PREPARATION_COMPLETED)) == 0
        assert len(_events_for(user.id, FUNNEL_EVENT_TASK_STARTED)) == 0
        assert len(_events_for(user.id, FUNNEL_EVENT_TASK_COMPLETED)) == 0
        failed = _events_for(user.id, FUNNEL_EVENT_TASK_FAILED)
        assert len(failed) == 1
        assert failed[0].metadata_json["task_type"] == TASK_TYPE_CLEIDE_BI
        assert failed[0].metadata_json["task_stage"] == "structural_validation"
        assert failed[0].metadata_json["error_code"] == "structural_validation_failed"


def test_cleide_bi_analytics_failure_emits_task_failed(app, ctx, monkeypatch, tmp_path):
    from app.cleide_upload_pipeline import process_cleide_upload

    _enable_session(app)
    monkeypatch.setattr(
        "app.cleide_upload_pipeline.get_cleide_config",
        lambda: SimpleNamespace(
            upload_total_max=10000,
            upload_max_file_size_bytes=2 * 1024 * 1024,
            upload_ttl_minutes=30,
            csv_delimiter_default=",",
            structural_max_rows=10000,
            structural_max_columns=120,
            analytics_max_rows=10000,
            analytics_group_limit=25,
        ),
    )
    monkeypatch.setattr("app.cleide_upload_store.get_cleide_upload_tmp_dir", lambda: str(tmp_path))
    monkeypatch.setattr(
        "app.cleide_upload_pipeline.apropriar_billing_upload_cleide",
        lambda **_k: None,
    )
    monkeypatch.setattr(
        "app.cleide_upload_pipeline.build_analytics_context",
        lambda **_k: (_ for _ in ()).throw(
            AnalyticsProcessingError(code="analytics_boom", message="falha analytics")
        ),
    )

    with app.app_context():
        user = _auth_user(monkeypatch, email="lote3b-cleide-bi-fail@test.com")
        with app.test_request_context(
            "/api/cleide/upload",
            method="POST",
            data={"file": (io.BytesIO(_xlsx_bytes()), "base.xlsx")},
            content_type="multipart/form-data",
            headers={"X-Execution-ID": "exec-cleide-bi-fail"},
        ):
            resp, status = process_cleide_upload()
            assert status == 400
            assert resp.get_json()["success"] is False

        failed = _events_for(user.id, FUNNEL_EVENT_TASK_FAILED)
        started = _events_for(user.id, FUNNEL_EVENT_TASK_STARTED)
        assert len(started) == 1
        assert len(failed) == 1
        assert failed[0].metadata_json["task_type"] == TASK_TYPE_CLEIDE_BI
        assert failed[0].metadata_json["error_code"] == "analytics_boom"
        assert FunnelEvent.query.filter_by(user_id=user.id, event_name=FUNNEL_EVENT_TASK_COMPLETED).count() == 0


def test_cleide_bi_growth_fail_open_preserves_functional_response(app, ctx, monkeypatch, tmp_path):
    from app.cleide_upload_pipeline import process_cleide_upload

    _enable_session(app)
    monkeypatch.setattr(
        "app.cleide_upload_pipeline.get_cleide_config",
        lambda: SimpleNamespace(
            upload_total_max=10000,
            upload_max_file_size_bytes=2 * 1024 * 1024,
            upload_ttl_minutes=30,
            csv_delimiter_default=",",
            structural_max_rows=10000,
            structural_max_columns=120,
            analytics_max_rows=10000,
            analytics_group_limit=25,
        ),
    )
    monkeypatch.setattr("app.cleide_upload_store.get_cleide_upload_tmp_dir", lambda: str(tmp_path))
    monkeypatch.setattr(
        "app.cleide_upload_pipeline.apropriar_billing_upload_cleide",
        lambda **_k: None,
    )
    monkeypatch.setattr(
        "app.funnel_event_service.record_funnel_event",
        lambda **_k: (_ for _ in ()).throw(RuntimeError("analytics down")),
    )

    with app.app_context():
        _auth_user(monkeypatch, email="lote3b-cleide-bi-failopen@test.com")
        with app.test_request_context(
            "/api/cleide/upload",
            method="POST",
            data={"file": (io.BytesIO(_xlsx_bytes()), "base.xlsx")},
            content_type="multipart/form-data",
            headers={"X-Execution-ID": "exec-cleide-bi-failopen"},
        ):
            resp, status = process_cleide_upload()
            assert status == 200
            assert resp.get_json()["success"] is True
            assert resp.get_json()["analytics_ready"] is True


def test_cleide_bi_chat_started_completed(app, ctx, monkeypatch):
    import app.cleide_controlled_chat as controlled_chat

    monkeypatch.setattr(
        controlled_chat,
        "resolve_cleide_ai_flags",
        lambda: SimpleNamespace(
            ai_enabled=False,
            environment="test",
            selected_flag="off",
            reason="test",
            has_api_key=False,
            api_key_label="",
        ),
    )
    monkeypatch.setattr(controlled_chat, "get_cleide_chat_context", lambda _s: _chat_context())
    monkeypatch.setattr(
        controlled_chat,
        "_resolve_authz",
        lambda: {"permitido": True, "modo_operacao": "normal"},
    )

    with app.app_context():
        user = _auth_user(monkeypatch, email="lote3b-cleide-chat-ok@test.com")
        with app.test_request_context():
            body, status = run_cleide_controlled_chat(
                question="resumo operacional",
                session_obj={},
            )
        assert status == 200
        assert body.get("reply")
        assert body.get("fallback_used") is False

        started = _events_for(user.id, FUNNEL_EVENT_TASK_STARTED)
        completed = _events_for(user.id, FUNNEL_EVENT_TASK_COMPLETED)
        assert len(started) == 1
        assert len(completed) == 1
        assert started[0].source == FUNNEL_SOURCE_CLEIDE_BI_CHAT
        assert started[0].metadata_json["task_type"] == TASK_TYPE_CLEIDE_BI_CHAT
        assert completed[0].metadata_json["task_type"] == TASK_TYPE_CLEIDE_BI_CHAT


def test_cleide_bi_chat_final_failure_emits_task_failed(app, ctx, monkeypatch):
    import app.cleide_controlled_chat as controlled_chat

    monkeypatch.setattr(
        controlled_chat,
        "resolve_cleide_ai_flags",
        lambda: SimpleNamespace(
            ai_enabled=False,
            environment="test",
            selected_flag="off",
            reason="test",
            has_api_key=False,
            api_key_label="",
        ),
    )
    monkeypatch.setattr(controlled_chat, "get_cleide_chat_context", lambda _s: _chat_context())
    monkeypatch.setattr(
        controlled_chat,
        "_resolve_authz",
        lambda: {"permitido": True, "modo_operacao": "normal"},
    )

    with app.app_context():
        user = _auth_user(monkeypatch, email="lote3b-cleide-chat-fail@test.com")
        with app.test_request_context():
            body, status = run_cleide_controlled_chat(
                question="xyz intenção totalmente desconhecida sem sentido",
                session_obj={},
            )
        assert status == 200
        assert body.get("fallback_used") is True

        started = _events_for(user.id, FUNNEL_EVENT_TASK_STARTED)
        failed = _events_for(user.id, FUNNEL_EVENT_TASK_FAILED)
        assert len(started) == 1
        assert len(failed) == 1
        assert failed[0].metadata_json["error_code"] == "unknown_intent"


def test_cleide_bi_chat_growth_fail_open(app, ctx, monkeypatch):
    import app.cleide_controlled_chat as controlled_chat

    monkeypatch.setattr(
        controlled_chat,
        "resolve_cleide_ai_flags",
        lambda: SimpleNamespace(
            ai_enabled=False,
            environment="test",
            selected_flag="off",
            reason="test",
            has_api_key=False,
            api_key_label="",
        ),
    )
    monkeypatch.setattr(controlled_chat, "get_cleide_chat_context", lambda _s: _chat_context())
    monkeypatch.setattr(
        controlled_chat,
        "_resolve_authz",
        lambda: {"permitido": True, "modo_operacao": "normal"},
    )
    monkeypatch.setattr(
        "app.funnel_event_service.record_funnel_event",
        lambda **_k: (_ for _ in ()).throw(RuntimeError("analytics down")),
    )

    with app.app_context():
        _auth_user(monkeypatch, email="lote3b-cleide-chat-failopen@test.com")
        with app.test_request_context():
            body, status = run_cleide_controlled_chat(
                question="resumo operacional",
                session_obj={},
            )
        assert status == 200
        assert body.get("reply")
        assert body.get("fallback_used") is False


def test_roberto_bi_started_completed_and_failed(app, ctx, monkeypatch):
    import app.roberto_bi as bi
    from app.upload_handler import SESSION_KEY_UPLOAD_REF, _try_record_roberto_bi_growth_task
    from app.funnel_event_service import FUNNEL_EVENT_TASK_FAILED

    _enable_session(app)
    with app.app_context():
        user = _auth_user(monkeypatch, email="lote3b-roberto-bi@test.com")
        upload_ref = "upload-ref-roberto-bi-1"

        # Falha funcional após preparação (sem linhas válidas) — via helper de upload.
        _try_record_roberto_bi_growth_task(
            event_name=FUNNEL_EVENT_TASK_FAILED,
            idempotency_key=f"growth:roberto_bi:exec-no-rows:task_failed",
            execution_id="exec-no-rows",
            task_stage="upload_processing",
            error_code="roberto_bi_no_valid_rows",
        )
        failed = _events_for(user.id, FUNNEL_EVENT_TASK_FAILED)
        assert len(failed) == 1
        assert failed[0].metadata_json["error_code"] == "roberto_bi_no_valid_rows"
        assert failed[0].source == FUNNEL_SOURCE_ROBERTO_BI

        monkeypatch.setattr(
            bi,
            "get_dados_upload_cliente",
            lambda: [
                {
                    "data_emissao": "2026-01-01",
                    "id_cidade_origem": 1,
                    "id_cidade_destino": 2,
                    "uf_origem": "SP",
                    "uf_destino": "RJ",
                    "peso_real": 10.0,
                    "valor_frete_total": 100.0,
                    "modal": "rodoviario",
                }
            ],
        )
        monkeypatch.setattr(
            "app.upload_handler.get_dados_upload_cliente",
            lambda: [
                {
                    "data_emissao": "2026-01-01",
                    "id_cidade_origem": 1,
                    "id_cidade_destino": 2,
                    "uf_origem": "SP",
                    "uf_destino": "RJ",
                    "peso_real": 10.0,
                    "valor_frete_total": 100.0,
                    "modal": "rodoviario",
                }
            ],
        )
        monkeypatch.setattr(
            bi,
            "_montar_contexto_bi_roberto",
            lambda: {
                "unidos": [{"peso_real": 10}],
                "qualidade_base": {"ok": True},
                "qualidade_previsao": None,
                "resultado_prever": None,
                "recomendacoes_analise": [],
                "serie_temporal": {"meses": [], "valores": []},
            },
        )

        with app.test_request_context("/api/roberto_bi/contexto_analitico"):
            from flask import session

            session[SESSION_KEY_UPLOAD_REF] = upload_ref
            ctx1 = bi.get_contexto_bi_roberto()
            ctx2 = bi.get_contexto_bi_roberto()
            assert ctx1["qualidade_base"]["ok"] is True
            assert ctx2 is ctx1  # memoizacao funcional intacta

        started = _events_for(user.id, FUNNEL_EVENT_TASK_STARTED)
        completed = _events_for(user.id, FUNNEL_EVENT_TASK_COMPLETED)
        assert len(started) == 1
        assert len(completed) == 1
        assert started[0].metadata_json["task_type"] == TASK_TYPE_ROBERTO_BI
        assert completed[0].source == FUNNEL_SOURCE_ROBERTO_BI
        # Mesma execução: started/completed compartilham o mesmo execution_id (UUID).
        assert started[0].execution_id == completed[0].execution_id
        assert started[0].execution_id != upload_ref
        assert started[0].idempotency_key == (
            f"growth:roberto_bi:{upload_ref}:{started[0].execution_id}:task_started"
        )
        assert completed[0].idempotency_key == (
            f"growth:roberto_bi:{upload_ref}:{completed[0].execution_id}:task_completed"
        )
        first_exec_id = started[0].execution_id

        # Nova requisição com o mesmo upload_ref = nova execução analítica legítima.
        with app.test_request_context("/api/roberto_bi/contexto_analitico"):
            from flask import g, session

            # Simula fim do ciclo HTTP anterior (g é app-scoped no harness de teste).
            if hasattr(g, "_roberto_bi_contexto"):
                delattr(g, "_roberto_bi_contexto")
            session[SESSION_KEY_UPLOAD_REF] = upload_ref
            ctx_second = bi.get_contexto_bi_roberto()
            assert ctx_second["qualidade_base"]["ok"] is True

        started_all = _events_for(user.id, FUNNEL_EVENT_TASK_STARTED)
        completed_all = _events_for(user.id, FUNNEL_EVENT_TASK_COMPLETED)
        assert len(started_all) == 2
        assert len(completed_all) == 2
        second_exec_id = started_all[1].execution_id
        assert second_exec_id != first_exec_id
        assert started_all[1].execution_id == completed_all[1].execution_id
        assert started_all[1].idempotency_key != started_all[0].idempotency_key

        # Análise com falha funcional.
        monkeypatch.setattr(
            bi,
            "_montar_contexto_bi_roberto",
            lambda: (_ for _ in ()).throw(RuntimeError("boom")),
        )
        with app.test_request_context("/api/roberto_bi/contexto_analitico"):
            from flask import g, session

            session[SESSION_KEY_UPLOAD_REF] = "upload-ref-roberto-bi-fail"
            if hasattr(g, "_roberto_bi_contexto"):
                delattr(g, "_roberto_bi_contexto")
            with pytest.raises(RuntimeError):
                bi.get_contexto_bi_roberto()

        failed_analysis = [
            e
            for e in _events_for(user.id, FUNNEL_EVENT_TASK_FAILED)
            if (e.metadata_json or {}).get("error_code") == "roberto_bi_analysis_failed"
        ]
        assert len(failed_analysis) == 1
        assert failed_analysis[0].execution_id
        assert failed_analysis[0].idempotency_key == (
            f"growth:roberto_bi:upload-ref-roberto-bi-fail:"
            f"{failed_analysis[0].execution_id}:task_failed"
        )


def test_roberto_bi_growth_fail_open(app, ctx, monkeypatch):
    import app.roberto_bi as bi
    from app.upload_handler import SESSION_KEY_UPLOAD_REF

    _enable_session(app)
    monkeypatch.setattr(
        "app.upload_handler.get_dados_upload_cliente",
        lambda: [{"peso_real": 1, "uf_origem": "SP", "uf_destino": "RJ", "id_cidade_origem": 1, "id_cidade_destino": 2}],
    )
    monkeypatch.setattr(
        bi,
        "_montar_contexto_bi_roberto",
        lambda: {
            "unidos": [],
            "qualidade_base": {},
            "qualidade_previsao": None,
            "resultado_prever": None,
            "recomendacoes_analise": [],
            "serie_temporal": {},
        },
    )
    monkeypatch.setattr(
        "app.funnel_event_service.record_funnel_event",
        lambda **_k: (_ for _ in ()).throw(RuntimeError("analytics down")),
    )

    with app.app_context():
        _auth_user(monkeypatch, email="lote3b-roberto-bi-failopen@test.com")
        with app.test_request_context("/api/roberto_bi/contexto_analitico"):
            from flask import session

            session[SESSION_KEY_UPLOAD_REF] = "upload-ref-failopen"
            ctx_out = bi.get_contexto_bi_roberto()
            assert "qualidade_base" in ctx_out


def test_roberto_chat_started_completed(app, ctx, monkeypatch):
    import app.run_roberto_chat as rc

    monkeypatch.setattr(
        rc,
        "get_contexto_bi_roberto_upload_only",
        lambda: {
            "unidos": [{"peso_real": 10.0, "valor_frete_total": 20.0, "uf_destino": "SP", "modal": "rodoviario"}],
            "serie_temporal": {"meses": ["2024-01"], "valores": [2.0]},
            "qualidade_base": {},
            "recomendacoes_analise": [],
        },
    )
    monkeypatch.setattr(rc, "_register_snapshot_processing_event", lambda **_k: None)
    monkeypatch.setattr(rc, "_get_client", lambda: object())
    monkeypatch.setattr(rc, "_get_model_candidates", lambda: ["gemini-2.5-flash"])
    monkeypatch.setattr(
        rc,
        "heatmap_brasil_upload_only",
        lambda: {"ufs": [], "valores": [], "nivel_temperatura": [], "tendencia_alta": [], "qualidade_uf": []},
    )
    monkeypatch.setattr(
        rc,
        "cleiton_governed_generate_content",
        lambda *a, **k: SimpleNamespace(text="Analise consolidada do Roberto."),
    )

    with app.app_context():
        user = _auth_user(monkeypatch, email="lote3b-roberto-chat-ok@test.com")
        with app.test_request_context():
            out = rc.chat_roberto_reply(
                "Analise custo medio",
                [],
                max_history=5,
                execution_id="exec-roberto-chat-ok",
            )
        assert "Analise consolidada" in out["reply"]

        started = _events_for(user.id, FUNNEL_EVENT_TASK_STARTED)
        completed = _events_for(user.id, FUNNEL_EVENT_TASK_COMPLETED)
        prep = _events_for(user.id, FUNNEL_EVENT_TASK_PREPARATION_STARTED)
        assert len(prep) == 1
        assert len(started) == 1
        assert len(completed) == 1
        assert started[0].source == FUNNEL_SOURCE_ROBERTO_CHAT
        assert started[0].metadata_json["task_type"] == TASK_TYPE_ROBERTO_CHAT
        assert completed[0].execution_id == "exec-roberto-chat-ok"


def test_roberto_chat_functional_failure(app, ctx, monkeypatch):
    import app.run_roberto_chat as rc

    monkeypatch.setattr(rc, "get_contexto_bi_roberto_upload_only", lambda: None)
    monkeypatch.setattr(rc, "_register_snapshot_processing_event", lambda **_k: None)

    with app.app_context():
        user = _auth_user(monkeypatch, email="lote3b-roberto-chat-fail@test.com")
        with app.test_request_context():
            out = rc.chat_roberto_reply(
                "Analise os dados",
                [],
                max_history=5,
                execution_id="exec-roberto-chat-fail",
            )
        assert out.get("requires_upload") is True

        failed = _events_for(user.id, FUNNEL_EVENT_TASK_FAILED)
        assert len(failed) == 1
        assert failed[0].metadata_json["error_code"] == "roberto_chat_requires_upload"
        assert failed[0].source == FUNNEL_SOURCE_ROBERTO_CHAT
        assert FunnelEvent.query.filter_by(user_id=user.id, event_name=FUNNEL_EVENT_TASK_STARTED).count() == 0


def test_roberto_chat_growth_fail_open(app, ctx, monkeypatch):
    import app.run_roberto_chat as rc

    monkeypatch.setattr(
        rc,
        "get_contexto_bi_roberto_upload_only",
        lambda: {
            "unidos": [{"peso_real": 1, "valor_frete_total": 1}],
            "serie_temporal": {},
            "qualidade_base": {},
            "recomendacoes_analise": [],
        },
    )
    monkeypatch.setattr(rc, "_register_snapshot_processing_event", lambda **_k: None)
    monkeypatch.setattr(rc, "_get_client", lambda: object())
    monkeypatch.setattr(rc, "_get_model_candidates", lambda: ["gemini-2.5-flash"])
    monkeypatch.setattr(
        rc,
        "heatmap_brasil_upload_only",
        lambda: {"ufs": [], "valores": [], "nivel_temperatura": [], "tendencia_alta": [], "qualidade_uf": []},
    )
    monkeypatch.setattr(
        rc,
        "cleiton_governed_generate_content",
        lambda *a, **k: SimpleNamespace(text="ok"),
    )
    monkeypatch.setattr(
        "app.funnel_event_service.record_funnel_event",
        lambda **_k: (_ for _ in ()).throw(RuntimeError("analytics down")),
    )

    with app.app_context():
        _auth_user(monkeypatch, email="lote3b-roberto-chat-failopen@test.com")
        with app.test_request_context():
            out = rc.chat_roberto_reply("Pergunta", [], max_history=3, execution_id="exec-failopen")
        assert out["reply"] == "ok"

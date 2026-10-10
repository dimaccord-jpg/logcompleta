"""Fixtures sintéticas do vertical slice. Nenhum arquivo de cliente."""
from __future__ import annotations

import io
import zipfile

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

SYNTHETIC_EMAIL = "operacao@sintetico.example"
SYNTHETIC_CNPJ = "04.252.011/0001-10"
SIDE_NOTE = "Acima de 500 kg aplicar excedente de R$ 0,85/kg"
SUL_COMMENT = "Aplicar mínimo de R$ 40,00"
SUL_MERGE_TEXT = "Região Sul"
FORMULA = "=A3+5"
CACHED_FORMULA_RESULT = 45
NUMERIC_CPF = 39053344705
NUMERIC_CPF_DECIMAL = 52998224725.0
NUMERIC_CNPJ = 10000000000064
NUMERIC_CARD = 4111111111111111
NUMERIC_BANK = 123456
NUMERIC_MONEY = 10.5
NUMERIC_WEIGHT = 7.25
MIXED_PDF_TEXT = "Trecho misto"


def build_tariff_pdf(*, draw_lines: bool = True) -> bytes:
    """Uma página com SP, três subcolunas, valores e nota lateral."""
    ops: list[str] = []

    def op(line: str) -> None:
        ops.append(line)

    def text(x: int, y: int, value: str, size: int = 11) -> None:
        escaped = value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        op("BT")
        op(f"/F1 {size} Tf")
        op(f"1 0 0 1 {x} {y} Tm")
        op(f"({escaped}) Tj")
        op("ET")

    if draw_lines:
        for y in (700, 674, 648, 620):
            op(f"72 {y} m 360 {y} l S")
        for x in (72, 168, 264, 360):
            op(f"{x} 620 m {x} 700 l S")
    text(72, 730, "SP", 16)
    text(90, 680, "Capital")
    text(180, 680, "Interior I")
    text(276, 680, "Interior II")
    text(100, 652, "12,50")
    text(190, 652, "18,00")
    text(286, 652, "22,00")
    text(100, 628, "15,00")
    text(190, 628, "20,00")
    text(286, 628, "25,00")
    text(400, 660, SIDE_NOTE, 9)
    text(72, 80, SYNTHETIC_EMAIL, 10)
    text(280, 80, SYNTHETIC_CNPJ, 10)
    return _pdf_from_stream("\n".join(ops).encode("latin-1"))


def build_visual_only_pdf() -> bytes:
    """Página com desenho e sem texto acessível."""
    stream = b"0.2 0.2 0.2 rg\n80 400 180 60 re\nf\n"
    return _pdf_from_stream(stream)


def build_mixed_visual_pdf() -> bytes:
    """Texto acessível e um retângulo pintado que o slice não interpreta."""
    ops: list[str] = ["0.2 0.4 0.6 rg", "80 400 180 60 re", "f"]
    escaped = MIXED_PDF_TEXT.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    ops.extend(["BT", "/F1 12 Tf", "1 0 0 1 72 700 Tm", f"({escaped}) Tj", "ET"])
    return _pdf_from_stream("\n".join(ops).encode("latin-1"))


def build_precision_xlsx() -> bytes:
    """Floats que o format(..., 'f') arredondava ou zerava, mais um CPF integral."""
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Num"
    rows = (
        ("zero", 0),
        ("meio", 10.5),
        ("peso", 7.25),
        ("pequeno", 0.0000001),
        ("fracao", 1.23456789),
        ("negativo", -0.0000001),
        ("grande", 123456789.123456),
        ("CPF", NUMERIC_CPF_DECIMAL),
    )
    for index, (label, value) in enumerate(rows, start=1):
        sheet.cell(row=index, column=1, value=label)
        sheet.cell(row=index, column=2, value=value)
    bio = io.BytesIO()
    wb.save(bio)
    wb.close()
    return bio.getvalue()


def build_oversized_grid_xlsx(*, sheets: int = 3, cells_per_sheet: int = 4001) -> bytes:
    """Cada aba fica abaixo de 8000 células; o documento inteiro não."""
    wb = Workbook()
    for index in range(sheets):
        sheet = wb.active if index == 0 else wb.create_sheet()
        sheet.title = f"Aba{index + 1}"
        for cell_index in range(cells_per_sheet):
            sheet.cell(row=cell_index + 1, column=1, value=cell_index)
    bio = io.BytesIO()
    wb.save(bio)
    wb.close()
    return bio.getvalue()


def build_giant_merge_xlsx() -> bytes:
    """Um único merge A1:GR200 escrito no XML, sem materializar MergedCell na montagem."""
    return _inject_merge_block(build_sparse_corners_xlsx(), {1: ["A1:GR200"]})


def build_many_small_merges_xlsx(*, sheets: int = 3, merges_per_sheet: int = 1000) -> bytes:
    """Milhares de merges de duas células, distribuídos entre abas."""
    wb = Workbook()
    for index in range(sheets):
        sheet = wb.active if index == 0 else wb.create_sheet()
        sheet.title = f"M{index + 1}"
        sheet["A1"] = "x"
    bio = io.BytesIO()
    wb.save(bio)
    wb.close()
    refs = {
        index: [f"A{row}:B{row}" for row in range(1, merges_per_sheet + 1)]
        for index in range(1, sheets + 1)
    }
    return _inject_merge_block(bio.getvalue(), refs)


def build_isolated_line_pdf() -> bytes:
    """Texto acessível e um traço gráfico longe de qualquer grade."""
    ops = ["40 220 m 520 220 l S"]
    escaped = "Trecho com linha".replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    ops.extend(["BT", "/F1 12 Tf", "1 0 0 1 72 700 Tm", f"({escaped}) Tj", "ET"])
    return _pdf_from_stream("\n".join(ops).encode("latin-1"))


def build_sparse_corners_xlsx() -> bytes:
    """Só A1 e GR200. O retângulo entre elas tem 40 mil coordenadas."""
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Raro"
    sheet["A1"] = "origem"
    sheet["GR200"] = "destino"
    bio = io.BytesIO()
    wb.save(bio)
    wb.close()
    return bio.getvalue()


def build_grouped_hidden_xlsx() -> bytes:
    """Colunas C:F ocultas num único intervalo, sem expandir cada letra."""
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Ocultas"
    sheet["A1"] = "visivel"
    sheet["C1"] = "meio"
    sheet["F1"] = "fim"
    sheet.column_dimensions.group("C", "F", hidden=True)
    bio = io.BytesIO()
    wb.save(bio)
    wb.close()
    return bio.getvalue()


def build_sensitive_numeric_xlsx() -> bytes:
    """Números de célula, inclusive identificadores que o Excel guarda como número."""
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Dados"
    rows = (
        ("CPF", NUMERIC_CPF),
        ("CNPJ", NUMERIC_CNPJ),
        ("Valor", NUMERIC_MONEY),
        ("Peso", NUMERIC_WEIGHT),
        ("Tarifa", 0),
        ("cartao", NUMERIC_CARD),
        ("agencia", NUMERIC_BANK),
    )
    for index, (label, value) in enumerate(rows, start=1):
        sheet.cell(row=index, column=1, value=label)
        sheet.cell(row=index, column=2, value=value)
    bio = io.BytesIO()
    wb.save(bio)
    wb.close()
    return bio.getvalue()


def build_tariff_xlsx() -> bytes:
    """Duas abas sintéticas com merge, comentário, fórmula, zero e vazio."""
    wb = Workbook()
    sp = wb.active
    sp.title = "SP"
    sp["A1"] = "SP"
    sp["A1"].font = Font(bold=True)
    sp["A2"] = "Capital"
    sp["B2"] = "Interior I"
    sp["C2"] = "Interior II"
    sp["A3"] = 10.5
    sp["B3"] = 12
    sp["C3"] = 0
    sp["C3"].number_format = "0.00"
    sp["A5"] = SYNTHETIC_EMAIL
    sp["B5"] = SYNTHETIC_CNPJ
    sp.row_dimensions[6].hidden = True
    sp.column_dimensions["E"].hidden = True

    sul = wb.create_sheet("Sul")
    sul.merge_cells("A1:C1")
    sul["A1"] = SUL_MERGE_TEXT
    sul["A2"] = "SC"
    sul["B2"] = "PR"
    sul["C2"] = "RS"
    sul["A3"] = 40
    sul["B3"] = 41
    sul["B3"].comment = Comment(SUL_COMMENT, "sintetico")
    sul["C3"] = 42
    sul["D3"] = FORMULA

    bio = io.BytesIO()
    wb.save(bio)
    wb.close()
    return _inject_cached_formula(bio.getvalue(), "A3+5", str(CACHED_FORMULA_RESULT))


def _pdf_from_stream(stream: bytes) -> bytes:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Count 1 /Kids [3 0 R] >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Rotate 0 "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode("ascii") + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode("ascii")
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode("ascii")
    out += (
        f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii")
    )
    return bytes(out)


def _inject_merge_block(raw: bytes, refs_by_sheet: dict[int, list[str]]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as src, zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            payload = src.read(info.filename)
            name = info.filename
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                number = int(name.removeprefix("xl/worksheets/sheet").removesuffix(".xml"))
                refs = refs_by_sheet.get(number) or []
                if refs:
                    text = payload.decode("utf-8")
                    block = (
                        f'<mergeCells count="{len(refs)}">'
                        + "".join(f'<mergeCell ref="{ref}"/>' for ref in refs)
                        + "</mergeCells>"
                    )
                    if "</worksheet>" not in text:
                        raise RuntimeError("worksheet_sem_fecho")
                    payload = text.replace("</worksheet>", block + "</worksheet>", 1).encode("utf-8")
            dst.writestr(name, payload)
    return out.getvalue()


def _inject_cached_formula(raw: bytes, formula_body: str, cached: str) -> bytes:
    needle = f"<f>{formula_body}</f>"
    out = io.BytesIO()
    replaced = False
    with zipfile.ZipFile(io.BytesIO(raw)) as src, zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            payload = src.read(info.filename)
            if info.filename.startswith("xl/worksheets/sheet") and info.filename.endswith(".xml"):
                text = payload.decode("utf-8")
                if needle in text:
                    text = _replace_cached_value(text, needle, cached)
                    payload = text.encode("utf-8")
                    replaced = True
            dst.writestr(info.filename, payload)
    if not replaced:
        raise RuntimeError("formula_cache_not_injected")
    return out.getvalue()


_OUTSIDE_MERGE_XML = """<?xml version="1.0" encoding="UTF-8"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <sheetData/>
  <mergeCells count="1"><mergeCell ref="A1:GR200"/></mergeCells>
</worksheet>
""".encode("utf-8")

_ENTITY_WORKSHEET_XML = b"""<?xml version="1.0"?>
<!DOCTYPE worksheet [<!ENTITY x "y">]>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">&x;</worksheet>
"""


def build_small_merge_xlsx() -> bytes:
    """A1, GR200 e um merge pequeno. A área do merge cabe no orçamento."""
    return _inject_merge_block(build_sparse_corners_xlsx(), {1: ["A1:B1"]})


def build_three_long_text_sheets_xlsx() -> bytes:
    """3 abas, uma célula de 3000 caracteres em cada uma."""
    wb = Workbook()
    rows = (("Sa", "alfa " * 600), ("Sb", "beta " * 600), ("Sc", "gama " * 600))
    for index, (name, value) in enumerate(rows):
        sheet = wb.active if index == 0 else wb.create_sheet()
        sheet.title = name
        sheet["A1"] = value
    return _save_workbook(wb)


def build_shared_text_budget_xlsx() -> bytes:
    """Valor, comentário, fórmula e cache no mesmo documento."""
    wb = Workbook()
    sheet = wb.active
    sheet.title = "S"
    sheet["A1"] = "V" * 12
    sheet["A1"].comment = Comment("K" * 12, "sintetico")
    sheet["B1"] = "=" + ("F" * 10)
    bio = io.BytesIO()
    wb.save(bio)
    wb.close()
    return _inject_cached_formula(bio.getvalue(), "F" * 10, "123456789")


def build_repeated_number_xlsx() -> bytes:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "N"
    for row in range(1, 4):
        sheet.cell(row=row, column=1, value=123456789)
    return _save_workbook(wb)


def build_precise_number_xlsx() -> bytes:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "N"
    sheet["A1"] = 1.23456789
    return _save_workbook(wb)


def build_long_sheet_name_xlsx() -> bytes:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "N" * 31
    sheet["A1"] = "C" * 20
    return _save_workbook(wb)


def build_many_sheets_xlsx(*, sheets: int = 51) -> bytes:
    wb = Workbook()
    wb.active.title = "S0"
    for index in range(1, sheets):
        wb.create_sheet(f"S{index}")
        wb.worksheets[-1]["A1"] = index
    return _save_workbook(wb)


def rename_worksheet_part(raw: bytes, new_part: str = "xl/worksheets/custom.xml") -> bytes:
    """Aponta o relacionamento do workbook para um part com nome não sequencial."""
    old_part = "xl/worksheets/sheet1.xml"
    out = io.BytesIO()
    found = False
    with zipfile.ZipFile(io.BytesIO(raw)) as src, zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            payload = src.read(info.filename)
            name = info.filename
            if name == old_part:
                name = new_part
                found = True
            elif name == "xl/_rels/workbook.xml.rels":
                payload = (
                    payload.decode("utf-8")
                    .replace("/xl/worksheets/sheet1.xml", f"/{new_part}")
                    .encode("utf-8")
                )
            elif name == "[Content_Types].xml":
                payload = (
                    payload.decode("utf-8")
                    .replace("/xl/worksheets/sheet1.xml", f"/{new_part}")
                    .encode("utf-8")
                )
            dst.writestr(name, payload)
    if not found:
        raise RuntimeError("worksheet_nao_encontrada")
    return out.getvalue()


def retarget_worksheet(
    raw: bytes,
    new_target: str,
    *,
    old_target: str = "/xl/worksheets/sheet1.xml",
    target_mode: str | None = None,
    duplicate: bool = False,
) -> bytes:
    needle = f'Target="{old_target}"'
    replacement = f'Target="{new_target}"'
    if target_mode:
        replacement += f' TargetMode="{target_mode}"'
    out = io.BytesIO()
    replaced = False
    with zipfile.ZipFile(io.BytesIO(raw)) as src, zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            payload = src.read(info.filename)
            if info.filename == "xl/_rels/workbook.xml.rels":
                text = payload.decode("utf-8")
                if needle not in text:
                    raise RuntimeError("rel_worksheet_ausente")
                text = text.replace(needle, replacement, 1)
                if duplicate:
                    text = _duplicate_first_relationship(text)
                payload = text.encode("utf-8")
                replaced = True
            dst.writestr(info.filename, payload)
    if not replaced:
        raise RuntimeError("rels_ausente")
    return out.getvalue()


def replace_zip_member(raw: bytes, name: str, payload: bytes) -> bytes:
    out = io.BytesIO()
    found = False
    with zipfile.ZipFile(io.BytesIO(raw)) as src, zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = payload if info.filename == name else src.read(info.filename)
            if info.filename == name:
                found = True
            dst.writestr(info.filename, body)
        if not found:
            dst.writestr(name, payload)
    return out.getvalue()


def duplicate_worksheet_target(raw: bytes) -> bytes:
    """Duas sheets apontam para o mesmo part, com Ids distintos."""
    rels_extra = (
        '<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
        'Target="/xl/worksheets/sheet1.xml" Id="rId9"/>'
    )
    sheet_extra = (
        '<sheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'name="Copia" sheetId="9" state="visible" r:id="rId9"/>'
    )
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as src, zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            payload = src.read(info.filename)
            if info.filename == "xl/_rels/workbook.xml.rels":
                text = payload.decode("utf-8")
                text = text.replace("</Relationships>", rels_extra + "</Relationships>", 1)
                payload = text.encode("utf-8")
            elif info.filename == "xl/workbook.xml":
                text = payload.decode("utf-8")
                text = text.replace("</sheets>", sheet_extra + "</sheets>", 1)
                payload = text.encode("utf-8")
            dst.writestr(info.filename, payload)
    return out.getvalue()


_OOXML_WORKSHEET_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"


def alias_worksheet_target(raw: bytes, *, total_sheets: int) -> bytes:
    """Várias entradas de sheet, cada uma com r:id próprio, todas no mesmo part."""
    if total_sheets < 1:
        raise ValueError("total_sheets")
    extra = total_sheets - 1
    if extra == 0:
        return raw
    rels: list[str] = []
    sheets: list[str] = []
    for index in range(1, extra + 1):
        rel_id = f"rIdAlias{index}"
        rels.append(
            "<Relationship "
            f'Type="{_OOXML_WORKSHEET_REL}" '
            'Target="/xl/worksheets/sheet1.xml" '
            f'Id="{rel_id}"/>'
        )
        sheets.append(
            '<sheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
            f'name="Copia{index}" sheetId="{index + 1}" state="visible" r:id="{rel_id}"/>'
        )
    return _patch_workbook_package(raw, rels="".join(rels), sheets="".join(sheets))


def build_aliased_merge_xlsx(*, sheets: int = 51, merges: int = 100) -> bytes:
    """Um worksheet com N merges, referenciado por `sheets` entradas do workbook."""
    base = _inject_merge_block(
        build_sparse_corners_xlsx(),
        {1: [f"A{row}:B{row}" for row in range(1, merges + 1)]},
    )
    return alias_worksheet_target(base, total_sheets=sheets)


def remove_sheet_relationship_id(raw: bytes, rel_id: str = "rId1") -> bytes:
    needle = f' r:id="{rel_id}"'
    return _rewrite_member_text(raw, "xl/workbook.xml", lambda text: _replace_once(text, needle, ""))


def rewrite_sheet_relationship_id(raw: bytes, old_id: str, new_id: str) -> bytes:
    return _rewrite_member_text(
        raw,
        "xl/workbook.xml",
        lambda text: _replace_once(text, f'r:id="{old_id}"', f'r:id="{new_id}"'),
    )


def retype_worksheet_relationship(raw: bytes, new_type: str) -> bytes:
    needle = f'Type="{_OOXML_WORKSHEET_REL}"'
    return _rewrite_member_text(
        raw,
        "xl/_rels/workbook.xml.rels",
        lambda text: _replace_once(text, needle, f'Type="{new_type}"'),
    )


def attach_orphan_chartsheet(raw: bytes) -> bytes:
    """Chartsheet fora da coleção de sheets. Não vira worksheet."""
    rel = (
        '<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chartsheet" '
        'Target="/xl/chartsheets/chartsheet1.xml" Id="rIdChart"/>'
    )
    content_type = (
        '<Override PartName="/xl/chartsheets/chartsheet1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.chartsheet+xml"/>'
    )
    chart = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<chartsheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        "<sheetPr/></chartsheet>"
    ).encode("utf-8")
    patched = _patch_workbook_package(raw, rels=rel, content_type=content_type)
    return replace_zip_member(patched, "xl/chartsheets/chartsheet1.xml", chart)


def build_traversal_worksheet_xlsx() -> bytes:
    retargeted = retarget_worksheet(build_sparse_corners_xlsx(), "../outside.xml")
    return replace_zip_member(retargeted, "outside.xml", _OUTSIDE_MERGE_XML)


def build_entity_worksheet_xlsx(*, renamed: bool = False) -> bytes:
    data = build_sparse_corners_xlsx()
    part = "xl/worksheets/sheet1.xml"
    if renamed:
        data = rename_worksheet_part(data)
        part = "xl/worksheets/custom.xml"
    return replace_zip_member(data, part, _ENTITY_WORKSHEET_XML)


def _save_workbook(wb: Workbook) -> bytes:
    bio = io.BytesIO()
    wb.save(bio)
    wb.close()
    return bio.getvalue()


def _patch_workbook_package(
    raw: bytes,
    *,
    rels: str = "",
    sheets: str = "",
    content_type: str = "",
) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as src, zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            payload = src.read(info.filename)
            if rels and info.filename == "xl/_rels/workbook.xml.rels":
                payload = _replace_once(payload.decode("utf-8"), "</Relationships>", rels + "</Relationships>").encode(
                    "utf-8"
                )
            elif sheets and info.filename == "xl/workbook.xml":
                payload = _replace_once(payload.decode("utf-8"), "</sheets>", sheets + "</sheets>").encode("utf-8")
            elif content_type and info.filename == "[Content_Types].xml":
                payload = _replace_once(payload.decode("utf-8"), "</Types>", content_type + "</Types>").encode("utf-8")
            dst.writestr(info.filename, payload)
    return out.getvalue()


def _rewrite_member_text(raw: bytes, name: str, edit) -> bytes:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        if name not in archive.namelist():
            raise RuntimeError("membro_ausente")
        updated = edit(archive.read(name).decode("utf-8")).encode("utf-8")
    return replace_zip_member(raw, name, updated)


def _replace_once(text: str, needle: str, replacement: str) -> str:
    if needle not in text:
        raise RuntimeError("trecho_ausente")
    return text.replace(needle, replacement, 1)


def _duplicate_first_relationship(text: str) -> str:
    start = text.find("<Relationship ")
    end = text.find("/>", start)
    if start < 0 or end < 0:
        raise RuntimeError("rel_ausente")
    element = text[start : end + 2]
    return text.replace("</Relationships>", element + "</Relationships>", 1)


def _replace_cached_value(xml: str, needle: str, cached: str) -> str:
    idx = xml.find(needle)
    after = idx + len(needle)
    empty = "<v></v>"
    if xml.startswith(empty, after):
        return xml[:after] + f"<v>{cached}</v>" + xml[after + len(empty) :]
    if xml.startswith("<v>", after):
        end = xml.find("</v>", after)
        return xml[:after] + f"<v>{cached}</v>" + xml[end + 4 :]
    return xml[:after] + f"<v>{cached}</v>" + xml[after:]
